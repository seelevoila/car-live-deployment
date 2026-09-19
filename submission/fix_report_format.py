from pathlib import Path
from datetime import datetime
import shutil
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '2026-chongqing-ai-competition'
NAMES = ['01-系统设计文档', '02-汽车垂直领域RAG知识库构建方案', '03-流式TTS动态改稿实现方案', '04-语音克隆拟人化优化方案', '05-标准化功能测试报告']


def child(parent, tag, **attrs):
    el = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        el.set(qn('w:' + key), str(value))
    parent.append(el)
    return el


def black_rpr(rpr):
    for el in list(rpr):
        if el.tag.split('}')[-1] in ('color', 'textFill', 'textOutline', 'shadow', 'glow', 'reflection'):
            rpr.remove(el)
    child(rpr, 'color', val='000000')


def black_all(doc):
    for part in doc.part.package.parts:
        root = getattr(part, '_element', None)
        if root is None:
            continue
        for rpr in root.iter(qn('w:rPr')):
            black_rpr(rpr)
        for run in root.iter(qn('w:r')):
            rpr = run.find(qn('w:rPr'))
            if rpr is None:
                rpr = OxmlElement('w:rPr')
                run.insert(0, rpr)
            black_rpr(rpr)
        for color in root.iter(qn('w:color')):
            color.attrib.clear()
            color.set(qn('w:val'), '000000')


def revision_table(doc):
    heading = next(p for p in doc.paragraphs if p.text.strip() == '文档修订记录')
    old = heading._p.getnext()
    while old is not None and old.tag != qn('w:tbl'):
        old = old.getnext()
    assert old is not None
    from docx.table import Table
    source = Table(old, doc)
    rows = [[c.text for c in r.cells] for r in source.rows]
    values = []
    for r in rows[1:]:
        if len(r) == 4:
            values.append([str(len(values) + 1), r[0], r[1], r[2], '', '', r[3]])
        else:
            values.append(r[:7])
    while len(values) < 6:
        values.append([''] * 7)
    table = doc.add_table(rows=1 + len(values), cols=7)
    for row, content in zip(table.rows, [['序号', '版本号', '修改时间', '修改人', '审核人', '批准人', '备注']] + values):
        for cell, text in zip(row.cells, content):
            cell.text = text
    old.addprevious(table._tbl)
    old.getparent().remove(old)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Pt(0)
    for r in heading.runs:
        r.font.name = 'SimSun'
        r.font.size = Pt(16)
        r.bold = False
        r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'SimSun')
    return table._tbl


def table_format(table, width, revision):
    grid = table._tbl.tblGrid
    original = [c.w for c in grid.gridCol_lst]
    weights = [1.4, 1.8, 2.6, 2, 2, 2, 4.4] if revision else original
    total = int(width / 635)
    widths = [round(total * float(v) / sum(weights)) for v in weights]
    widths[-1] += total - sum(widths)
    # Reset inherited table and cell border overrides to one continuous grid.
    pr = table._tbl.tblPr
    for node in list(pr):
        pr.remove(node)
    child(pr, 'tblW', w=total, type='dxa')
    child(pr, 'jc', val='center')
    borders = child(pr, 'tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        child(borders, edge, val='single', sz=4, color='000000', space=0)
    child(pr, 'tblLayout', type='fixed')
    for col, size in zip(grid.gridCol_lst, widths):
        col.set(qn('w:w'), str(size))
    for idx, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        for node in list(trpr):
            trpr.remove(node)
        if idx == 0:
            child(trpr, 'tblHeader')
        child(trpr, 'cantSplit')
        if revision:
            child(trpr, 'trHeight', val=480, hRule='atLeast')
        position = 0
        for tc in row._tr.tc_lst:
            tcpr = tc.get_or_add_tcPr()
            span = tcpr.find(qn('w:gridSpan'))
            count = int(span.get(qn('w:val'))) if span is not None else 1
            for node in list(tcpr):
                if node.tag not in (qn('w:gridSpan'), qn('w:vMerge')):
                    tcpr.remove(node)
            child(tcpr, 'tcW', w=sum(widths[position:position + count]), type='dxa')
            position += count
            if revision and idx == 0:
                child(tcpr, 'shd', val='clear', color='auto', fill='E7E7E7')
            margins = child(tcpr, 'tcMar')
            for edge, size in [('top', 70), ('bottom', 70), ('left', 100), ('right', 100)]:
                child(margins, edge, w=size, type='dxa')
            child(tcpr, 'vAlign', val='center')
            for p in tc.iter(qn('w:p')):
                ppr = p.find(qn('w:pPr'))
                if ppr is None:
                    ppr = OxmlElement('w:pPr')
                    p.insert(0, ppr)
                for node in list(ppr):
                    ppr.remove(node)
                child(ppr, 'spacing', before=0, after=0, line=260, lineRule='auto')
                child(ppr, 'ind', left=0, right=0, firstLine=0)
                child(ppr, 'jc', val='center' if revision and idx > 0 and position <= 6 else 'left')
                for r in p.iter(qn('w:r')):
                    rpr = r.find(qn('w:rPr'))
                    if rpr is None:
                        rpr = OxmlElement('w:rPr')
                        r.insert(0, rpr)
                    for node in list(rpr):
                        if node.tag.split('}')[-1] in ('sz','szCs','rFonts','b','bCs'):
                            rpr.remove(node)
                    child(rpr, 'rFonts', ascii='Times New Roman', hAnsi='Times New Roman', eastAsia='SimSun')
                    child(rpr, 'sz', val=21 if len(widths) < 6 or revision else 19)
                    child(rpr, 'b', val=1 if idx == 0 and not revision else 0)


def process(name, backup):
    path = OUT / (name + '.docx')
    shutil.copy2(path, backup / path.name)
    doc = Document(path)
    revised = revision_table(doc)
    sec = doc.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    for t in doc.tables:
        table_format(t, width, t._tbl is revised)
    for part in doc.part.package.parts:
        if str(part.partname).startswith('/word/header'):
            root = part._element
            for t in root.iter(qn('w:t')):
                t.text = (t.text or '').replace('LOGO', '').strip()
            for ppr in root.iter(qn('w:pPr')):
                ind = ppr.find(qn('w:ind'))
                if ind is not None:
                    ppr.remove(ind)
                child(ppr, 'ind', left=0, right=0, firstLine=0)
            for border in root.iter(qn('w:bottom')):
                border.attrib.clear()
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '4')
                border.set(qn('w:color'), '000000')
    black_all(doc)
    doc.save(path)
    shutil.copy2(path, OUT / (name + '-完整版.docx'))
    print(name)


if __name__ == '__main__':
    backup = ROOT / ('_before_black_grid_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
    backup.mkdir()
    for name in NAMES:
        process(name, backup)
