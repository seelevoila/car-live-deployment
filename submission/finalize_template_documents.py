from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "2026-chongqing-ai-competition"

DOC_NAMES = [
    "01-系统设计文档-完整版.docx",
    "02-汽车垂直领域RAG知识库构建方案-完整版.docx",
    "03-流式TTS动态改稿实现方案-完整版.docx",
    "04-语音克隆拟人化优化方案-完整版.docx",
    "05-标准化功能测试报告-完整版.docx",
]

BLUE = "2F75B5"
LIGHT_GRAY = "E7E6E6"
GRID = "808080"
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, name: str, size: float, bold: bool | None = None, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def set_style_font(style, name: str, size: float, bold: bool = False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")


def set_paragraph_border(paragraph, color=BLUE, size="8", space="4"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = pbdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pbdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def clear_container(container):
    element = container._element
    for child in list(element):
        element.remove(child)


def add_page_field(run):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_toc_field(paragraph):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在 Word 中更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, placeholder, end])
    set_run_font(run, "SimSun", 10.5)


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="4"):
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def mark_header_row(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def style_table(table, revision=False):
    table.style = "Table Grid"
    set_table_borders(table)
    if table.rows:
        mark_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, LIGHT_GRAY if revision and row_index == 0 else "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    set_run_font(run, "SimSun", 10.5, bold=revision and row_index == 0)


def set_update_fields(doc):
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def first_heading(doc):
    for p in doc.paragraphs:
        if p.style and p.style.name == "Heading 1":
            return p
    raise ValueError("no Heading 1 found")


def insert_front_matter(doc):
    if any(p.text.strip() == "文档修订记录" for p in doc.paragraphs):
        return
    anchor = first_heading(doc)
    page_break_1 = doc.add_paragraph()
    page_break_1.add_run().add_break(WD_BREAK.PAGE)
    revision = doc.add_paragraph("文档修订记录", style="Heading 1")
    revision.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=7)
    headers = ["序号", "版本号", "修改时间", "修改人", "审核人", "批准人", "备注"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    row = table.add_row()
    values = ["1", "V1.0", "2026-09-15", "项目组", "", "", "按详细设计说明书模板统一整理，补充目录、修订记录和版式。"]
    for cell, value in zip(row.cells, values):
        cell.text = value
    for _ in range(5):
        table.add_row()
    style_table(table, revision=True)
    page_break_2 = doc.add_paragraph()
    page_break_2.add_run().add_break(WD_BREAK.PAGE)
    toc_heading = doc.add_paragraph("目录", style="Heading 1")
    toc_heading.paragraph_format.keep_with_next = True
    toc = doc.add_paragraph()
    toc.paragraph_format.line_spacing = 1.15
    add_toc_field(toc)
    page_break_3 = doc.add_paragraph()
    page_break_3.add_run().add_break(WD_BREAK.PAGE)
    nodes = [
        page_break_1._p,
        revision._p,
        table._tbl,
        page_break_2._p,
        toc_heading._p,
        toc._p,
        page_break_3._p,
    ]
    for node in nodes:
        anchor._p.addprevious(node)


def update_revision_table(doc):
    for table in doc.tables:
        if not table.rows:
            continue
        first_row = [cell.text.strip() for cell in table.rows[0].cells]
        if not first_row or first_row[0] not in ("版本", "序号"):
            continue
        while len(table.columns) < 7:
            table.add_column(Cm(2.0))
        while len(table.rows) < 7:
            table.add_row()
        headers = ["序号", "版本号", "修改时间", "修改人", "审核人", "批准人", "备注"]
        values = ["1", "V1.0", "2026-09-15", "项目组", "", "", "按详细设计说明书模板统一整理，补充目录、修订记录和版式。"]
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, values):
            cell.text = value
        for row in table.rows[2:]:
            for cell in row.cells:
                cell.text = ""
        style_table(table, revision=True)
        return


def exclude_front_matter_from_toc(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() not in ("文档修订记录", "目录"):
            continue
        paragraph.style = doc.styles["Normal"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(15)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, "SimHei", 16, bold=True)
        ppr = paragraph._p.get_or_add_pPr()
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            ppr.append(outline)
        outline.set(qn("w:val"), "9")


def style_header_footer(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

        header = section.header
        clear_container(header)
        hp = header.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        run = hp.add_run("汽车直播智能体项目")
        set_run_font(run, "SimHei", 9.5, bold=True, color=BLACK)
        set_paragraph_border(hp)

        footer = section.footer
        clear_container(footer)
        fp = footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(0)
        run = fp.add_run("第 ")
        set_run_font(run, "SimSun", 9)
        page_run = fp.add_run()
        add_page_field(page_run)
        set_run_font(page_run, "SimSun", 9)
        run = fp.add_run(" 页")
        set_run_font(run, "SimSun", 9)


def style_document_text(doc):
    styles = doc.styles
    set_style_font(styles["Normal"], "SimSun", 10.5)
    normal = styles["Normal"]
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    for name, size, before, after in (
        ("Heading 1", 16, 15, 8),
        ("Heading 2", 13.5, 10, 5),
        ("Heading 3", 11.5, 7, 3),
    ):
        if name in styles:
            style = styles[name]
            set_style_font(style, "SimHei", size, bold=True)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.first_line_indent = Cm(0)

    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""
        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(65)
            paragraph.paragraph_format.space_after = Pt(22)
            for run in paragraph.runs:
                set_run_font(run, "SimHei", 15, bold=True)
        elif index == 1:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(16)
            for run in paragraph.runs:
                set_run_font(run, "SimHei", 25, bold=True)
        elif index == 2:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(22)
            for run in paragraph.runs:
                set_run_font(run, "SimSun", 11)
        elif style_name == "Heading 1":
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, "SimHei", 16, bold=True)
        elif style_name == "Heading 2":
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, "SimHei", 13.5, bold=True)
        elif style_name == "Heading 3":
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, "SimHei", 11.5, bold=True)
        elif text.startswith("图") or text.startswith("表"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                set_run_font(run, "SimSun", 9.5)
        else:
            if style_name not in ("No Spacing", "Normal"):
                paragraph.paragraph_format.first_line_indent = Cm(0)
            for run in paragraph.runs:
                set_run_font(run, "SimSun", 10.5)

    for table in doc.tables:
        style_table(table)


def trim_trailing_empty_paragraphs(doc):
    body = doc.element.body
    while True:
        children = list(body.iterchildren())
        if len(children) < 2 or children[-1].tag != qn("w:sectPr"):
            return
        candidate = children[-2]
        if candidate.tag != qn("w:p"):
            return
        paragraph = Paragraph(candidate, doc)
        if paragraph.text.strip():
            return
        body.remove(candidate)


def replace_heading_prefixes(doc, replacements: Iterable[tuple[str, str]]):
    replacements = list(replacements)
    for paragraph in doc.paragraphs:
        if not paragraph.style or not paragraph.style.name.startswith("Heading"):
            continue
        stripped = paragraph.text.strip()
        for old, new in replacements:
            if stripped == old or stripped.startswith(old + " "):
                suffix = stripped[len(old):]
                paragraph.text = new + suffix
                break


def prepare_system(doc):
    replace_heading_prefixes(doc, [("2.5.5 易用性、灵活性和经济性", "2.5.5 易用性")])
    if not any(p.text.strip().startswith("2.5.6 ") for p in doc.paragraphs):
        target = next(p for p in doc.paragraphs if p.text.strip().startswith("3 系统功能模块详细设计"))
        h = doc.add_paragraph("2.5.6 灵活性和可扩展性", style="Heading 3")
        b = doc.add_paragraph("模块通过 HTTP/JSON、模型配置和版本化数据结构保留替换空间，便于后续扩展多模型、多音色和多终端能力。")
        target._p.addprevious(b._p)
        target._p.addprevious(h._p)
    if not any(p.text.strip().startswith("2.5.7 ") for p in doc.paragraphs):
        target = next(p for p in doc.paragraphs if p.text.strip().startswith("3 系统功能模块详细设计"))
        h = doc.add_paragraph("2.5.7 经济性和投资保护", style="Heading 3")
        b = doc.add_paragraph("优先复用现有本地设备、开源组件和已有数据，控制部署成本，同时通过标准接口保护既有配置、索引和测试资产。")
        target._p.addprevious(b._p)
        target._p.addprevious(h._p)
    replace_heading_prefixes(doc, [
        ("系统设计文档扩充内容", "B.0 扩充内容说明"),
        ("3.4", "B.1"), ("3.5", "B.2"), ("3.6", "B.3"),
        ("3.7", "B.4"), ("3.8", "B.5"),
    ])


def prepare_rag(doc):
    replace_heading_prefixes(doc, [
        ("1  场景与问题定义", "1 引言"),
        ("2  数据处理链路", "2 知识库总体设计"),
        ("3  检索策略", "3 数据处理与检索设计"),
        ("4  版本、溯源与管理", "4 数据结构、版本与溯源"),
        ("5  测试设计与结果", "5 性能与测试设计"),
        ("6  现场演示脚本", "6 运行与演示规定"),
        ("参考资料", "附录 A 参考资料"),
        ("扩充内容", "附录 B 扩充内容"),
        ("RAG知识库构建方案扩充内容", "B.0 扩充内容说明"),
        ("2.5", "B.1"), ("2.6", "B.2"), ("2.7", "B.3"),
        ("2.8", "B.4"), ("2.9", "B.5"),
    ])


def prepare_tts(doc):
    replace_heading_prefixes(doc, [
        ("1  需求拆解", "1 引言"),
        ("2  播报时序", "2 播报总体设计"),
        ("3  动态改稿算法", "3 动态改稿与状态设计"),
        ("4  电子嗒声与卡顿治理", "4 性能设计"),
        ("5  性能与验收", "5 接口与数据结构"),
        ("6  现场操作与故障回退", "6 运行、故障与回退"),
        ("参考资料", "附录 A 参考资料"),
        ("扩充内容", "附录 B 扩充内容"),
        ("流式TTS动态改稿方案扩充内容", "B.0 扩充内容说明"),
        ("2.5", "B.1"), ("2.6", "B.2"), ("2.7", "B.3"),
        ("2.8", "B.4"), ("2.9", "B.5"),
    ])


def prepare_voice(doc):
    replace_heading_prefixes(doc, [
        ("1  参考音频规范", "1 引言"),
        ("2  样本质量诊断", "2 音色克隆总体设计"),
        ("3  音色 profile 与稳定性", "3 参考音频质量与质量门禁"),
        ("4  多参考音频策略", "4 音色 Profile 与数据结构"),
        ("5  与流式播报打通", "5 流式播报接口"),
        ("6  当前验证与主观评测", "6 性能与评测"),
        ("7  合规与风险", "7 合规、错误处理与运行规定"),
        ("参考资料", "附录 A 参考资料"),
        ("扩充内容", "附录 B 扩充内容"),
        ("语音克隆拟人化优化方案扩充内容", "B.0 扩充内容说明"),
        ("2.5", "B.1"), ("2.6", "B.2"), ("2.7", "B.3"),
        ("2.8", "B.4"), ("2.9", "B.5"),
    ])


def prepare_test(doc):
    replace_heading_prefixes(doc, [
        ("参考资料", "附录 A 参考资料"),
        ("扩充内容", "附录 B 扩充内容"),
        ("标准化功能测试报告扩充内容", "B.0 扩充内容说明"),
        ("4.", "B.1"), ("5.", "B.2"), ("6.", "B.3"),
        ("7.", "B.4"), ("8.", "B.5"),
    ])


def process(path: Path):
    doc = Document(path)
    if path.name.startswith("01-"):
        prepare_system(doc)
    elif path.name.startswith("02-"):
        prepare_rag(doc)
    elif path.name.startswith("03-"):
        prepare_tts(doc)
    elif path.name.startswith("04-"):
        prepare_voice(doc)
    else:
        prepare_test(doc)
    insert_front_matter(doc)
    update_revision_table(doc)
    style_header_footer(doc)
    style_document_text(doc)
    exclude_front_matter_from_toc(doc)
    trim_trailing_empty_paragraphs(doc)
    set_update_fields(doc)
    doc.save(path)
    print(path.name)


if __name__ == "__main__":
    for name in DOC_NAMES:
        process(OUT / name)
