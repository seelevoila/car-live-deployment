# -*- coding: utf-8 -*-
"""
自动化整合竞赛提交材料
将所有扩充内容整合到DOCX文档中
"""
import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 路径配置
SUBMISSION_DIR = Path(__file__).parent
ROOT_DIR = SUBMISSION_DIR.parent
MATERIALS_DIR = SUBMISSION_DIR / "2026-chongqing-ai-competition"

# 扩充内容文件
EXPANDED_CONTENTS = {
    "系统设计": SUBMISSION_DIR / "系统设计文档扩充.md",
    "RAG方案": SUBMISSION_DIR / "RAG方案扩充.md",
    "TTS方案": SUBMISSION_DIR / "TTS方案扩充.md",
    "测试报告": SUBMISSION_DIR / "测试报告扩充.md",
    "语音克隆": SUBMISSION_DIR / "语音克隆方案扩充.md"
}

# 目标DOCX文件
TARGET_DOCS = {
    "系统设计": MATERIALS_DIR / "01-系统设计文档.docx",
    "RAG方案": MATERIALS_DIR / "02-汽车垂直领域RAG知识库构建方案.docx",
    "TTS方案": MATERIALS_DIR / "03-流式TTS动态改稿实现方案.docx",
    "语音克隆": MATERIALS_DIR / "04-语音克隆拟人化优化方案.docx",
    "测试报告": MATERIALS_DIR / "05-标准化功能测试报告.docx"
}

def read_markdown(file_path):
    """读取Markdown文件"""
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()

def parse_markdown_sections(content):
    """解析Markdown内容为章节"""
    lines = content.split('\n')
    sections = []
    current_section = None
    current_content = []

    for line in lines:
        # 检测标题
        if line.startswith('#'):
            # 保存上一节
            if current_section:
                sections.append({
                    'level': current_section['level'],
                    'title': current_section['title'],
                    'content': '\n'.join(current_content)
                })

            # 开始新节
            level = len(line) - len(line.lstrip('#'))
            title = line.lstrip('#').strip()
            current_section = {'level': level, 'title': title}
            current_content = []
        else:
            current_content.append(line)

    # 保存最后一节
    if current_section:
        sections.append({
            'level': current_section['level'],
            'title': current_section['title'],
            'content': '\n'.join(current_content)
        })

    return sections

def add_markdown_to_docx(doc, markdown_content):
    """将Markdown内容添加到DOCX文档"""
    sections = parse_markdown_sections(markdown_content)

    for section in sections:
        level = section['level']
        title = section['title']
        content = section['content'].strip()

        # 跳过元信息标题
        if title in ["新增章节", "扩充内容"]:
            continue

        # 添加标题
        if level <= 3:
            heading = doc.add_heading(title, level=level)
            heading.style.font.name = 'Microsoft YaHei'
            heading.style.font.size = Pt(18 - level * 2)

        # 处理内容
        if not content:
            continue

        # 检测表格
        if '|' in content and content.count('|') > 4:
            add_table_from_markdown(doc, content)
        # 检测代码块
        elif '```' in content:
            add_code_blocks(doc, content)
        # 普通文本
        else:
            add_formatted_text(doc, content)

def add_table_from_markdown(doc, content):
    """从Markdown表格创建DOCX表格"""
    lines = [line.strip() for line in content.split('\n') if '|' in line]
    if len(lines) < 2:
        return

    # 解析表头
    header_line = lines[0]
    headers = [cell.strip() for cell in header_line.split('|') if cell.strip()]

    # 跳过分隔线
    data_lines = [line for line in lines[2:] if not line.startswith('|--')]

    if not data_lines:
        return

    # 解析数据行
    rows = []
    for line in data_lines:
        cells = [cell.strip() for cell in line.split('|') if cell.strip()]
        if len(cells) == len(headers):
            rows.append(cells)

    if not rows:
        return

    # 创建表格
    table = doc.add_table(rows=len(rows)+1, cols=len(headers))
    table.style = 'Light Grid Accent 1'

    # 填充表头
    header_row = table.rows[0]
    for i, header_text in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = header_text
        # 表头加粗
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.name = 'Microsoft YaHei'
                run.font.size = Pt(10)

    # 填充数据
    for row_idx, row_data in enumerate(rows):
        table_row = table.rows[row_idx + 1]
        for col_idx, cell_text in enumerate(row_data):
            cell = table_row.cells[col_idx]
            cell.text = cell_text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Microsoft YaHei'
                    run.font.size = Pt(9)

    doc.add_paragraph()  # 表格后空行

def add_code_blocks(doc, content):
    """添加代码块"""
    parts = content.split('```')
    for i, part in enumerate(parts):
        if i % 2 == 0:
            # 普通文本
            if part.strip():
                add_formatted_text(doc, part)
        else:
            # 代码块
            lines = part.split('\n')
            # 跳过语言标识行
            code_lines = lines[1:] if lines else []
            code_text = '\n'.join(code_lines)

            if code_text.strip():
                p = doc.add_paragraph(code_text)
                p.style = 'No Spacing'
                for run in p.runs:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                # 背景色（通过shading）
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn
                shading_elm = OxmlElement('w:shd')
                shading_elm.set(qn('w:fill'), 'F0F0F0')
                p._element.get_or_add_pPr().append(shading_elm)

def add_formatted_text(doc, text):
    """添加格式化文本"""
    lines = text.split('\n')

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # 检测列表项
        if line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(line[2:], style='List Bullet')
        elif line[0:2].replace('.', '').replace('、', '').isdigit():
            p = doc.add_paragraph(line.split('.', 1)[-1].split('、', 1)[-1].strip(), style='List Number')
        else:
            p = doc.add_paragraph(line)

        # 设置字体
        for run in p.runs:
            run.font.name = 'Microsoft YaHei'
            run.font.size = Pt(11)

        # 处理加粗
        if '**' in line:
            parts = line.split('**')
            p.clear()
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.font.name = 'Microsoft YaHei'
                run.font.size = Pt(11)
                if i % 2 == 1:
                    run.font.bold = True

def integrate_document(doc_type):
    """整合单个文档"""
    print(f"\n正在整合: {doc_type}")

    # 读取扩充内容
    expanded_file = EXPANDED_CONTENTS.get(doc_type)
    if not expanded_file or not expanded_file.exists():
        print(f"  ⚠️  扩充内容不存在: {expanded_file}")
        return False

    markdown_content = read_markdown(expanded_file)

    # 打开目标文档
    target_doc = TARGET_DOCS.get(doc_type)
    if not target_doc or not target_doc.exists():
        print(f"  ⚠️  目标文档不存在: {target_doc}")
        return False

    doc = Document(target_doc)

    # 添加分隔符
    doc.add_page_break()
    doc.add_heading('扩充内容', level=1)

    # 添加扩充内容
    add_markdown_to_docx(doc, markdown_content)

    # 保存
    output_file = target_doc.parent / f"{target_doc.stem}-扩充版.docx"
    doc.save(output_file)

    print(f"  ✓ 已保存: {output_file}")
    return True

def generate_summary_report():
    """生成整合报告"""
    print("\n" + "="*60)
    print("竞赛材料自动整合完成")
    print("="*60)

    print("\n📁 已生成文件:")
    for doc_type in TARGET_DOCS.keys():
        output_file = TARGET_DOCS[doc_type].parent / f"{TARGET_DOCS[doc_type].stem}-扩充版.docx"
        if output_file.exists():
            size_kb = output_file.stat().st_size / 1024
            print(f"  ✓ {output_file.name} ({size_kb:.1f} KB)")

    print("\n📊 内容统计:")
    for name, file_path in EXPANDED_CONTENTS.items():
        if file_path.exists():
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
                lines = len([l for l in content.split('\n') if l.strip()])
                chars = len(content)
                print(f"  • {name}: {lines}行, {chars}字符")

    print("\n📋 下一步操作:")
    print("  1. 打开生成的'-扩充版.docx'文档")
    print("  2. 调整格式和排版")
    print("  3. 插入架构图（从 架构图和流程图.md 生成）")
    print("  4. 另存为PDF")
    print("  5. 制作PPT（参考 优化计划.md）")

def main():
    """主函数"""
    print("="*60)
    print("竞赛提交材料自动整合工具")
    print("="*60)

    # 检查扩充内容文件
    print("\n检查扩充内容文件...")
    missing_files = []
    for name, file_path in EXPANDED_CONTENTS.items():
        if file_path.exists():
            print(f"  ✓ {name}: {file_path.name}")
        else:
            print(f"  ✗ {name}: 文件不存在")
            missing_files.append(name)

    if missing_files:
        print(f"\n⚠️  缺少扩充内容文件，无法继续")
        return

    # 检查目标文档
    print("\n检查目标DOCX文档...")
    missing_docs = []
    for name, file_path in TARGET_DOCS.items():
        if file_path.exists():
            print(f"  ✓ {name}: {file_path.name}")
        else:
            print(f"  ✗ {name}: 文件不存在")
            missing_docs.append(name)

    if missing_docs:
        print(f"\n⚠️  缺少目标文档，无法继续")
        return

    # 开始整合
    print("\n开始整合文档...")
    success_count = 0
    for doc_type in TARGET_DOCS.keys():
        if integrate_document(doc_type):
            success_count += 1

    # 生成报告
    generate_summary_report()

    print(f"\n✅ 成功整合 {success_count}/{len(TARGET_DOCS)} 个文档")

if __name__ == '__main__':
    main()
