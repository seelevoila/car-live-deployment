from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SUBMISSION_DIR = Path(__file__).resolve().parent
DOCS_DIR = SUBMISSION_DIR / "2026-chongqing-ai-competition"
DIAGRAMS_DIR = SUBMISSION_DIR / "diagrams"
BACKUP_DIR = SUBMISSION_DIR / "_docx_backup_before_finalize"

DOC_SPECS = [
    {
        "name": "01-系统设计文档-完整版.docx",
        "figures": [("2总体架构", "01-系统架构图", "系统架构图"), ("3.6部署架构与扩展性", "06-数据流向", "数据流向")],
    },
    {
        "name": "02-汽车垂直领域RAG知识库构建方案-完整版.docx",
        "figures": [("3检索策略", "02-RAG检索Pipeline", "RAG检索Pipeline")],
    },
    {
        "name": "03-流式TTS动态改稿实现方案-完整版.docx",
        "figures": [("2.6延迟优化深度分析", "03-TTS时序图", "TTS流式合成时序"), ("2.7动态改稿算法详解", "04-改稿状态机", "动态改稿状态机")],
    },
    {
        "name": "04-语音克隆拟人化优化方案-完整版.docx",
        "figures": [("2.6质量评估完整Pipeline", "05-质量门禁", "质量门禁流程")],
    },
    {"name": "05-标准化功能测试报告-完整版.docx", "figures": []},
]

DROP_EXACT = {
    "这是第一部分扩充内容。我将继续完成其他章节。",
    "这是RAG方案文档的扩充内容。继续完成其他部分...",
    "这是TTS方案文档的完整扩充内容。",
    "这是语音克隆方案的完整扩充内容。",
    "这是测试报告的完整扩充内容。",
    "---",
}


def stripped(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def set_rfonts(target, font_name: str) -> None:
    """Set all OOXML font slots, including the East Asian slot Word uses for Chinese."""
    rpr = target._element.get_or_add_rPr() if hasattr(target, "_element") else target.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{slot}"), font_name)


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = name
    set_rfonts(run, name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, name: str, size: float | None = None, bold: bool | None = None) -> None:
    style.font.name = name
    set_rfonts(style, name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def set_spacing(fmt, line=1.5, before=0, after=0) -> None:
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if line == 1.5 else WD_LINE_SPACING.SINGLE
    fmt.line_spacing = line
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def clear_borders(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pborder = ppr.find(qn("w:pBdr"))
    if pborder is not None:
        ppr.remove(pborder)


def is_heading(paragraph) -> bool:
    return bool(paragraph.style and paragraph.style.name.startswith("Heading"))


def heading_level(text: str) -> int | None:
    text = text.strip()
    if re.match(r"^\d+(?:\.\d+)*\s", text):
        dots = text.split()[0].count(".")
        return min(dots + 1, 3)
    if re.match(r"^\d+\.\s", text):
        return 2
    if text == "参考资料" or text == "扩充内容":
        return 1
    if "扩充内容" in text:
        return 2
    return None


def remove_appendix_and_noise(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    appendix_start = None
    for p in paragraphs:
        if "架构图插入指南" in p.text:
            appendix_start = p
            break
    if appendix_start is not None:
        # Remove the appendix and its preceding page-break paragraph.
        previous = appendix_start._element.getprevious()
        while previous is not None:
            prev_text = "".join(previous.itertext()).strip()
            if not prev_text:
                previous = previous.getprevious()
                continue
            if previous.tag == qn("w:p"):
                prev_p = next((p for p in doc.paragraphs if p._element is previous), None)
                if prev_p is not None and previous.find(".//" + qn("w:br")) is not None:
                    remove_paragraph(prev_p)
            break
        for p in list(doc.paragraphs):
            if p._element is appendix_start._element or p._element.getprevious() is not None:
                pass
        started = False
        for p in list(doc.paragraphs):
            if p._element is appendix_start._element:
                started = True
            if started:
                remove_paragraph(p)

    for p in list(doc.paragraphs):
        text = p.text.strip()
        if text in DROP_EXACT:
            remove_paragraph(p)


def normalize_headings(doc: Document) -> None:
    for p in doc.paragraphs:
        if not is_heading(p):
            continue
        level = heading_level(p.text)
        if level is None:
            p.style = doc.styles["Normal"]
            for run in p.runs:
                set_run_font(run, "SimSun", 11)
            continue
        p.style = doc.styles[f"Heading {level}"]


def find_insertion_target(doc: Document, key: str):
    key = stripped(key)
    paragraphs = list(doc.paragraphs)
    for index, p in enumerate(paragraphs):
        if stripped(p.text) != key:
            continue
        target = p
        for candidate in paragraphs[index + 1 :]:
            text = candidate.text.strip()
            if not text:
                continue
            if text.startswith("┌") or text.startswith("+") or text.startswith("|"):
                break
            if is_heading(candidate):
                break
            if 20 <= len(text) <= 800:
                target = candidate
            break
        return target
    return None


def insert_after(paragraph, new_paragraph):
    paragraph._p.addnext(new_paragraph._p)


def add_figure_after(doc: Document, target, image_path: Path, caption: str, number: int) -> None:
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(4)
    image_paragraph.paragraph_format.space_after = Pt(2)
    image_paragraph.paragraph_format.keep_with_next = True

    run = image_paragraph.add_run()
    picture = run.add_picture(str(image_path), width=Cm(14))
    max_height = Cm(16)
    if picture.height > max_height:
        ratio = max_height / picture.height
        picture.height = int(picture.height * ratio)
        picture.width = int(picture.width * ratio)

    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(2)
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_run = caption_paragraph.add_run(f"图{number} {caption}")
    set_run_font(caption_run, "SimSun", 10)

    insert_after(target, image_paragraph)
    insert_after(image_paragraph, caption_paragraph)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, "SimSun", 11)
    set_spacing(normal.paragraph_format, 1.5, 0, 0)

    title = styles["Title"]
    set_style_font(title, "SimHei", 22, True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(title.paragraph_format, 1.0, 0, 10)

    for level, size in ((1, 16), (2, 13.5), (3, 12)):
        style = styles[f"Heading {level}"]
        set_style_font(style, "SimHei", size, True)
        style.paragraph_format.keep_with_next = True
        if level == 1:
            set_spacing(style.paragraph_format, 1.0, 12, 6)
        elif level == 2:
            set_spacing(style.paragraph_format, 1.0, 9, 4)
        else:
            set_spacing(style.paragraph_format, 1.0, 6, 3)

    no_spacing = styles["No Spacing"]
    set_style_font(no_spacing, "Microsoft YaHei", 9)
    set_spacing(no_spacing.paragraph_format, 1.0, 0, 0)

    for name in ("List Bullet", "List Number"):
        if name in styles:
            set_style_font(styles[name], "SimSun", 11)
            set_spacing(styles[name].paragraph_format, 1.5, 0, 0)

    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)


def normalize_body_and_tables(doc: Document) -> None:
    configure_styles(doc)
    paragraphs = doc.paragraphs
    for index, p in enumerate(paragraphs):
        if index < 3:
            clear_borders(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if index == 0:
                size, font, bold = 14, "SimHei", False
            elif index == 1:
                size, font, bold = 22, "SimHei", True
            else:
                size, font, bold = 12, "SimSun", False
            p.paragraph_format.space_before = Pt(0 if index else 8)
            p.paragraph_format.space_after = Pt(8 if index == 1 else 4)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                set_run_font(run, font, size, bold)
            continue

        style_name = p.style.name if p.style else "Normal"
        if style_name == "No Spacing":
            font, size, line = "Microsoft YaHei", 9, 1.0
        elif style_name.startswith("Heading"):
            font, size, line = "SimHei", {"Heading 1": 16, "Heading 2": 13.5, "Heading 3": 12}.get(style_name, 12), 1.0
        else:
            font, size, line = "SimSun", 11, 1.5
        p.paragraph_format.line_spacing = line
        for run in p.runs:
            set_run_font(run, font, size)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        set_run_font(run, "Microsoft YaHei", 10)

    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                for run in p.runs:
                    set_run_font(run, "SimSun", 9)


def process(spec) -> None:
    source = DOCS_DIR / spec["name"]
    if not source.exists():
        raise FileNotFoundError(source)
    BACKUP_DIR.mkdir(exist_ok=True)
    backup = BACKUP_DIR / source.name
    if not backup.exists():
        shutil.copy2(source, backup)

    doc = Document(source)
    remove_appendix_and_noise(doc)
    normalize_headings(doc)

    for number, (heading_key, image_stem, caption) in enumerate(spec["figures"], start=1):
        image_path = DIAGRAMS_DIR / f"{image_stem}.png"
        target = find_insertion_target(doc, heading_key)
        if target is None:
            raise RuntimeError(f"Heading not found in {source.name}: {heading_key}")
        add_figure_after(doc, target, image_path, caption, number)

    normalize_body_and_tables(doc)
    doc.save(source)
    print(f"OK {source.name}")


def main() -> None:
    for spec in DOC_SPECS:
        process(spec)


if __name__ == "__main__":
    main()
