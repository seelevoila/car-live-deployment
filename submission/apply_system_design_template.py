from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
DOC_PATH = ROOT / "2026-chongqing-ai-competition" / "01-系统设计文档-完整版.docx"
BACKUP = ROOT / "_docx_backup_before_template" / DOC_PATH.name


def text_of(element) -> str:
    return "".join(element.itertext()).replace("\u00a0", " ").strip()


def paragraph_for(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"paragraph not found: {prefix}")


def is_h1_element(element, doc: Document) -> bool:
    if element.tag != qn("w:p"):
        return False
    paragraph = Paragraph(element, doc)
    return bool(paragraph.style and paragraph.style.name == "Heading 1")


def section_ranges(doc: Document):
    body = doc.element.body
    children = list(body.iterchildren())
    starts = []
    for i, element in enumerate(children):
        if not is_h1_element(element, doc):
            continue
        t = text_of(element)
        if t.startswith(("1 ", "2 ", "3 ", "4 ", "5 ", "6 ", "7 ", "参考资料", "扩充内容")):
            starts.append((i, t, element))
    ranges = {}
    for pos, (start, title, element) in enumerate(starts):
        end = starts[pos + 1][0] if pos + 1 < len(starts) else len(children) - 1
        ranges[title] = children[start:end]
    return children, starts, ranges


def replace_heading(doc: Document, old_prefix: str, new_text: str) -> Paragraph:
    paragraph = paragraph_for(doc, old_prefix)
    paragraph.text = new_text
    return paragraph


def make_paragraph(doc: Document, text: str, style: str = "Normal", align=None) -> Paragraph:
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    if align is not None:
        p.alignment = align
    return p


def insert_before(target: Paragraph, new_paragraph: Paragraph) -> None:
    target._p.addprevious(new_paragraph._p)


def insert_after(target: Paragraph, new_paragraph: Paragraph) -> None:
    target._p.addnext(new_paragraph._p)


def move_to_body_end(doc: Document, paragraph: Paragraph) -> None:
    doc.element.body.insert(-1, paragraph._p)


def set_keep(paragraph: Paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_heading_before(doc: Document, target: Paragraph, text: str, level: int) -> Paragraph:
    p = make_paragraph(doc, text, f"Heading {level}")
    insert_before(target, p)
    set_keep(p)
    return p


def add_heading_after(doc: Document, target: Paragraph, text: str, level: int) -> Paragraph:
    p = make_paragraph(doc, text, f"Heading {level}")
    insert_after(target, p)
    set_keep(p)
    return p


def reorder_main_sections(doc: Document) -> None:
    children, starts, ranges = section_ranges(doc)
    if len(starts) < 9:
        raise RuntimeError(f"unexpected section count: {len(starts)}")

    # Keep the cover and metadata in place, then place the standard sections in order.
    first_h1 = starts[0][0]
    cover = children[:first_h1]
    by_prefix = {}
    for title, block in ranges.items():
        for prefix in ("1 ", "2 ", "3 ", "4 ", "5 ", "6 ", "7 ", "参考资料", "扩充内容"):
            if title.startswith(prefix):
                by_prefix[prefix] = block
                break

    ordered = (
        cover
        + by_prefix["1 "]
        + by_prefix["2 "]
        + by_prefix["3 "]
        + by_prefix["6 "]
        + by_prefix["4 "]
        + by_prefix["5 "]
        + by_prefix["7 "]
        + by_prefix["参考资料"]
        + by_prefix["扩充内容"]
    )

    body = doc.element.body
    sectpr = body.sectPr
    for element in list(body.iterchildren()):
        if element is not sectpr:
            body.remove(element)
    for element in ordered:
        body.insert(len(body) - 1, element)


def add_standard_subsections(doc: Document) -> None:
    # Capture the original section objects before changing any heading text.
    h1 = paragraph_for(doc, "1  ")
    h1_1 = paragraph_for(doc, "1.1")
    h1_2 = paragraph_for(doc, "1.2")
    h2 = paragraph_for(doc, "2.1")
    h_main_2 = paragraph_for(doc, "2  ")
    h1_3 = paragraph_for(doc, "3  ")
    h1_4 = paragraph_for(doc, "4 ")
    h1_5 = paragraph_for(doc, "5 ")
    h1_6 = paragraph_for(doc, "6 ")
    h1_7 = paragraph_for(doc, "7 ")
    h_ref = paragraph_for(doc, "参考资料")
    h_expanded = paragraph_for(doc, "扩充内容")

    h1.text = "1 引言"
    h1_1.text = "1.1 编写目的"
    h1_2.text = "1.2 项目背景与业务边界"
    p = make_paragraph(doc, "1.3 参考材料", "Heading 2")
    insert_before(h_main_2, p)
    set_keep(p)
    ref = make_paragraph(doc, "项目需求与提交规范、GPT-SoVITS 官方文档及本项目已有配置与测试记录。")
    insert_before(h_main_2, ref)

    h_main_2.text = "2 系统总体设计"
    h2.text = "2.3 整体技术架构与关键数据流"
    add_heading_after(doc, h_main_2, "2.1 整体架构", 2)
    caption = paragraph_for(doc, "图1 系统架构图")
    add_heading_after(doc, caption, "2.2 整体功能架构", 2)

    additions = [
        ("2.4 设计目标", 2, "系统以可追溯问答、连续播报、边界安全改稿和本地可运行为主要设计目标。"),
        ("2.5 设计原则", 2, "设计在满足现场可用性的基础上，优先保证数据可追溯、接口可替换和故障可恢复。"),
        ("2.5.1 总体原则", 3, "模块职责清晰，数据流向可观测，关键状态和版本均可复现。"),
        ("2.5.2 实用性和先进性", 3, "采用成熟的 FastAPI、SQLite、FAISS、BM25 和 GPT-SoVITS 组件，兼顾本地部署成本与后续升级空间。"),
        ("2.5.3 标准化、开放性、兼容性", 3, "通过 HTTP/JSON、PCM/WAV 和通用文件格式连接前后端及模型服务，降低替换组件的成本。"),
        ("2.5.4 高可靠性、稳定性", 3, "通过健康检查、超时、重试、缓存和分级回退，保证单个服务异常时仍能保留文字答案。"),
        ("2.5.5 易用性、灵活性和经济性", 3, "以少量配置完成资料导入、检索、播报和改稿，优先复用本地资源并为分布式扩展预留接口。"),
    ]
    for title, level, body in additions:
        heading = make_paragraph(doc, title, f"Heading {level}")
        insert_before(h1_3, heading)
        set_keep(heading)
        body_paragraph = make_paragraph(doc, body)
        insert_after(heading, body_paragraph)

    h1_3.text = "3 系统功能模块详细设计"
    h1_6.text = "4 性能设计"
    cursor = h1_6
    for title in ("4.1 响应时间", "4.2 并发用户数"):
        heading = make_paragraph(doc, title, "Heading 2")
        insert_after(cursor, heading)
        set_keep(heading)
        cursor = heading

    db_heading = make_paragraph(doc, "5 数据库设计", "Heading 1")
    insert_before(h1_4, db_heading)
    db_text = make_paragraph(doc, "系统采用 SQLite 保存文档、片段、版本、会话和音色 Profile 等可追踪记录；FAISS 保存语义向量索引，BM25 保存中文词法索引。数据库记录与索引通过文档版本和片段标识关联，支持来源追溯、重建索引和后续迁移 PostgreSQL。")
    insert_before(h1_4, db_text)

    h1_4.text = "6 接口设计"
    add_heading_after(doc, h1_4, "6.1 接口清单", 2)
    h1_5.text = "7 系统出错处理设计"
    add_heading_after(doc, h1_5, "7.1 出错信息", 2)
    h1_7.text = "8 系统处理规定"
    cursor = h1_7
    for title, body in [
        ("8.1 输入输出要求", "输入支持 PDF、DOCX、TXT 和脚本文本；输出包括可溯源答案、PCM 音频及状态信息。"),
        ("8.2 数据管理能力要求", "资料、片段、版本、索引和音色 Profile 均应具备唯一标识、来源和更新时间。"),
        ("8.3 故障处理要求", "服务异常时保留已播放内容，停止未提交任务，并按健康检查和回退策略恢复。"),
        ("8.4 其他专门要求", "涉及个人信息和参考音频时，必须完成脱敏、授权确认和本地存储控制。"),
    ]:
        heading = make_paragraph(doc, title, "Heading 2")
        insert_after(cursor, heading)
        set_keep(heading)
        body_paragraph = make_paragraph(doc, body)
        insert_after(heading, body_paragraph)
        cursor = body_paragraph

    h_ref.text = "附录 A 参考资料"
    h_expanded.text = "附录 B 扩充内容"


def main() -> None:
    BACKUP.parent.mkdir(exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(DOC_PATH, BACKUP)
    doc = Document(DOC_PATH)
    reorder_main_sections(doc)
    add_standard_subsections(doc)
    doc.save(DOC_PATH)
    print(f"saved {DOC_PATH}")


if __name__ == "__main__":
    main()
