from pathlib import Path
import json
import struct
import threading
import time
from tempfile import TemporaryDirectory
import wave

from docx import Document

from app import main
from app.main import INSUFFICIENT_ANSWER, _audio_quality, _iter_pcm_wav_payload, _live_unit_params, _llm_answer, _normalize_prompt_text, _normalize_reference_audio, _prompt_alignment, _short_answer, _tts_params, _tts_stream_units, _tts_text_units, _unlink_owned_upload, _voice_model_profile, _wav_data_offset, extract, normalize_tts_text
from app.db import conn, init_db
from app.config import settings


def test_unknown_question_does_not_raise_and_uses_source_fallback():
    sources = [{"content": "车型定位是紧凑型SUV。", "score": 0.4}]
    result = _short_answer("适合家用吗？", sources)
    assert "知识库中暂时没有足够依据" in result


def test_structured_vehicle_answers_disambiguate_similar_fields():
    sources = [{"content": "最大功率(kW)是：150\n电动机(Ps)是：204\n对外交流放电功率(kW)是：6"}]
    assert "204" in _short_answer("电动机功率是多少？", sources)
    assert "150" in _short_answer("最大功率是多少？", sources)
    assert "6" in _short_answer("对外放电功率是多少？", sources)


def test_structured_vehicle_answers_support_configuration_fields():
    sources = [{"content": "车道偏离预警系统是：标配\nOTA升级是：标配\nWi-Fi热点是：标配"}]
    assert "标配" in _short_answer("OTA升级是什么配置？", sources)
    assert "标配" in _short_answer("车道偏离预警是标配吗？", sources)
    assert _short_answer("这款车适合家用吗？", sources) == INSUFFICIENT_ANSWER


def test_tts_normalization_handles_vehicle_units_and_dates():
    result = normalize_tts_text("2026款，续航580km，快充0.33小时")
    assert "二零二六款" in result
    assert "五百八十公里" in result
    assert "零点三三小时" in result
    assert " " not in result


def test_tts_normalization_removes_formatting_spaces_from_model_names():
    result = normalize_tts_text("欧拉 5 EV 2026 款 580 km")
    assert result == "欧拉五EV二零二六款五百八十公里"


def test_prompt_normalization_removes_subtitle_whitespace_without_merging_english_words():
    result = _normalize_prompt_text("这是\n GPT SoVITS \n模型")
    assert result == "这是GPT SoVITS模型。"


def test_prompt_normalization_preserves_ellipsis_ending():
    assert _normalize_prompt_text("没钱…") == "没钱…"


def test_audio_quality_reports_pcm_reference():
    root = Path(__file__).resolve().parents[2] / "data" / "uploads" / "voices"
    path = root / "07c2ddaa9e9c48b38effbbe6713386f8.optimized.wav"
    if not path.is_file():
        path = root / "07c2ddaa9e9c48b38effbbe6713386f8.normalized.clean.wav"
    if not path.is_file():
        return
    quality = _audio_quality(str(path))
    expected = "ready" if "optimized" in path.stem else "needs-review"
    assert quality["status"] == expected
    assert quality["channels"] == 1
    assert quality["sample_rate"] == 24000


def test_prompt_alignment_flags_transcript_that_is_too_long_for_reference():
    result = _prompt_alignment("这是一个非常长的参考文本，明显超过了几秒录音能自然说完的长度。", 3.0)
    assert result["status"] == "review"
    assert "偏长" in result["message"]


def test_prompt_alignment_accepts_normal_reference_density():
    result = _prompt_alignment("这是命运的邂逅吗？还是，久别重逢呢。", 3.6)
    assert result["status"] == "ok"


def test_audio_quality_keeps_natural_internal_pause_as_warning():
    with TemporaryDirectory() as folder:
        path = Path(folder) / "pause.wav"
        sample_rate = 24000
        frames = []
        # A speech-like tone with a 700 ms phrase break. The break should be
        # reported to the operator but must not be time-compressed.
        for index in range(int(sample_rate * 1.6)):
            frames.append(int(5000 * (index % 80) / 80))
        frames.extend([0] * int(sample_rate * 0.7))
        for index in range(int(sample_rate * 1.6)):
            frames.append(int(5000 * (index % 80) / 80))
        with wave.open(str(path), "wb") as audio:
            audio.setnchannels(1)
            audio.setsampwidth(2)
            audio.setframerate(sample_rate)
            audio.writeframes(struct.pack("<" + "h" * len(frames), *frames))
        quality = _audio_quality(str(path))
    assert quality["status"] == "ready"
    assert quality["longest_silence"] >= 0.65
    assert quality["warnings"]


def test_reference_normalization_does_not_use_internal_silence_compression(monkeypatch, tmp_path):
    source = tmp_path / "quiet.wav"
    sample_rate = 24000
    speech = [int(7000 * ((index % 100) / 100 - 0.5)) for index in range(int(sample_rate * 1.6))]
    samples = speech + [0] * int(sample_rate * 0.7) + speech
    with wave.open(str(source), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(sample_rate)
        audio.writeframes(struct.pack("<" + "h" * len(samples), *samples))
    ffmpeg = main._ffmpeg_binary()
    if not ffmpeg:
        return
    normalized = _normalize_reference_audio(source, force=True)
    try:
        assert normalized.is_file()
        # The 700 ms phrase break must survive normalization; the old
        # silenceremove graph collapsed it and shortened the prompt.
        assert _audio_quality(str(normalized))["duration"] >= 3.7
        assert normalized.name.endswith(".optimized.wav")
    finally:
        normalized.unlink(missing_ok=True)


def test_docx_extract_includes_table_cells():
    with TemporaryDirectory() as folder:
        path = Path(folder) / "vehicle.docx"
        document = Document()
        document.add_paragraph("车型参数")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "续航"
        table.cell(0, 1).text = "580公里"
        table.cell(1, 0).text = "轴距"
        table.cell(1, 1).text = "2720毫米"
        document.save(path)
        content = extract(path)[0][1]
    assert "580公里" in content
    assert "2720毫米" in content


def test_delete_guard_does_not_remove_sample_files():
    sample = Path(__file__).resolve().parents[2] / "data" / "sample" / "protected-test.txt"
    sample.write_text("fixture", encoding="utf-8")
    try:
        assert _unlink_owned_upload(str(sample)) is False
        assert sample.is_file()
    finally:
        sample.unlink(missing_ok=True)


def test_stream_text_units_bound_first_audio_size():
    units = _tts_text_units("今天为大家介绍这款车型的续航和智能配置，以及安全辅助功能。")
    assert units
    assert max(map(len, units)) <= 48
    assert min(map(len, units)) >= 8


def test_tts_scheduler_pauses_background_work_for_foreground_ticket():
    # A live stream keeps one ticket across all phrase units. Optional
    # calibration must yield immediately instead of switching profiles between
    # those units.
    with main._foreground_ticket():
        assert main._acquire_tts_lock("background", timeout=0.02) is False
    assert main.TTS_FOREGROUND_ACTIVE == 0
    assert main.TTS_FOREGROUND_WAITERS == 0


def test_tts_scheduler_foreground_waiter_wins_after_background_request():
    assert main._acquire_tts_lock("background", timeout=0.1)
    foreground_acquired = threading.Event()
    release_foreground = threading.Event()
    background_result = []

    def try_foreground():
        acquired = main._acquire_tts_lock("foreground", timeout=0.5)
        if acquired:
            foreground_acquired.set()
            release_foreground.wait(0.5)
            main._release_tts_lock()

    def try_background():
        background_result.append(main._acquire_tts_lock("background", timeout=0.2))

    foreground = threading.Thread(target=try_foreground)
    background = threading.Thread(target=try_background)
    foreground.start()
    time.sleep(0.02)
    background.start()
    try:
        main._release_tts_lock("background")
        assert foreground_acquired.wait(0.5)
        foreground.join(0.5)
        release_foreground.set()
        foreground.join(0.5)
        background.join(0.5)
    finally:
        release_foreground.set()
        if foreground.is_alive():
            foreground.join(0.5)
        if background.is_alive():
            background.join(0.5)
        if main.TTS_LOCK.locked():
            main._release_tts_lock()
    assert background_result == [False]
    assert main.TTS_FOREGROUND_ACTIVE == 0


def test_stream_text_units_do_not_split_number_and_unit():
    units = _tts_text_units("欧拉5 EV 2026 款 580km 激光雷达版，续航580公里，轴距2720毫米。")
    joined = "".join(units)
    assert "五百八十公里" in joined
    assert "二千七百二十毫米" in joined
    assert all(not unit.strip().endswith(("五百八十", "二千七百二十")) for unit in units)


def test_stream_text_units_keep_short_live_salutation_with_following_clause():
    units = _tts_text_units("老板，今天为大家介绍欧拉5 EV 2026 款 580km 激光雷达版。")
    assert units
    # Keep the salutation out of a standalone 2-character request, while
    # bounding the first unit so it can produce PCM within three seconds.
    assert units[0].startswith("老板，")
    assert len(units[0]) > 2
    assert len(units[0]) <= 48
    assert "激光雷达" in units[0]


def test_stream_text_units_merge_two_full_clauses_to_prevent_live_gap():
    units = _tts_text_units("它的 CLTC 纯电续航为五百八十公里，轴距二千七百二十毫米，最大功率一百五十千瓦。")
    assert units
    assert "轴距" in "".join(units)
    assert max(map(len, units)) <= 48
    assert all(len(unit.strip('，。')) >= 8 for unit in units)
    assert any("续航为五百八十公里，轴距" in unit for unit in units)


def test_unitized_stream_request_preserves_frontend_phrase():
    request = main.TTSRequest(text="老板，今天为大家介绍欧拉五 EV 二零二六款。", unitized=True)
    assert _tts_stream_units(request) == [normalize_tts_text(request.text)]


def test_non_unitized_stream_request_keeps_server_fallback_splitter():
    request = main.TTSRequest(text="今天为大家介绍这款车型的续航和智能配置，以及安全辅助功能。")
    assert _tts_stream_units(request) == _tts_text_units(request.text)


def test_reference_quality_flags_long_internal_silence():
    path = Path(__file__).resolve().parents[2] / "data" / "uploads" / "voices" / "07c2ddaa9e9c48b38effbbe6713386f8.normalized.clean.wav"
    if not path.is_file():
        return
    quality = _audio_quality(str(path))
    assert quality["longest_silence"] >= 0.55
    assert quality["status"] == "needs-review"


def test_wav_data_offset_handles_extra_chunks():
    fmt = b"fmt " + (16).to_bytes(4, "little") + (1).to_bytes(2, "little") + (1).to_bytes(2, "little")
    fmt += (24000).to_bytes(4, "little") + (48000).to_bytes(4, "little") + (2).to_bytes(2, "little") + (16).to_bytes(2, "little")
    junk = b"JUNK" + (4).to_bytes(4, "little") + b"test"
    data = b"data" + (4).to_bytes(4, "little") + b"\x00\x00\x00\x00"
    wav = b"RIFF" + (4 + len(fmt) + len(junk) + len(data)).to_bytes(4, "little") + b"WAVE" + fmt + junk + data
    assert _wav_data_offset(wav) == 12 + len(fmt) + len(junk) + 8


def test_stream_payload_removes_headers_between_units():
    header = (
        b"RIFF\x24\x00\x00\x00WAVE"
        b"fmt \x10\x00\x00\x00\x01\x00\x01\x00"
        b"\x80\xbb\x00\x00\x00\x77\x01\x00\x02\x00\x10\x00"
        b"data\x00\x00\x00\x00"
    )
    first = header + b"\x01\x00\x02\x00"
    second = header + b"\x03\x00\x04\x00"
    output = b"".join(_iter_pcm_wav_payload([first[:17], first[17:]], include_header=True))
    output += b"".join(_iter_pcm_wav_payload([second[:9], second[9:]], include_header=False))
    assert output == header + b"\x01\x00\x02\x00\x03\x00\x04\x00"


def test_live_tts_uses_incremental_fragment_profile(monkeypatch):
    monkeypatch.setattr(main, "_voice_config", lambda _voice_id: ("ref.wav", "参考文本。", "zh"))
    params = _tts_params(main.TTSRequest(text="测试"), streaming=True)
    assert params["parallel_infer"] is True
    assert params["streaming_mode"] == 1
    assert params["fragment_interval"] == 0.0


def test_tts_params_preserve_multi_reference_payload_for_post_json(monkeypatch):
    monkeypatch.setattr(main, "_voice_config", lambda _voice_id: ("primary.wav", "参考文本。", "zh"))
    monkeypatch.setattr(main, "_voice_aux_reference_paths", lambda _voice_id: ["aux-a.wav", "aux-b.wav"])
    params = _tts_params(main.TTSRequest(text="测试", voice_id="voice-multi"))
    assert params["ref_audio_path"] == "primary.wav"
    assert params["aux_ref_audio_paths"] == ["aux-a.wav", "aux-b.wav"]


def test_preset_voices_keep_fixed_sampler_and_user_speed(monkeypatch):
    for voice_id in main.PRESET_VOICES:
        base = _tts_params(main.TTSRequest(text="测试", voice_id=voice_id))
        styled = _tts_params(main.TTSRequest(text="测试", voice_id=voice_id,
            delivery='lively', temperature=1.2, top_p=1, speed_factor=1.1))
        assert base['temperature'] == styled['temperature']
        assert base['top_p'] == styled['top_p']
        assert base['speed_factor'] == 1.0
        assert styled['speed_factor'] == 1.1
        assert base['seed'] == styled['seed']


def test_tts_defaults_use_reproducible_stable_sampling(monkeypatch):
    monkeypatch.setattr(main, "_voice_config", lambda _voice_id: ("ref.wav", "参考文本。", "zh"))
    first = _tts_params(main.TTSRequest(text="今天介绍这款车。", voice_id="voice-new"))
    retry = _tts_params(main.TTSRequest(text="今天介绍这款车。", voice_id="voice-new"))
    other = _tts_params(main.TTSRequest(text="今天介绍另一款车。", voice_id="voice-new"))
    calibrated = _tts_params(main.TTSRequest(text="今天介绍这款车。", voice_id="voice-new"), seed_override=12345)
    assert first["top_k"] == 15
    assert first["top_p"] == 0.65
    assert first["temperature"] == 0.60
    assert first["repetition_penalty"] == 1.35
    assert first["sample_steps"] == 32
    assert first["seed"] == retry["seed"]
    assert first["seed"] == other["seed"]
    assert calibrated["seed"] == 12345


def test_tts_uses_per_voice_calibrated_sampling_unless_manually_overridden(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "database_path", tmp_path / "voice-sampling.db")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    init_db()
    with conn() as c:
        c.execute(
            "INSERT INTO voices(id,name,style,provider,reference_path,prompt_text,prompt_lang,"
            "sampling_seed,sampling_top_k,sampling_top_p,sampling_temperature,sampling_model_version,calibration_details,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?,?, ?,datetime('now'))",
            ("voice-adaptive", "自适应", "自然", "gpt-sovits", "ref.wav", "参考文本。", "zh", 88, 10, 0.82, 0.62, settings.gpt_sovits_model_version, json.dumps({"calibration_version": main.CALIBRATION_VERSION})),
        )
    monkeypatch.setattr(main, "_voice_config", lambda _voice_id: ("ref.wav", "参考文本。", "zh"))
    adaptive = _tts_params(main.TTSRequest(text="测试", voice_id="voice-adaptive"))
    manual = _tts_params(
        main.TTSRequest(
            text="测试",
            voice_id="voice-adaptive",
            top_p=0.96,
            temperature=0.88,
            repetition_penalty=1.4,
        )
    )
    assert adaptive["top_p"] == 0.82
    assert adaptive["top_k"] == 10
    assert adaptive["temperature"] == 0.62
    assert adaptive["seed"] == 88
    assert manual["top_p"] == 0.96
    assert manual["temperature"] == 0.88
    assert manual["repetition_penalty"] == 1.4


def test_new_voice_routes_to_base_model(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "database_path", tmp_path / "voice-profile.db")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    init_db()
    reference = tmp_path / "clone.wav"
    with wave.open(str(reference), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(24000)
        audio.writeframes(struct.pack("<h", 1200) * (24000 * 3))
    with conn() as c:
        c.execute(
            "INSERT INTO voices(id,name,style,provider,reference_path,prompt_text,prompt_lang,model_profile,created_at) VALUES(?,?,?,?,?,?,?,?,datetime('now'))",
            ("voice-new", "星期日", "亲民", "gpt-sovits", str(reference), "参考文本。", "zh", "base"),
        )
    assert _voice_model_profile("voice-new") == "base"


def test_tts_params_include_same_speaker_auxiliary_references(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "database_path", tmp_path / "voice-aux.db")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    init_db()
    references = []
    for index in range(3):
        reference = tmp_path / f"ref-{index}.wav"
        with wave.open(str(reference), "wb") as audio:
            audio.setnchannels(1)
            audio.setsampwidth(2)
            audio.setframerate(24000)
            audio.writeframes(struct.pack("<h", 1200 + index * 100) * (24000 * 3))
        references.append(reference)
    with conn() as c:
        c.execute(
            "INSERT INTO voices(id,name,style,provider,reference_path,prompt_text,prompt_lang,"
            "aux_reference_paths,validated_aux_reference_paths,model_profile,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,datetime('now'))",
            ("voice-aux", "多参考", "自然", "gpt-sovits", str(references[0]), "参考文本。", "zh",
             json.dumps([str(references[1]), str(references[2])]),
             json.dumps([str(references[1]), str(references[2])]), "base"),
        )
    monkeypatch.setattr(main, "_voice_config", lambda _voice_id: (str(references[0]), "参考文本。", "zh"))
    params = _tts_params(main.TTSRequest(text="测试", voice_id="voice-aux"))
    assert params["aux_ref_audio_paths"] == [str(references[1]), str(references[2])]


def test_tts_params_honor_empty_screened_auxiliary_set(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "database_path", tmp_path / "voice-screened-aux.db")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    init_db()
    references = []
    for index in range(2):
        reference = tmp_path / f"screened-ref-{index}.wav"
        with wave.open(str(reference), "wb") as audio:
            audio.setnchannels(1)
            audio.setsampwidth(2)
            audio.setframerate(24000)
            audio.writeframes(struct.pack("<h", 1200 + index * 100) * (24000 * 3))
        references.append(reference)
    with conn() as c:
        c.execute(
            "INSERT INTO voices(id,name,style,provider,reference_path,prompt_text,prompt_lang,"
            "aux_reference_paths,validated_aux_reference_paths,model_profile,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,datetime('now'))",
            ("voice-screened", "筛选辅助参考", "自然", "gpt-sovits", str(references[0]),
             "参考文本。", "zh", json.dumps([str(references[1])]), "[]", "base"),
        )
    monkeypatch.setattr(main, "_voice_config", lambda _voice_id: (str(references[0]), "参考文本。", "zh"))
    params = _tts_params(main.TTSRequest(text="测试", voice_id="voice-screened"))
    assert params["aux_ref_audio_paths"] == []


def test_preview_tts_uses_non_streaming_quality_profile(monkeypatch):
    monkeypatch.setattr(main, "_voice_config", lambda _voice_id: ("ref.wav", "参考文本。", "zh"))
    params = _tts_params(main.TTSRequest(text="试听测试"), streaming=False)
    assert params["streaming_mode"] is False
    assert params["parallel_infer"] is True
    assert params["text_split_method"] == "cut2"
    assert params["fragment_interval"] == 0.0


def test_live_opening_unit_always_uses_incremental_profile(monkeypatch):
    monkeypatch.setattr(main, "_voice_config", lambda _voice_id: ("ref.wav", "参考文本。", "zh"))
    long_request = main.TTSRequest(text="老板，今天为大家介绍欧拉五 EV 二零二六款五百八十公里激光雷达版。")
    short_request = main.TTSRequest(text="欢迎来到直播间。")
    assert _live_unit_params(long_request, 0)["streaming_mode"] == 2
    assert _live_unit_params(long_request, 1)["streaming_mode"] == 2
    assert _live_unit_params(short_request, 0)["streaming_mode"] == 2
    assert _live_unit_params(short_request, 0)["parallel_infer"] is False


def test_streaming_mode_defaults_to_quality_profile():
    assert settings.gpt_sovits_streaming_mode == 1


def test_reference_optimization_resolves_original_upload(tmp_path):
    original = tmp_path / "speaker.wav"
    derivative = tmp_path / "speaker.normalized.clean.optimized.wav"
    original.write_bytes(b"original")
    derivative.write_bytes(b"derived")
    assert main._original_reference_audio(derivative) == original


def test_optional_llm_without_configuration_returns_local_fallback(monkeypatch):
    monkeypatch.setattr(settings, "llm_base_url", "")
    monkeypatch.setattr(settings, "llm_model", "")
    assert _llm_answer("续航是多少？", [{"content": "续航：580公里"}]) is None


def test_session_persists_selected_voice(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "database_path", tmp_path / "session.db")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    init_db()
    saved = main.create_session(main.Session(vehicle="欧拉5 EV", script="测试稿", voice_id="voice-demo"))
    assert saved["voice_id"] == "voice-demo"
    with conn() as c:
        assert c.execute("SELECT voice_id FROM sessions WHERE id=?", (saved["id"],)).fetchone()[0] == "voice-demo"


def test_updating_reference_text_invalidates_sampling_calibration(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "database_path", tmp_path / "voice-prompt-update.db")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    monkeypatch.setattr(settings, "gpt_sovits_url", "")
    init_db()
    reference = tmp_path / "clone.wav"
    with wave.open(str(reference), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(24000)
        audio.writeframes(struct.pack("<h", 1200) * (24000 * 3))
    with conn() as c:
        c.execute(
            "INSERT INTO voices(id,name,style,provider,reference_path,prompt_text,prompt_lang,"
            "sampling_seed,validated_aux_reference_paths,created_at) VALUES(?,?,?,?,?,?,?,?,?,datetime('now'))",
            ("voice-prompt-update", "更新测试", "自然", "gpt-sovits", str(reference),
             "旧参考文本。", "zh", 1234, "[]"),
        )
    saved = main.update_voice("voice-prompt-update", main.VoiceUpdate(prompt_text="新的参考文本。"))
    assert saved["prompt_text"] == "新的参考文本。"
    with conn() as c:
        row = c.execute(
            "SELECT sampling_seed,sampling_top_p,sampling_temperature,calibration_details "
            "FROM voices WHERE id=?",
            ("voice-prompt-update",),
        ).fetchone()
    assert tuple(row) == (None, None, None, None)


def test_default_tts_keeps_builtin_after_a_clone_is_added(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "database_path", tmp_path / "voice.db")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    monkeypatch.setattr(settings, "gpt_sovits_ref_audio", "")
    monkeypatch.setattr(settings, "gpt_sovits_prompt_text", "")
    init_db()
    reference = tmp_path / "clone.wav"
    with wave.open(str(reference), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(24000)
        audio.writeframes(struct.pack("<h", 1200) * (24000 * 3))
    with conn() as c:
        c.execute(
            "INSERT INTO voices(id,name,style,provider,reference_path,prompt_text,prompt_lang,created_at) VALUES(?,?,?,?,?,?,?,datetime('now'))",
            ("voice-ready", "测试克隆", "亲民", "gpt-sovits", str(reference), "测试参考文本。", "zh"),
        )
    ref, prompt, _ = main._voice_config("steady")
    assert ref == main.preset_reference_paths("steady")[0]
    assert prompt == main.PRESET_VOICES["steady"]["prompt_text"]
    clone_ref, clone_prompt, _ = main._voice_config("voice-ready")
    assert clone_ref == str(reference)
    assert clone_prompt == "测试参考文本。"


def test_default_tts_does_not_use_promptless_global_reference(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "database_path", tmp_path / "voice-global.db")
    monkeypatch.setattr(settings, "upload_dir", tmp_path / "uploads")
    monkeypatch.setattr(settings, "gpt_sovits_ref_audio", str(tmp_path / "global.wav"))
    monkeypatch.setattr(settings, "gpt_sovits_prompt_text", "")
    init_db()
    global_reference = tmp_path / "global.wav"
    clone_reference = tmp_path / "clone.wav"
    for reference in (global_reference, clone_reference):
        with wave.open(str(reference), "wb") as audio:
            audio.setnchannels(1)
            audio.setsampwidth(2)
            audio.setframerate(24000)
            audio.writeframes(struct.pack("<h", 1200) * (24000 * 3))
    with conn() as c:
        c.execute(
            "INSERT INTO voices(id,name,style,provider,reference_path,prompt_text,prompt_lang,created_at) VALUES(?,?,?,?,?,?,?,datetime('now'))",
            ("voice-global-fallback", "测试克隆", "亲民", "gpt-sovits", str(clone_reference), "真实参考文本。", "zh"),
        )
    ref, prompt, _ = main._voice_config("steady")
    assert ref == main.preset_reference_paths("steady")[0]
    assert prompt == main.PRESET_VOICES["steady"]["prompt_text"]
