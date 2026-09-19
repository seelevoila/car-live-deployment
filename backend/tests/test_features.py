import asyncio
from pathlib import Path
import json
import pytest
from fastapi.testclient import TestClient
from app import main, llm_gateway
from app.config import settings
from app.db import init_db, conn
from app.rag import retrieve


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'database_path',tmp_path/'test.sqlite3')
    monkeypatch.setattr(settings,'upload_dir',tmp_path/'uploads')
    monkeypatch.setattr(settings,'llm_base_url','')
    monkeypatch.setattr(settings,'llm_model','')
    init_db()
    return TestClient(main.app)


def upload(client,text,name='car.txt'):
    response=client.post('/api/documents',files={'file':(name,text.encode(),'text/plain')},data={'brand':'测试','series':'测试车','year':'2026'})
    assert response.status_code==200,response.text
    return response.json()['id']


def test_personal_data_is_removed_from_stored_file_and_index(client):
    did=upload(client,'轴距(mm)是：2700\n联系人：13800000000\n邮箱：demo@example.com\n身份证110101199001011234仅用于测试')
    with conn() as c:
        doc=c.execute('SELECT * FROM documents WHERE id=?',(did,)).fetchone()
        chunk=c.execute('SELECT * FROM chunks WHERE document_id=?',(did,)).fetchone()
    assert '13800000000' not in Path(doc['path']).read_text(encoding='utf-8')
    assert 'demo@example.com' not in chunk['content']
    assert '110101199001011234' not in chunk['content']
    assert '2700' in chunk['content']
    from app.rag_models import dimension
    assert len(chunk['embedding']) == dimension() * 4


def test_batch_metadata_versions_and_atomic_missing_id(client):
    ids=[upload(client,'轴距(mm)是：2700',f'car{i}.txt') for i in range(2)]
    response=client.patch('/api/knowledge/batch',json={'ids':ids,'license':'自有授权','source_url':'https://example.com/vehicle'})
    assert response.status_code==200
    for did in ids:
        data=client.get(f'/api/knowledge/{did}/content').json()
        assert data['document']['version']==2
        assert data['document']['license']=='自有授权'
    assert client.post('/api/knowledge/batch-delete',json={'ids':ids+[9999]}).status_code==404
    assert len(client.get('/api/documents').json())==2
    assert client.post('/api/knowledge/batch-delete',json={'ids':ids}).json()['deleted']==2
    with conn() as c:
        assert c.execute('SELECT count(*) FROM chunks').fetchone()[0]==0


def test_provenance_scores_are_bounded_and_versions_returned(client):
    upload(client,'CLTC纯电续航里程(km)是：580\n轴距(mm)是：2720')
    sources=retrieve('CLTC纯电续航是多少','测试','测试车','2026')
    assert sources and 0<sources[0]['score']<=1
    assert sources[0]['version']==1 and sources[0]['chunk_id']


def test_analytics_counts_real_events_and_deduplicates_retries(client):
    upload(client,'轴距(mm)是：2720')
    query=client.post('/api/query',json={'question':'轴距是多少','brand':'测试','series':'测试车'})
    assert '2720' in query.json()['answer']
    for _ in range(2):
        assert client.post('/api/analytics/events',json={'id':'event1','kind':'playback','seconds':5,'vehicle':'测试车'}).status_code==200
    data=client.get('/api/analytics').json()
    assert data['question_count']==1 and data['playback_seconds']==5
    assert data['retrieval_hotspots'][0]['count']==1


def test_no_fake_human_scores_and_real_scores_are_validated(client):
    assert client.get('/api/voice-evaluations').json()['records']==[]
    body={'voice_id':'friendly','reviewer':'听测者A','text':'大家好，欢迎来到汽车直播间。','similarity':4,'clarity':5,'pause':4,'emotion':3,'overall':4}
    assert client.post('/api/voice-evaluations',json={**body,'overall':6}).status_code==422
    assert client.post('/api/voice-evaluations',json=body).status_code==200
    data=client.get('/api/voice-evaluations').json()
    assert data['voices'][0]['averages']['overall']==4
    assert client.get('/api/reports/export').status_code==200


def test_unconfigured_model_status_is_honest_and_no_request_sent(client,monkeypatch):
    monkeypatch.setattr(llm_gateway.httpx.Client,'post',lambda *a,**k:pytest.fail('unexpected external request'))
    assert client.get('/api/llm/status').json()['configured'] is False
    assert llm_gateway.generate('问题',[{'content':'资料'}]) is None


def test_cancelled_async_lock_waiter_does_not_orphan_the_inference_lock():
    async def scenario():
        with main._tts_lock():
            entered=[]
            async def waiter():
                async with main._async_tts_lock(): entered.append(True)
            task=asyncio.create_task(waiter())
            await asyncio.sleep(.04)
            task.cancel()
            with pytest.raises(asyncio.CancelledError): await task
            assert not entered
        async with main._async_tts_lock(timeout=.2): pass
        assert not main.TTS_LOCK.locked()
    asyncio.run(scenario())


def test_cancelling_lock_owner_releases_for_next_request():
    async def scenario():
        entered=asyncio.Event()
        async def owner():
            async with main._async_tts_lock():
                entered.set()
                await asyncio.sleep(10)
        task=asyncio.create_task(owner())
        await entered.wait();task.cancel()
        with pytest.raises(asyncio.CancelledError): await task
        async with main._async_tts_lock(timeout=.2): pass
        assert not main.TTS_LOCK.locked()
    asyncio.run(scenario())


def test_expired_documents_are_excluded_from_current_retrieval(client):
    did=upload(client,'CLTC纯电续航里程(km)是：580')
    assert client.patch('/api/knowledge/batch',json={'ids':[did],'valid_until':'2020-01-01'}).status_code==200
    assert retrieve('续航','测试','测试车','2026')==[]


def test_comparison_uses_technical_evidence_not_marketing_filler():
    sources=[{'content':'核心亮点：三动力自由选，省钱省心，性价比高！\n纯电：150kW永磁电机；燃油：1.5T发动机和7速双离合。\n纯电：百公里电耗11kWh；燃油：WLTC油耗6.4L。'}]
    answer=main._short_answer('纯电和燃油在动力和能耗上有什么区别？',sources)
    assert '150kW' in answer and '6.4L' in answer
    assert '性价比高' not in answer


def test_conflicting_trade_in_policy_is_not_presented_as_a_single_price():
    sources=[{'content':'至高 10000 元置换补贴。\n最高7000元置换补贴，不限品牌。'}]
    assert '冲突' in main._short_answer('置换补贴最高是多少？',sources)


def test_legacy_binding_cannot_change_builtin_identity(client,monkeypatch):
    with conn() as c:
        c.execute("INSERT INTO voices(id,name,reference_path,prompt_text,model_profile) VALUES('clone-test','测试音色','test.wav','参考文字。','base')")
        c.execute("UPDATE voices SET reference_voice_id='clone-test' WHERE id='steady'")
    assert main._voice_config('steady')[0] == main.preset_reference_paths('steady')[0]
    assert main._voice_model_profile('steady')=='base'
    response=client.put('/api/voice-presets',json=dict.fromkeys(['steady','energetic','friendly'],'clone-test'))
    assert response.status_code==409
    assert client.get('/api/voice-presets').json()


def test_lively_delivery_preserves_facts_and_keeps_speed_bounded():
    from app.tts_gpt_sovits import delivery_text
    assert delivery_text('大家好，欢迎来到直播间。','lively',1)[0].endswith('！')
    text,speed=delivery_text('续航580公里，功率150千瓦。','warm',1)
    assert text=='续航580公里，功率150千瓦。'
    assert .9<speed<1


def test_multiformat_batch_redacts_stored_files_and_keeps_tables_and_pages(client):
    import io
    import fitz
    from docx import Document
    pdf=fitz.open()
    page=pdf.new_page()
    page.insert_text((40,60),'验收车续航(km)是：666',fontname='china-s')
    page.insert_text((40,100),'13800000000 fixture@example.com')
    pdf_bytes=pdf.tobytes()
    pdf.close()
    doc=Document()
    doc.add_paragraph('验收车资料 联系人：13800000000 fixture@example.com')
    table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='轴距(mm)是：'
    table.cell(0,1).text='2888'
    buffer=io.BytesIO()
    doc.save(buffer)
    files=[('files',('car.pdf',pdf_bytes,'application/pdf')),
           ('files',('car.docx',buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
           ('files',('car.txt','续航(km)是：666 联系人：13800000000 fixture@example.com'.encode(),'text/plain'))]
    result=client.post('/api/documents/batch',files=files,data={'brand':'多格式验收'})
    assert result.status_code==200 and result.json()['count']==3,result.text
    with conn() as c:
        docs=c.execute('SELECT * FROM documents').fetchall()
        chunks=c.execute('SELECT * FROM chunks').fetchall()
    for doc in docs:
        text='\n'.join(text for _,text in main.extract(Path(doc['path'])))
        assert '13800000000' not in text and 'fixture@example.com' not in text
    assert any('2888' in c['content'] for c in chunks)
    assert any(c['page']==1 and '666' in c['content'] for c in chunks)
    from app.rag_models import dimension
    assert all(len(c['embedding']) == dimension() * 4 for c in chunks)


def test_configured_model_adapter_handles_success_failures_and_unsupported_numbers(client,monkeypatch):
    import httpx
    monkeypatch.setattr(settings,'llm_base_url','https://model.example/v1')
    monkeypatch.setattr(settings,'llm_model','user-selected-model')
    calls=[]
    def response(_client,url,**kwargs):
        calls.append((url,kwargs['json']))
        return httpx.Response(200,request=httpx.Request('POST',url),json={'choices':[{'message':{'content':'根据资料，续航为580公里。[1]'}}]})
    monkeypatch.setattr(httpx.Client,'post',response)
    sources=[{'content':'续航580公里','document_name':'车型资料','version':2}]
    assert '580' in llm_gateway.generate('续航是多少',sources)
    assert calls[0][0]=='https://model.example/v1/chat/completions'
    assert calls[0][1]['model']=='user-selected-model'
    def unsupported(_client,url,**kwargs):
        return httpx.Response(200,request=httpx.Request('POST',url),json={'choices':[{'message':{'content':'续航999公里'}}]})
    monkeypatch.setattr(httpx.Client,'post',unsupported)
    assert llm_gateway.generate('续航是多少',sources) is None
    def timeout(*args,**kwargs): raise httpx.ReadTimeout('test timeout')
    monkeypatch.setattr(httpx.Client,'post',timeout)
    assert llm_gateway.generate('续航是多少',sources) is None
