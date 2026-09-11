"""Redact personal identifiers before storing newly uploaded documents."""
from pathlib import Path
import re

PATTERNS = [
    (re.compile(r'(?<!\d)1[3-9]\d{9}(?!\d)'), '[手机号已脱敏]'),
    (re.compile(r'(?<![0-9A-Za-z])\d{17}[\dXx](?![0-9A-Za-z])'), '[身份证已脱敏]'),
    (re.compile(r'[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}'), '[邮箱已脱敏]'),
]


def redact(text):
    for pattern,replacement in PATTERNS:
        text=pattern.sub(replacement,text)
    return text


def sanitize_upload(path: Path):
    count=0
    if path.suffix.lower()=='.txt':
        raw=path.read_text(encoding='utf-8-sig',errors='replace')
        count=sum(len(p.findall(raw)) for p,_ in PATTERNS)
        path.write_text(redact(raw),encoding='utf-8')
    elif path.suffix.lower()=='.docx':
        from docx import Document
        source=Document(path)
        clean=Document()
        for p in source.paragraphs:
            count+=sum(len(pattern.findall(p.text)) for pattern,_ in PATTERNS)
            if p.text: clean.add_paragraph(redact(p.text))
        for table in source.tables:
            output=clean.add_table(rows=len(table.rows),cols=len(table.columns))
            for i,row in enumerate(table.rows):
                for j,cell in enumerate(row.cells):
                    count+=sum(len(pattern.findall(cell.text)) for pattern,_ in PATTERNS)
                    output.cell(i,j).text=redact(cell.text)
        clean.save(path)  # normalized document omits comments, authors and embedded attachments
    elif path.suffix.lower()=='.pdf':
        import fitz
        with fitz.open(path) as doc:
            for page in doc:
                for pattern,_ in PATTERNS:
                    for match in set(pattern.findall(page.get_text())):
                        for rect in page.search_for(match):
                            page.add_redact_annot(rect,fill=(1,1,1))
                            count+=1
                page.apply_redactions(images=2)
            doc.set_metadata({})
            data=doc.tobytes(garbage=4,deflate=True)
        path.write_bytes(data)
    return count
