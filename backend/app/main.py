import asyncio
import io
import json
import logging
import math
import re
import statistics
import sqlite3
import struct
import subprocess
import threading
import time
import wave
import zlib
from functools import lru_cache
from copy import deepcopy
from contextlib import asynccontextmanager, contextmanager
from pathlib import Path
from typing import List
from uuid import uuid4

import httpx
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel, Field

from .config import settings
from .db import conn, init_db, now
from .rag import split_text, tokens, retrieve, retrieve_with_trace, encode_vectors, INDEX_VERSION, index_status, data_quality
from .rag_chunking import build_chunks
from .rag_answers import prepare_context, commercial_answer, applicable_warnings, live_data_answer, appliance_answer
from .compliance import redact, sanitize_upload
from . import llm_gateway
from .local_answers import normalize_question, excerpt_answer
from .tts_gpt_sovits import gpt_sovits_engine, delivery_text, model_profiles, profile_weights
from .preset_voices import PRESET_VOICES, reference_paths as preset_reference_paths
from .runtime_paths import detect_ffmpeg, speaker_model_root
from .audio_metrics import wav_metrics
from . import prompt_guard
from . import voice_adaptation


def _legacy_probe_failure(details):
    """Recognize old calibrations that accepted only silence/runaway audio."""
    scores = details.get("scores") if isinstance(details, dict) else None
    if not isinstance(scores, list) or not scores:
        return None
    valid = []
    for item in scores:
        prosody = item.get("prosody") if isinstance(item, dict) else None
        if not isinstance(prosody, dict) or not prosody.get("available"):
            continue
        duration = float(prosody.get("duration") or 0)
        voiced_ratio = float(prosody.get("voiced_ratio") or 0)
        valid.append(0.6 <= duration <= 12.0 and voiced_ratio >= 0.20)
    if valid and not any(valid):
        return "历史校准结果显示合成音频过短、过长或缺少有效人声，请重新录制参考音频"
    return None


def _migrate_voice_synthesis_validation():
    """Invalidate legacy clones whose calibration never checked real speech."""
    try:
        with conn() as c:
            rows = c.execute(
                "SELECT id,calibration_details,synthesis_status FROM voices "
                "WHERE reference_path<>'' AND COALESCE(synthesis_status,'ready')='ready'"
            ).fetchall()
            for row in rows:
                try:
                    details = json.loads(row[1] or "{}")
                except (TypeError, ValueError):
                    details = {}
                message = _legacy_probe_failure(details)
                if message:
                    c.execute(
                        "UPDATE voices SET synthesis_status='failed',synthesis_message=?,"
                        "synthesis_checked_at=? WHERE id=?",
                        (message, now(), row[0]),
                    )
    except sqlite3.Error:
        return


def _set_voice_synthesis_status(voice_id: str, status: str, message: str = "", duration: float | None = None):
    try:
        with conn() as c:
            c.execute(
                "UPDATE voices SET synthesis_status=?,synthesis_message=?,synthesis_duration=?,"
                "synthesis_checked_at=? WHERE id=?",
                (status, message[:500], duration, now(), voice_id),
            )
    except sqlite3.Error:
        return


def _voice_synthesis_gate(voice_id: str):
    """Prevent direct API callers from using a clone before probe acceptance."""
    if voice_adaptation.BUSY.is_set():
        raise HTTPException(409, '正在创建专属音色，播报暂时等待训练和验收完成')
    if voice_id in PRESET_VOICES:
        return
    try:
        with conn() as c:
            row = c.execute(
                "SELECT synthesis_status,synthesis_message FROM voices WHERE id=?",
                (voice_id,),
            ).fetchone()
    except sqlite3.Error:
        row = None
    if not row:
        return
    status = row[0] or "ready"
    if status == "failed":
        raise HTTPException(422, row[1] or "该克隆音色未通过实际合成验收，请重新录制")
    if status == "pending":
        raise HTTPException(409, row[1] or "该克隆音色正在进行实际合成验收，请稍候")


def _tts_provider():
    """Return the project's only supported speech engine."""
    return "gpt-sovits"


def _tts_global_reference_audio() -> str:
    return settings.gpt_sovits_ref_audio


@asynccontextmanager
async def lifespan(_app: FastAPI):
    init_db()
    _migrate_voice_synthesis_validation()
    voice_adaptation.recover()
    with conn() as c:
        count = c.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
    if count == 0:
        for path in sorted(settings.sample_dir.glob("*.txt")):
            ingest(path, path.name, "欧拉", "欧拉5 EV", "2026")
    with conn() as c:
        outdated = c.execute("SELECT DISTINCT d.id,d.path FROM documents d LEFT JOIN chunks ch ON d.id=ch.document_id WHERE ch.embedding IS NULL OR ch.embedding_model<>?",(INDEX_VERSION,)).fetchall()
        for item in outdated:
            if Path(item['path']).is_file():
                c.execute('DELETE FROM chunks WHERE document_id=?',(item['id'],))
                total=_write_chunks(c,item['id'],Path(item['path']))
                c.execute('UPDATE documents SET chunks=? WHERE id=?',(total,item['id']))
    TTS_WARMUP.clear()
    LIVE_PATH_WARMED.clear()
    async with httpx.AsyncClient(timeout=httpx.Timeout(120, connect=10, read=60), trust_env=False) as client:
        _app.state.tts_client = client
        warmup_task = asyncio.create_task(_warmup_tts())
        try:
            yield
        finally:
            if warmup_task:
                warmup_task.cancel()
                try:
                    await warmup_task
                except asyncio.CancelledError:
                    pass
            _app.state.tts_client = None


app = FastAPI(title="汽车直播智能体", version="0.3.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Frontend uses same-origin /api, allow any host
    allow_methods=["*"],
    allow_headers=["*"],
)


LIVE2D_ALLOWED_SUFFIXES = {".json", ".moc3", ".physics3.json", ".cdi3.json", ".exp3.json", ".png", ".jpg", ".jpeg"}


def _live2d_asset_path(asset_path: str) -> Path:
    """Resolve one avatar asset while keeping traversal outside the model root impossible."""
    root = Path(settings.live2d_model_root).expanduser().resolve()
    candidate = (root / asset_path).resolve()
    if candidate != root and root not in candidate.parents:
        raise HTTPException(404, "Live2D 资源不存在")
    if not candidate.is_file():
        raise HTTPException(404, "Live2D 资源不存在")
    lower_name = candidate.name.lower()
    if not any(lower_name.endswith(suffix) for suffix in LIVE2D_ALLOWED_SUFFIXES):
        raise HTTPException(404, "Live2D 资源类型不受支持")
    return candidate


class Query(BaseModel):
    question: str = Field(min_length=2)
    brand: str = ""
    series: str = ""
    year: str = ""
    top_k: int = Field(default=5, ge=1, le=20)


class Session(BaseModel):
    vehicle: str = Field(default="", max_length=200)
    script: str = Field(default="", max_length=20000)
    voice_id: str = "browser-default"


class Revision(BaseModel):
    script: str = Field(min_length=1, max_length=20000)
    sentence: int = Field(default=0, ge=0)


class State(BaseModel):
    status: str
    sentence: int = Field(default=0, ge=0)


class TTSRequest(BaseModel):
    # The live UI sends short units, while the API also supports direct long
    # requests and splits them at natural pauses server-side. `unitized` is
    # true only when the caller already applied the UI's punctuation-aware
    # queueing rules; in that case the server must preserve the unit intact.
    text: str = Field(min_length=1, max_length=20000)
    voice_id: str = "steady"
    unitized: bool = False
    # A bounded live batch is already split by the browser for progress and
    # revision purposes. Keep it as one upstream request so GPT-SoVITS can
    # preserve its prompt/semantic state across the contained phrases.
    stream_batch: bool = False
    delivery: str = Field(default='natural', pattern='^(natural|warm|lively|calm)$')
    speed_factor: float = Field(default=1.0, ge=0.6, le=1.6)
    # Conservative defaults are important for short reference cloning.  A
    # high-temperature sample can suddenly jump an octave even when the
    # speaker identity is otherwise correct.
    # None means use the per-voice calibrated value. An explicit value remains
    # available for callers that deliberately want manual sampling.
    top_k: int | None = Field(default=None, ge=1, le=30)
    # ``None`` means "use the profile calibrated for this reference voice".
    # Callers can still supply explicit values for deliberate manual tuning.
    top_p: float | None = Field(default=None, ge=0.1, le=1.0)
    temperature: float | None = Field(default=None, ge=0.1, le=2.0)
    repetition_penalty: float | None = Field(default=None, ge=0.5, le=3.0)


class VoiceUpdate(BaseModel):
    prompt_policy: str | None = Field(default=None, pattern='^(off|checked)$')
    name: str | None = None
    style: str | None = None
    prompt_text: str | None = None
    prompt_lang: str | None = None


class ScriptRequest(BaseModel):
    brand: str = "欧拉"
    series: str = "欧拉5 EV"
    year: str = "2026"
    selling_points: str = Field(min_length=2)
    delivery: str = Field(default='natural', max_length=40)
    duration_seconds: int = Field(default=60, ge=20, le=180)


ALLOWED_DOCS = {".pdf", ".docx", ".txt"}
ALLOWED_AUDIO = {".wav", ".mp3", ".flac", ".ogg", ".m4a"}
REFERENCE_MIN_SECONDS = 3.0
REFERENCE_MAX_SECONDS = 10.0
REFERENCE_RECOMMENDED_MIN_SECONDS = 5.0
MAX_AUX_REFERENCE_AUDIO = 2
# Use a neutral sentence with enough phonetic coverage for speaker scoring.
# A two-second greeting is too short for reliable candidate ranking, especially
# for lower-pitched voices whose identity embedding needs more voiced frames.
CALIBRATION_TEXT = "大家好，欢迎来到直播间，今天我们介绍这款车型。"
# A real first live phrase catches clones that pass the tiny "你好" warmup but
# emit silence or an abnormally long semantic tail on normal copy.
CLONE_VALIDATION_TEXT = "大家好，欢迎来到直播间。"
WARMUP_TEXT = "你好。"
CALIBRATION_VERSION = 8
# GPT-SoVITS' v2ProPlus API and WebUI both use conservative sampling as the
# stable starting point. Per-voice calibration can still select a more open
# candidate when its measured speaker similarity is materially better.
DEFAULT_SAMPLING_PROFILE = {"top_p": 0.65, "temperature": 0.60, "repetition_penalty": 1.35}
# Keep calibration to four generations. Each candidate explores a seed
# and a sampling regime at once, so prosody-aware selection remains a
# background optimization rather than a user-facing clone delay.
CALIBRATION_CANDIDATES = (
    # Keep the seed fixed for the first two candidates so their difference is
    # attributable to sampler entropy, not a different semantic trajectory.
    # A lower top-k is useful for low-pitched voices whose consonants tend to
    # drift when several near-equivalent tokens are available.
    {"label": "conservative", "seed_offset": 0, "top_k": 10, "top_p": 0.55, "temperature": 0.48, "repetition_penalty": 1.35},
    {"label": "balanced", "seed_offset": 0, "top_k": 15, "top_p": 0.65, "temperature": 0.60, "repetition_penalty": 1.35},
    {"label": "stable", "seed_offset": 7919, "top_k": 20, "top_p": 0.72, "temperature": 0.66, "repetition_penalty": 1.35},
    # Keep one wider candidate for voices that lose important consonants when
    # sampling is too focused; it is selected only when scoring supports it.
    {"label": "expressive", "seed_offset": 1543, "top_k": 25, "top_p": 0.82, "temperature": 0.72, "repetition_penalty": 1.35},
)
TTS_LOCK = threading.Lock()
# GPT-SoVITS exposes one process-global inference worker.  A plain mutex lets
# the background calibration thread win the race immediately before a live
# phrase, which turns a normally fast request into a weight-switch outlier.
# Keep all lock acquisition behind this condition so foreground work always
# wins over optional warmup/calibration work.
TTS_SCHEDULER = threading.Condition()
TTS_FOREGROUND_WAITERS = 0
TTS_FOREGROUND_ACTIVE = 0
TTS_BACKGROUND_ACTIVE = 0
_ACTIVE_MODEL_PROFILE = None
_TTS_LAST_SUCCESS_AT = 0.0
_TTS_LAST_FOREGROUND_AT = 0.0
TTS_WARMUP = threading.Event()
LIVE_PATH_WARMED = threading.Event()
LIVE_PATH_WARMUP_ERROR = ''
LIVE_PATH_WARMUP_MS = None
VOICE_WARMING = set()
VOICE_WARMED = set()
VOICE_PROFILE_WARMING = set()
VOICE_CALIBRATING = set()
VOICE_CALIBRATION_QUEUE = set()
VOICE_CALIBRATION_FORCED = set()
VOICE_CALIBRATION_QUEUE_LOCK = threading.Lock()
VOICE_CALIBRATION_WORKER_RUNNING = False
CALIBRATION_DEFERRED = object()



def _acquire_tts_lock(priority: str = "foreground", timeout: float | None = None):
    """Acquire the single GPT-SoVITS worker with foreground priority.

    All inference paths in this process go through this helper.  Acquiring the
    raw mutex while holding ``TTS_SCHEDULER`` makes the foreground waiter check
    and lock acquisition atomic, so calibration cannot slip in between them.
    """
    if priority not in {"foreground", "background"}:
        raise ValueError(f"unknown TTS lock priority: {priority}")
    deadline = None if timeout is None else time.monotonic() + max(0.0, timeout)
    global TTS_FOREGROUND_WAITERS, TTS_FOREGROUND_ACTIVE, TTS_BACKGROUND_ACTIVE
    with TTS_SCHEDULER:
        if priority == "foreground":
            TTS_FOREGROUND_WAITERS += 1
            acquired = False
            try:
                while TTS_LOCK.locked():
                    remaining = None if deadline is None else deadline - time.monotonic()
                    if remaining is not None and remaining <= 0:
                        return False
                    TTS_SCHEDULER.wait(timeout=remaining)
                TTS_LOCK.acquire()
                TTS_FOREGROUND_ACTIVE += 1
                acquired = True
                return True
            finally:
                TTS_FOREGROUND_WAITERS -= 1
                if not acquired:
                    TTS_SCHEDULER.notify_all()
        while TTS_LOCK.locked() or TTS_FOREGROUND_WAITERS or TTS_FOREGROUND_ACTIVE:
            remaining = None if deadline is None else deadline - time.monotonic()
            if remaining is not None and remaining <= 0:
                return False
            TTS_SCHEDULER.wait(timeout=remaining)
        TTS_LOCK.acquire()
        TTS_BACKGROUND_ACTIVE += 1
        return True


def _release_tts_lock(priority: str = "foreground"):
    global TTS_FOREGROUND_ACTIVE, TTS_BACKGROUND_ACTIVE
    with TTS_SCHEDULER:
        if not TTS_LOCK.locked():
            return
        TTS_LOCK.release()
        if priority == "background":
            TTS_BACKGROUND_ACTIVE = max(0, TTS_BACKGROUND_ACTIVE - 1)
        elif TTS_FOREGROUND_ACTIVE:
            TTS_FOREGROUND_ACTIVE -= 1
        TTS_SCHEDULER.notify_all()


@contextmanager
def _tts_lock(*, priority: str = "foreground", timeout: float | None = None):
    acquired = _acquire_tts_lock(priority, timeout)
    if not acquired:
        raise TimeoutError("TTS inference worker is busy")
    try:
        yield
    finally:
        _release_tts_lock(priority)


@contextmanager
def _foreground_ticket():
    """Keep optional background jobs paused for one complete user request."""
    global TTS_FOREGROUND_ACTIVE, _TTS_LAST_FOREGROUND_AT
    with TTS_SCHEDULER:
        _TTS_LAST_FOREGROUND_AT = time.monotonic()
        TTS_FOREGROUND_ACTIVE += 1
        TTS_SCHEDULER.notify_all()
    try:
        yield
    finally:
        with TTS_SCHEDULER:
            _TTS_LAST_FOREGROUND_AT = time.monotonic()
            TTS_FOREGROUND_ACTIVE = max(0, TTS_FOREGROUND_ACTIVE - 1)
            TTS_SCHEDULER.notify_all()


@asynccontextmanager
async def _async_tts_lock(timeout=20):
    """Acquire without an orphan worker thread when a stream is cancelled."""
    deadline = time.monotonic() + timeout
    while not _acquire_tts_lock("foreground", 0):
        if time.monotonic() >= deadline:
            raise TimeoutError("语音引擎繁忙，请稍后重试")
        await asyncio.sleep(0.02)
    try:
        yield
    finally:
        _release_tts_lock()


def _tts_has_reference():
    global_reference = _tts_global_reference_audio()
    if global_reference and Path(global_reference).is_file():
        quality = _audio_quality(global_reference)
        if quality.get("status") in {"ready", "unverified"}:
            return True
    try:
        with conn() as c:
            rows = c.execute("SELECT reference_path,prompt_text FROM voices WHERE reference_path<>''").fetchall()
        return any(
            Path(row[0]).is_file()
            and _audio_quality(row[0]).get("status") == "ready"
            and row[1].strip()
            for row in rows
        )
    except Exception:
        return False


def _gpt_sovits_reachable():
    # GPT-SoVITS runs one inference worker. During clone calibration its HTTP
    # loop can be busy long enough for a separate /docs probe to time out even
    # though synthesis is healthy. Preserve a recent proven success so a page
    # refresh cannot incorrectly switch the studio to browser speech.
    if not settings.gpt_sovits_url:
        return False
    if time.monotonic() - _TTS_LAST_SUCCESS_AT < 300:
        return True
    return gpt_sovits_engine.probe(settings)


def _gpt_sovits_runtime_status():
    """Read the upstream process state for diagnostics and profile evidence."""
    if not settings.gpt_sovits_url:
        return {}
    base = settings.gpt_sovits_url.rstrip('/')
    try:
        with httpx.Client(timeout=httpx.Timeout(1.5, connect=0.4), trust_env=False) as client:
            response = client.get(base + '/runtime/status')
            response.raise_for_status()
            payload = response.json()
            return payload if isinstance(payload, dict) else {}
    except (httpx.HTTPError, ValueError, TypeError):
        return {}


def _mark_tts_success():
    global _TTS_LAST_SUCCESS_AT
    _TTS_LAST_SUCCESS_AT = time.monotonic()


def _tts_is_ready():
    return _tts_has_reference() and _gpt_sovits_reachable()


def extract(path: Path):
    if path.suffix.lower() == ".txt":
        return [(None, path.read_text(encoding="utf-8-sig", errors="ignore"))]
    if path.suffix.lower() == ".docx":
        from docx import Document

        d = Document(path)
        blocks = [p.text.strip() for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(c for c in cells if c))
        return [(None, "\n".join(blocks))]
    if path.suffix.lower() == ".pdf":
        import fitz

        with fitz.open(path) as d:
            return [(i + 1, p.get_text()) for i, p in enumerate(d)]
    raise ValueError("仅支持 PDF、DOCX、TXT")


def _write_chunks(c, document_id, source):
    count = 0
    document = c.execute('SELECT brand,series,year FROM documents WHERE id=?', (document_id,)).fetchone()
    for page, text in extract(source):
        records = build_chunks(redact(text), document['brand'], document['series'], document['year'])
        vectors = encode_vectors([r['search_text'] for r in records])
        for chunk, vector in zip(records, vectors):
            chunk['metadata']['parent_key'] = f"{page}:{chunk['metadata']['parent_key']}"
            c.execute(
                "INSERT INTO chunks(document_id,content,tokens,page,embedding,embedding_model,search_text,parent_content,metadata) VALUES(?,?,?,?,?,?,?,?,?)",
                (document_id, chunk['content'], " ".join(tokens(chunk['search_text'])), page, vector, INDEX_VERSION,
                 chunk['search_text'], chunk['parent_content'], json.dumps(chunk['metadata'], ensure_ascii=False)),
            )
            count += 1
    if count == 0:
        raise ValueError("资料内容为空，无法建立知识片段")
    return count


def ingest(source: Path, name: str, brand: str = "", series: str = "", year: str = ""):
    name=redact(name)
    created = now()
    with conn() as c:
        cur = c.execute(
            "INSERT INTO documents(name,path,type,size,brand,series,year,chunks,version,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
            (name, str(source), source.suffix[1:].upper(), source.stat().st_size, brand, series, year, 0, 1, created, created),
        )
        did = cur.lastrowid
        count = _write_chunks(c, did, source)
        c.execute("UPDATE documents SET chunks=? WHERE id=?", (count, did))
        c.execute(
            "INSERT INTO document_versions(document_id,version,name,path,size,chunks,created_at) VALUES(?,?,?,?,?,?,?)",
            (did, 1, name, str(source), source.stat().st_size, count, created),
        )
    return did


def replace_document(did: int, source: Path, name: str, brand: str, series: str, year: str):
    with conn() as c:
        row = c.execute("SELECT version FROM documents WHERE id=?", (did,)).fetchone()
        if not row:
            raise HTTPException(404, "资料不存在")
        version = int(row["version"]) + 1
        c.execute('UPDATE documents SET brand=?,series=?,year=? WHERE id=?', (brand, series, year, did))
        c.execute("DELETE FROM chunks WHERE document_id=?", (did,))
        count = _write_chunks(c, did, source)
        updated = now()
        c.execute(
            "UPDATE documents SET name=?,path=?,type=?,size=?,brand=?,series=?,year=?,chunks=?,version=?,updated_at=? WHERE id=?",
            (name, str(source), source.suffix[1:].upper(), source.stat().st_size, brand, series, year, count, version, updated, did),
        )
        c.execute(
            "INSERT INTO document_versions(document_id,version,name,path,size,chunks,created_at) VALUES(?,?,?,?,?,?,?)",
            (did, version, name, str(source), source.stat().st_size, count, updated),
        )
    return did


async def _warm_live_path_once():
    """Dispatch through the actual ASGI route and fully drain its WAV/PCM stream.

    No browser and no analytics event endpoint is involved. ASGI dispatch avoids
    assuming an externally configured backend port, while executing the same
    async transport, lock, weight check, decoder and WAV forwarding as playback.
    """
    global LIVE_PATH_WARMUP_MS
    start = time.perf_counter()
    # Chinese-only warmup left the first mixed-language answer at 4.285 s
    # (repeat: 1.394 s). Include the automotive acronyms used in real playback
    # so the English text frontend also executes before declaring readiness.
    body = {'voice_id':'steady', 'text':'这款EV车型的CLTC续航为五百八十公里。', 'unitized':True,
            'stream_batch':True, 'speed_factor':1.0, 'delivery':'natural'}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://warmup', trust_env=False) as client:
        async with client.stream('POST', '/api/tts/stream', json=body) as response:
            response.raise_for_status()
            header, received, offset = bytearray(), 0, None
            async for chunk in response.aiter_bytes():
                received += len(chunk)
                if offset is None:
                    header.extend(chunk)
                    offset = _wav_data_offset(header)
                    if offset is not None: header.clear()
            if offset is None or received-offset < 256:
                raise RuntimeError('实时预热未收到有效 PCM')
    LIVE_PATH_WARMUP_MS = round((time.perf_counter()-start)*1000, 1)
    LIVE_PATH_WARMED.set()
    VOICE_WARMED.add('steady')
    logging.getLogger('uvicorn.error').info('live_path_warmup complete: %.1f ms', LIVE_PATH_WARMUP_MS)


async def _warmup_tts():
    """Warm the engine once and invalidate readiness after a runtime restart."""
    global LIVE_PATH_WARMUP_ERROR, _ACTIVE_MODEL_PROFILE, _TTS_LAST_SUCCESS_AT
    instance = None
    if not settings.gpt_sovits_url:
        TTS_WARMUP.set()
        return
    async with httpx.AsyncClient(timeout=2, trust_env=False) as probe:
        while True:
            try:
                response = await probe.get(settings.gpt_sovits_url.rstrip('/') + '/runtime/status')
                response.raise_for_status()
                current = response.json().get('instance_id', 'legacy')
                if current != instance:
                    LIVE_PATH_WARMED.clear()
                    TTS_WARMUP.clear()
                    VOICE_WARMED.clear()
                    _ACTIVE_MODEL_PROFILE = None
                    _TTS_LAST_SUCCESS_AT = 0
                    instance = current
                if not LIVE_PATH_WARMED.is_set():
                    LIVE_PATH_WARMUP_ERROR = ''
                    await _warm_live_path_once()
                    TTS_WARMUP.set()
                    _schedule_incomplete_calibrations()
            except httpx.ReadTimeout:
                # A busy single-worker runtime can delay status replies. This
                # is not evidence of a restart and must not invalidate warmth.
                pass
            except Exception as exc:
                # ASGI stream failures may be wrapped in an ExceptionGroup.
                # Keep the retry task alive and expose the concrete failure.
                LIVE_PATH_WARMUP_ERROR = f'{type(exc).__name__}: {exc}'[:300]
                LIVE_PATH_WARMED.clear()
                TTS_WARMUP.clear()
                _TTS_LAST_SUCCESS_AT = 0
                logging.getLogger('uvicorn.error').warning('live_path_warmup pending: %s', exc)
            await asyncio.sleep(2)


@app.get("/api/health")
def health():
    return {"status": "ok", "service": "汽车直播智能体", "version": "0.3.0"}


@app.get("/api/live2d/status")
def live2d_status():
    root = Path(settings.live2d_model_root).expanduser()
    model = root / settings.live2d_model_file
    return {
        "enabled": bool(settings.live2d_enabled),
        "configured": bool(settings.live2d_enabled and model.is_file()),
        "name": "胡桃",
        "model_url": "/api/live2d/assets/" + settings.live2d_model_file.replace("\\", "/"),
    }


@app.get("/api/live2d/assets/{asset_path:path}")
def live2d_asset(asset_path: str):
    if not settings.live2d_enabled:
        raise HTTPException(404, "Live2D 功能未启用")
    path = _live2d_asset_path(asset_path)
    return FileResponse(path)


@app.get("/api/config/status")
def config():
    tts = {"mode": "gpt-sovits", "provider": "GPT-SoVITS", "ready": _tts_is_ready()}
    llm_ready = bool(_llm_endpoint() and settings.llm_model)
    return {
        "llm": llm_gateway.status(),
        "tts": tts,
        "retrieval": index_status(),
    }


@app.get("/api/tts/status")
def tts_status():
    configured = _tts_has_reference()
    reachable = _gpt_sovits_reachable()
    runtime = _gpt_sovits_runtime_status()
    # Profile priming is an opportunistic foreground warm-up. It may briefly
    # hold the inference lock, but it must not make the whole provider look
    # unavailable and force the browser into a long polling loop.
    warming_up = not TTS_WARMUP.is_set() or bool(VOICE_WARMING)
    # Keep the provider identity stable while the model is loading. `ready`
    # means an inference can start immediately; it must not be confused with
    # the older browser fallback state during a normal page refresh.
    ready = configured and reachable and LIVE_PATH_WARMED.is_set() and not voice_adaptation.BUSY.is_set()
    with conn() as c:
        rows = c.execute("SELECT reference_path,prompt_text FROM voices WHERE reference_path<>''").fetchall()
    cloned = sum(
        bool(row[1].strip() and Path(row[0]).is_file() and _audio_quality(row[0]).get("status") == "ready")
        for row in rows
    )
    return {
        "provider": "gpt-sovits",
        "provider_label": "GPT-SoVITS",
        "ready": ready,
        "configured": configured,
        "reachable": reachable,
        "cloned_voices": cloned,
        "endpoint": settings.gpt_sovits_url or None,
        "model_version": settings.gpt_sovits_model_version,
        "streaming": ready,
        "stream_mode": "pcm-wav-natural",
        "streaming_mode": settings.gpt_sovits_streaming_mode,
        "live_streaming_mode": settings.gpt_sovits_live_streaming_mode,
        "live_path_warmed": LIVE_PATH_WARMED.is_set(),
        "live_path_warmup_ms": LIVE_PATH_WARMUP_MS,
        "live_path_warmup_error": LIVE_PATH_WARMUP_ERROR or None,
        "active_model_profile": _ACTIVE_MODEL_PROFILE or "unknown",
        "runtime_instance_id": runtime.get("instance_id"),
        "runtime_gpt_weights": runtime.get("gpt_weights"),
        "runtime_sovits_weights": runtime.get("sovits_weights"),
        "runtime_reference_cache_size": runtime.get("reference_cache_size"),
        "runtime_reference_cache_hits": runtime.get("reference_cache_hits"),
        "model_profiles": list(model_profiles(settings)),
        "warming_up": warming_up,
        # Calibration is an optional background refinement. It must be
        # observable for the UI, but it must not make an already warmed clone
        # appear unavailable or route it to browser speech.
        "calibrating_voice_ids": sorted(VOICE_CALIBRATING),
        "warmed_voice_ids": sorted(VOICE_WARMED),
        "prompt_checks": dict(prompt_guard.STATUS),
        "adapting": voice_adaptation.BUSY.is_set(),
        "adaptation_enabled": settings.gpt_sovits_adaptation_enabled,
    }


def _audio_quality(path: str):
    try:
        file=Path(path)
        stat=file.stat()
        return deepcopy(_cached_audio_quality(str(file.resolve()),stat.st_mtime_ns,stat.st_size))
    except (OSError,ValueError,TypeError):
        return _inspect_audio_quality(path)


@lru_cache(maxsize=96)
def _cached_audio_quality(path,mtime,size):
    return _inspect_audio_quality(path)


def _inspect_audio_quality(path: str):
    """Inspect WAV references before they reach GPT-SoVITS.

    A surprising number of phone exports carry an MP4/AAC payload behind a
    .wav suffix. GPT-SoVITS may decode such a file, but speaker identity and
    prosody become noticeably worse. We report that state instead of silently
    accepting a misleading reference.
    """
    if not path:
        return {"status": "missing", "message": "未配置参考音频"}
    audio_path = Path(path)
    if not audio_path.exists():
        return {"status": "missing", "message": "参考音频文件不存在"}
    if audio_path.suffix.lower() != ".wav":
        return {"status": "unverified", "message": "非 WAV 格式，将由 GPT-SoVITS 解码"}
    try:
        with wave.open(str(audio_path), "rb") as stream:
            channels = stream.getnchannels()
            sample_width = stream.getsampwidth()
            sample_rate = stream.getframerate()
            duration = stream.getnframes() / float(sample_rate or 1)
            frames = stream.readframes(min(stream.getnframes(), sample_rate * 30))
    except (wave.Error, EOFError, OSError):
        return {"status": "invalid", "message": "不是标准 PCM WAV，请重新导出或上传 MP3/WAV 音频"}
    active_ratio = None
    rms_db = None
    peak_db = None
    longest_silence = None
    silence_ratio = None
    clipped_ratio = None
    snr_db = None
    if sample_width == 2 and frames:
        import array

        samples = array.array("h")
        samples.frombytes(frames)
        if samples:
            peak = max(abs(sample) for sample in samples) or 1
            rms = math.sqrt(sum(sample * sample for sample in samples) / len(samples)) or 1
            active_threshold = 32768 * (10 ** (-45 / 20))
            active_ratio = sum(abs(sample) >= active_threshold for sample in samples) / len(samples)
            clipped_ratio = sum(abs(sample) >= 32700 for sample in samples) / len(samples)
            rms_db = 20 * math.log10(rms / 32768)
            peak_db = 20 * math.log10(peak / 32768)
            # Long internal silence is copied by GPT-SoVITS as unnatural gaps.
            # Measure short RMS windows instead of individual samples so quiet
            # consonants are not mistaken for silence.
            frame_size = max(1, int(sample_rate * 0.02))
            quiet_threshold = 32768 * (10 ** (-42 / 20))
            quiet_frames = []
            frame_rms_values = []
            for start in range(0, len(samples), frame_size):
                window = samples[start:start + frame_size]
                if not window:
                    continue
                window_rms = math.sqrt(sum(sample * sample for sample in window) / len(window))
                frame_rms_values.append(window_rms)
                quiet_frames.append(window_rms < quiet_threshold)
            silence_ratio = sum(quiet_frames) / len(quiet_frames) if quiet_frames else 0
            longest_frames = current_frames = 0
            for is_quiet in quiet_frames:
                if is_quiet:
                    current_frames += 1
                    longest_frames = max(longest_frames, current_frames)
                else:
                    current_frames = 0
            longest_silence = longest_frames * 0.02
            active_rms_values = [value for value, quiet in zip(frame_rms_values, quiet_frames) if not quiet]
            quiet_rms_values = [value for value, quiet in zip(frame_rms_values, quiet_frames) if quiet]
            if active_rms_values:
                active_sorted = sorted(active_rms_values)
                quiet_sorted = sorted(quiet_rms_values)
                speech_level = active_sorted[int(len(active_sorted) * 0.75)]
                # Only estimate SNR when the recording provides a genuine
                # quiet floor. A clip made of continuous speech (or a flat
                # synthetic fixture) has no reliable noise reference and must
                # not be rejected based on an arbitrary low-percentile frame.
                if quiet_sorted:
                    noise_level = quiet_sorted[len(quiet_sorted) // 2]
                    snr_db = 20 * math.log10((speech_level + 1) / (noise_level + 1))
    issues = []
    if channels != 1:
        issues.append("建议单声道")
    if sample_width != 2:
        issues.append("建议 16-bit")
    if not 16000 <= sample_rate <= 48000:
        issues.append("采样率建议 16kHz-48kHz")
    if not REFERENCE_MIN_SECONDS <= duration <= REFERENCE_MAX_SECONDS:
        issues.append("时长建议 3-10 秒")
    if active_ratio is not None and active_ratio < 0.45:
        issues.append("有效人声比例偏低，建议去除长静音")
    if rms_db is not None and rms_db < -35:
        issues.append("人声音量偏低，建议统一响度")
    warnings = []
    if REFERENCE_MIN_SECONDS <= duration < REFERENCE_RECOMMENDED_MIN_SECONDS:
        warnings.append("参考音频少于 5 秒，建议补充到 5-10 秒以提高音色稳定性")
    if snr_db is not None:
        if snr_db < 8:
            issues.append("背景噪声过高，建议使用单人、安静环境录音")
        elif snr_db < 14:
            warnings.append("背景噪声偏高，可能降低克隆音色相似度")
    if longest_silence is not None:
        if longest_silence > 1.5:
            issues.append("检测到过长内部静音，建议重新录制连续人声")
        elif longest_silence > 0.55:
            # Keep a natural phrase break. Re-timing the middle of a prompt
            # changes the alignment between prompt_text and the waveform.
            warnings.append("保留了较长自然停顿，请确认参考文本中的标点与录音一致")
    if clipped_ratio is not None and clipped_ratio > 0.01:
        issues.append("检测到录音削波失真，建议降低麦克风增益后重新录制")
    messages = issues or warnings
    return {
        "status": "ready" if not issues else "needs-review",
        "message": "标准 PCM WAV" if not messages else "；".join(messages),
        "duration": round(duration, 2),
        "channels": channels,
        "sample_width": sample_width,
        "sample_rate": sample_rate,
        "active_ratio": round(active_ratio, 3) if active_ratio is not None else None,
        "silence_ratio": round(silence_ratio, 3) if silence_ratio is not None else None,
        "longest_silence": round(longest_silence, 2) if longest_silence is not None else None,
        "clipped_ratio": round(clipped_ratio, 4) if clipped_ratio is not None else None,
        "rms_db": round(rms_db, 1) if rms_db is not None else None,
        "peak_db": round(peak_db, 1) if peak_db is not None else None,
        "snr_db": round(snr_db, 1) if snr_db is not None else None,
        "warnings": warnings,
    }


def _probe_audio_payload(payload: bytes, *, expected_text: str | None = None):
    """Validate a generated WAV as speech, rather than merely as a file.

    Reference checks intentionally allow a 3-10 second recording. Generated
    probes need different bounds: a normal greeting can be shorter, while a
    runaway semantic sequence can be many seconds long. The signal checks
    below reject both the silent/"sigh" output and the pathological long tail.
    """
    if not payload or len(payload) <= 128:
        return {"status": "failed", "message": "GPT-SoVITS 未返回有效音频"}
    try:
        with wave.open(io.BytesIO(payload), "rb") as stream:
            channels = stream.getnchannels()
            sample_width = stream.getsampwidth()
            sample_rate = stream.getframerate()
            frame_count = stream.getnframes()
            duration = frame_count / float(sample_rate or 1)
            frames = stream.readframes(frame_count)
            duration = len(frames) / float(sample_rate * channels * sample_width or 1)
    except (wave.Error, EOFError, OSError):
        return {"status": "failed", "message": "GPT-SoVITS 返回的不是标准 WAV 音频"}
    if sample_width != 2 or not sample_rate or not frames:
        return {"status": "failed", "message": "GPT-SoVITS 返回的音频格式不可播放"}
    import array

    samples = array.array("h")
    samples.frombytes(frames)
    if channels > 1:
        samples = array.array("h", samples[::channels])
    if not samples:
        return {"status": "failed", "message": "GPT-SoVITS 返回空音频"}
    peak = max(abs(sample) for sample in samples) or 1
    rms = math.sqrt(sum(sample * sample for sample in samples) / len(samples)) or 1
    threshold = 32768 * (10 ** (-45 / 20))
    active_ratio = sum(abs(sample) >= threshold for sample in samples) / len(samples)
    frame_size = max(1, int(sample_rate * 0.02))
    quiet_threshold = 32768 * (10 ** (-42 / 20))
    quiet_frames = []
    for start in range(0, len(samples), frame_size):
        window = samples[start:start + frame_size]
        if window:
            quiet_frames.append(math.sqrt(sum(sample * sample for sample in window) / len(window)) < quiet_threshold)
    longest_frames = current_frames = 0
    for quiet in quiet_frames:
        if quiet:
            current_frames += 1
            longest_frames = max(longest_frames, current_frames)
        else:
            current_frames = 0
    silence_ratio = sum(quiet_frames) / len(quiet_frames) if quiet_frames else 1.0
    rms_db = 20 * math.log10(rms / 32768)
    peak_db = 20 * math.log10(peak / 32768)
    issues = []
    expected = max(.6, len(re.sub(r'[\W_]+', '', expected_text, flags=re.UNICODE))/4.5) if expected_text else None
    minimum = max(.25, expected*.3) if expected is not None else .6
    maximum = prompt_guard.duration_limit(expected_text) if expected_text else 12.0
    if duration < minimum:
        issues.append("合成结果过短，几乎没有有效语音")
    elif duration > maximum:
        issues.append("合成结果异常过长，疑似语义推理失控")
    if active_ratio < 0.20:
        issues.append("有效人声比例过低，可能只有叹息或静音")
    if rms_db < -38:
        issues.append("合成音量过低，未检测到稳定人声")
    if silence_ratio > 0.82:
        issues.append("合成结果包含过多静音")
    return {
        "status": "ready" if not issues else "failed",
        "message": "生成音频通过人声验收" if not issues else "；".join(issues),
        "duration": round(duration, 2),
        "active_ratio": round(active_ratio, 3),
        "silence_ratio": round(silence_ratio, 3),
        "longest_silence": round(longest_frames * 0.02, 2),
        "rms_db": round(rms_db, 1),
        "peak_db": round(peak_db, 1),
    }


def _ffmpeg_binary():
    """Locate ffmpeg binary from PATH, environment variable, or GPT-SoVITS runtime."""
    # Use the same discovery as startup, including installations found beside
    # the project when GPT_SOVITS_ROOT and PATH have not been configured.
    binary = detect_ffmpeg()
    return str(binary) if binary else None


def _decoded_peak_db(source: Path):
    """Measure decoded peak for compressed references (OGG/MP3/M4A)."""
    ffmpeg = _ffmpeg_binary()
    if not ffmpeg:
        return None
    try:
        completed = subprocess.run(
            [ffmpeg, "-hide_banner", "-loglevel", "error", "-i", str(source),
             "-vn", "-ac", "1", "-ar", "24000", "-f", "s16le", "-"],
            capture_output=True, timeout=45,
        )
    except (OSError, subprocess.TimeoutExpired):
        return None
    if completed.returncode != 0 or not completed.stdout:
        return None
    import array

    samples = array.array("h")
    samples.frombytes(completed.stdout[: len(completed.stdout) - (len(completed.stdout) % 2)])
    peak = max((abs(sample) for sample in samples), default=0)
    return 20 * math.log10(peak / 32768) if peak else None


@lru_cache(maxsize=128)
def _decoded_levels(path: str, mtime: int, size: int):
    """Measure the same mono/24kHz signal used for upload and microphone input."""
    ffmpeg = _ffmpeg_binary()
    if not ffmpeg:
        quality = _audio_quality(path)
        return {key: quality.get(key) for key in ('peak_db', 'rms_db')}
    try:
        result = subprocess.run([ffmpeg, '-hide_banner', '-loglevel', 'error', '-i', path,
            '-vn', '-ac', '1', '-ar', '24000', '-f', 'f32le', '-'], capture_output=True, timeout=45)
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise HTTPException(422, '参考音频电平检测失败，请检查音频文件或 ffmpeg 后重试') from exc
    if result.returncode or not result.stdout:
        raise HTTPException(422, '参考音频无法解码或预处理失败')
    import array
    samples = array.array('f', result.stdout)
    peak = max(map(abs, samples), default=0)
    rms = math.sqrt(sum(x*x for x in samples) / max(1, len(samples)))
    return {'peak_db': 20*math.log10(max(peak, 1e-10)), 'rms_db': 20*math.log10(max(rms, 1e-10))}


def _reference_levels(path: Path):
    stat = path.stat()
    return _decoded_levels(str(path.resolve()), stat.st_mtime_ns, stat.st_size)


@lru_cache(maxsize=1)
def _reference_target_rms():
    # Six packaged references measured in docs/LIVE_TTS_AUDIT_20260918.md:
    # Median of quality's one-decimal RMS readings: -18.05 dBFS. Compute from
    # the installed catalog, not speaker gender (full precision: -18.0395).
    levels = [_audio_quality(path).get('rms_db') for voice in PRESET_VOICES for path in preset_reference_paths(voice)]
    values = [level for level in levels if level is not None]
    return round(statistics.median(values), 2) if values else -20.0


def _reference_quality(path: Path):
    quality = _audio_quality(str(path))
    if not path.is_file() or quality.get('status') == 'invalid':
        return quality
    original = _original_reference_audio(path)
    if original.is_file():
        before = _reference_levels(original)
        after = _reference_levels(path)
        quality['normalization'] = {
            'target_rms_db': _reference_target_rms(), 'peak_ceiling_db': -3.0,
            'before': {k: round(v, 2) if v is not None else None for k,v in before.items()},
            'after': {k: round(v, 2) if v is not None else None for k,v in after.items()},
            'rms_change_db': round(after['rms_db']-before['rms_db'], 2) if after['rms_db'] is not None and before['rms_db'] is not None else None,
        }
    return quality


def _normalize_prompt_text(text: str):
    # Line breaks copied from subtitles are layout, not spoken pauses. Remove
    # whitespace around CJK/Latin boundaries while retaining spaces between
    # English words, which the GPT-SoVITS frontend needs for pronunciation.
    text = str(text or "").strip()
    text = re.sub(r"(?<=[\u3400-\u9fff])\s+(?=[A-Za-z0-9])", "", text)
    text = re.sub(r"(?<=[A-Za-z0-9])\s+(?=[\u3400-\u9fff])", "", text)
    text = re.sub(r"(?<=[\u3400-\u9fff])\s+(?=[\u3400-\u9fff])", "", text)
    text = re.sub(r"\s+", " ", text)
    if text and text[-1] not in "。！？；…,.!?;":
        text += "。"
    return text


def _prompt_alignment(prompt_text: str, duration: float | None):
    """Return non-blocking guidance when the transcript density looks wrong.

    GPT-SoVITS needs the supplied text to describe the reference audio exactly.
    We cannot prove semantic equality without ASR, but a speaking-rate check
    catches the common mistake of pasting a paragraph for a short clip.
    """
    if not prompt_text or not duration or duration <= 0:
        return {"status": "unknown", "message": "请确认参考文本与录音逐字一致"}
    speech_chars = len(re.sub(r"[\s\W_]+", "", prompt_text, flags=re.UNICODE))
    if not speech_chars:
        return {"status": "invalid", "message": "参考文本需要包含实际说出的汉字或数字"}
    chars_per_second = speech_chars / duration
    if chars_per_second < 1.4:
        return {"status": "review", "message": "参考文本相对录音偏短；请确认没有漏写句子或保留过长停顿", "chars_per_second": round(chars_per_second, 2)}
    if chars_per_second > 8.5:
        return {"status": "review", "message": "参考文本相对录音偏长；请确认没有粘贴未录入的内容", "chars_per_second": round(chars_per_second, 2)}
    return {"status": "ok", "message": "参考文本长度与录音时长匹配", "chars_per_second": round(chars_per_second, 2)}


def _original_reference_audio(source: Path) -> Path:
    """Resolve a generated reference derivative back to its uploaded source."""
    stem = source.stem
    base = stem
    bases = [stem]
    derivative = re.compile(r"(?:\.normalized|\.clean|\.optimized\d*|\.processed|\.denoised|\.lite|\.pauses|\.level-v2)$", re.IGNORECASE)
    while True:
        stripped = derivative.sub("", base)
        if stripped == base:
            break
        base = stripped
        bases.append(base)
    if len(bases) == 1:
        return source
    for candidate_base in sorted(set(bases), key=len):
        for suffix in sorted(ALLOWED_AUDIO):
            candidate = source.with_name(candidate_base + suffix)
            if candidate.is_file():
                return candidate
    return source


def _normalize_reference_audio(source: Path, *, force: bool = True):
    """Create a model-compatible WAV while preserving the speaker recording.

    GPT-SoVITS already rescales and resamples the reference internally.  Fixed
    FFT denoising, low-pass filtering, and aggressive
    silence removal are destructive here: they change formants and often turn
    a clean recording into the electronic hiss heard in the clone.  The
    preprocessing step therefore only removes clearly excessive silence,
    decodes the source, folds it to mono, and writes deterministic 24 kHz/
    16-bit PCM with measured RMS alignment and peak protection.
    """
    quality = _audio_quality(str(source))
    levels = _reference_levels(source)
    source_peak_db, source_rms_db = levels['peak_db'], levels['rms_db']
    if source_rms_db is not None and source_rms_db < -45:
        raise HTTPException(422, '参考音频人声电平过低，请靠近麦克风重新录制')
    # Match the packaged references' measured median RMS. A lookahead limiter
    # protects peaks; no denoising, pitch changes or internal pause compression.
    gain_db = _reference_target_rms() - source_rms_db if source_rms_db is not None else 0.0
    peak_requires_adjustment = abs(gain_db) >= 0.1
    if (
        quality["status"] == "ready"
        and quality.get("channels") == 1
        and quality.get("sample_rate") == 24000
        and quality.get("sample_width") == 2
        and not peak_requires_adjustment
        and (source_peak_db is None or source_peak_db <= -3.0)
    ):
        return source
    ffmpeg = _ffmpeg_binary()
    if not ffmpeg:
        if quality["status"] == "ready":
            return source
        raise HTTPException(422, "找不到音频预处理工具，无法把参考音频转换为标准 PCM WAV")
    target = source.with_name(source.stem + ".optimized.wav")
    if target == source:
        target = source.with_name(source.stem + ".processed.wav")
    # Only remove clearly non-speech leading/trailing sections. Internal
    # silences stay untouched: compressing them changes prompt-text alignment
    # and is a common source of metallic or clipped cloning artifacts.
    filter_graph = (
        "silenceremove=start_periods=1:start_duration=0.30:start_threshold=-50dB,"
        "areverse,silenceremove=start_periods=1:start_duration=0.80:start_threshold=-50dB,areverse"
    )
    if peak_requires_adjustment:
        filter_graph += f",volume={gain_db:.2f}dB"
    # Bundled ffmpeg predates alimiter's latency option. Compensate its 5 ms
    # lookahead explicitly, retaining the reference/transcript alignment.
    filter_graph += ',apad=pad_dur=0.005,alimiter=limit=0.7079458:attack=5:release=50:level=false,atrim=start=0.005,asetpts=PTS-STARTPTS'
    try:
        # Trimming and peak limiting change RMS. Re-measure the actual file,
        # adjusting the gain against the original input, never cascading
        # normalization through already limited derivatives.
        for attempt in range(4):
            completed = subprocess.run(
                [ffmpeg, "-y", "-hide_banner", "-loglevel", "error", "-i", str(source),
                 "-af", filter_graph, "-map_metadata", "-1", "-ac", "1", "-ar", "24000", "-sample_fmt", "s16", str(target)],
                capture_output=True, text=True, timeout=90,
            )
            if completed.returncode or not target.is_file():
                break
            actual = _reference_levels(target)['rms_db']
            correction = _reference_target_rms() - actual if actual is not None else 0
            if abs(correction) <= .25 or attempt == 3:
                break
            gain_db += correction
            if ',volume=' in filter_graph:
                filter_graph = re.sub(r',volume=[^,]+', f',volume={gain_db:.3f}dB', filter_graph)
            else:
                filter_graph = filter_graph.replace(',apad=', f',volume={gain_db:.3f}dB,apad=')
    except (OSError, subprocess.TimeoutExpired) as exc:
        target.unlink(missing_ok=True)
        raise HTTPException(422, f"参考音频预处理失败：{exc}") from exc
    if completed.returncode != 0 or not target.is_file():
        target.unlink(missing_ok=True)
        detail = (completed.stderr or "未知错误").strip()[-300:]
        raise HTTPException(422, f"参考音频无法解码或预处理失败：{detail}")
    normalized_quality = _audio_quality(str(target))
    if normalized_quality["status"] == "invalid" or not normalized_quality.get("duration"):
        target.unlink(missing_ok=True)
        raise HTTPException(422, "参考音频预处理后仍不是有效 PCM WAV")
    if normalized_quality.get("duration", 0) < REFERENCE_MIN_SECONDS or normalized_quality.get("duration", 0) > REFERENCE_MAX_SECONDS:
        target.unlink(missing_ok=True)
        raise HTTPException(422, "参考音频有效时长需为 3-10 秒；请提供完整的单人录音")
    return target


def _calibration_status(voice_id: str, status: str, message: str):
    logging.getLogger('uvicorn.error').info('voice_calibration voice=%s status=%s reason=%s', voice_id, status, message)
    with conn() as c:
        c.execute('UPDATE voices SET calibration_status=?,calibration_message=?,calibration_checked_at=? WHERE id=?',
                  (status, message[:1000], now(), voice_id))


def _calibration_complete(row):
    try:
        details = json.loads(row['calibration_details'] or '{}')
    except (TypeError, ValueError):
        details = {}
    return (all(row[key] is not None for key in ('sampling_seed', 'sampling_top_k', 'sampling_top_p',
            'sampling_temperature', 'validated_aux_reference_paths'))
            and row['sampling_model_version'] == settings.gpt_sovits_model_version
            and details.get('calibration_version') == CALIBRATION_VERSION)


def _schedule_incomplete_calibrations():
    if not settings.gpt_sovits_calibration_enabled:
        return
    with conn() as c:
        rows = c.execute("SELECT * FROM voices WHERE reference_path<>'' AND model_profile='base' "
                         "AND synthesis_status='ready'").fetchall()
    for row in rows:
        if row['id'] not in PRESET_VOICES and not _calibration_complete(row):
            _schedule_voice_calibration(row['id'])


def _calibrate_voice(voice_id: str, *, force=False):
    """Filter references and jointly select identity/prosody sampling.

    Calibration is deliberately cooperative with live speech. Older versions
    held ``TTS_LOCK`` while a helper process generated every candidate, which
    made an otherwise ready clone wait behind the whole calibration pass. Each
    candidate now acquires the lock only for its own HTTP request; embedding
    scoring runs outside the lock and cannot interrupt playback.
    """
    if voice_id in PRESET_VOICES or not settings.gpt_sovits_calibration_enabled or voice_id in VOICE_CALIBRATING:
        logging.getLogger('uvicorn.error').debug('voice_calibration skip voice=%s disabled/preset/already-running', voice_id)
        return None
    with conn() as c:
        row = c.execute(
            "SELECT reference_path,model_profile,sampling_seed,validated_aux_reference_paths,aux_reference_paths,prompt_text,"
            "sampling_top_k,sampling_top_p,sampling_temperature,sampling_model_version,calibration_details "
            "FROM voices WHERE id=?",
            (voice_id,),
        ).fetchone()
    try:
        calibration_details = json.loads(row[10] or "{}") if row else {}
    except (TypeError, ValueError):
        calibration_details = {}
    if not row or not row[0] or row[1] != "base" or (not force and
        row[2] is not None and row[3] is not None and row[7] is not None and row[8] is not None
        and row[9] == settings.gpt_sovits_model_version and calibration_details.get("calibration_version") == CALIBRATION_VERSION
    ):
        return None
    script = Path(__file__).with_name("voice_similarity.py")
    python = Path(settings.gpt_sovits_python)
    checkpoint = Path(settings.gpt_sovits_speaker_model)
    if not (script.is_file() and python.is_file() and checkpoint.is_file()):
        _calibration_status(voice_id, 'failed', '校准依赖缺失：请检查 Python、评分脚本与说话人权重路径')
        return None
    try:
        raw_aux = json.loads(row[4] or "[]")
    except (TypeError, ValueError):
        raw_aux = []
    references = [row[0]] + [
        str(path) for path in raw_aux[:MAX_AUX_REFERENCE_AUDIO]
        if isinstance(path, str) and Path(path).is_file() and _audio_quality(path).get("status") == "ready"
    ]
    calibration_signature = (row[0], row[4] or "[]", row[5] or "")
    calibration_dir = settings.upload_dir / "voice-calibration"
    calibration_dir.mkdir(parents=True, exist_ok=True)
    base_seed = _voice_sampling_seed(voice_id)
    candidate_specs = [
        {
            "label": profile["label"],
            "seed": (base_seed + profile["seed_offset"]) & 0x7FFFFFFF,
            "top_k": profile["top_k"],
            "top_p": profile["top_p"],
            "temperature": profile["temperature"],
            "repetition_penalty": DEFAULT_SAMPLING_PROFILE["repetition_penalty"],
        }
        for profile in CALIBRATION_CANDIDATES
    ]
    run_id = uuid4().hex[:10]
    candidate_dir = calibration_dir / f"{voice_id}-{run_id}"
    candidate_dir.mkdir(parents=True, exist_ok=True)
    VOICE_CALIBRATING.add(voice_id)
    _calibration_status(voice_id, 'running', '正在筛选参考与生成候选')
    try:
        def run_similarity(extra_args):
            command = [
                str(python), str(script),
                "--gpt-root", str(speaker_model_root(python, checkpoint)),
                "--checkpoint", str(checkpoint),
                "--aux-threshold", str(settings.gpt_sovits_aux_similarity_threshold),
            ]
            for reference in references:
                command.extend(("--reference", str(reference)))
            command.extend(extra_args)
            completed = subprocess.run(command, capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=120)
            phase = 'scores' if extra_args else 'references'
            (candidate_dir / f'{phase}.json').write_text(json.dumps({
                'command': command, 'returncode': completed.returncode,
                'stdout': completed.stdout, 'stderr': completed.stderr}, ensure_ascii=False, indent=2), encoding='utf-8')
            if completed.returncode != 0:
                raise RuntimeError(f'评分进程失败（{phase}，exit={completed.returncode}）：{completed.stderr[-600:]}')
            lines = (completed.stdout or "").strip().splitlines()
            if not lines:
                raise RuntimeError(f'评分进程没有返回结果（{phase}）')
            return json.loads(lines[-1])

        # Screen auxiliaries before synthesis so a mismatched clip never
        # influences the candidate audio. This subprocess performs only local
        # speaker-embedding work and does not touch the inference worker.
        screened = run_similarity([])
        if not screened:
            raise RuntimeError('参考筛选没有返回结果')
        accepted_aux = screened.get("accepted_aux_paths", [])
        # Identity screening is necessary but does not certify continuity.
        # Keep discontinuous auxiliaries out until their pauses are repaired
        # and the resulting files are screened again.
        withheld_aux = []
        for path in list(accepted_aux):
            metrics = wav_metrics(Path(path).read_bytes())
            if metrics['max_internal_silence_s'] > 0.34:
                accepted_aux.remove(path)
                withheld_aux.append({'path': path, 'reason': 'internal-pause-over-0.34s', **metrics})
        references = [references[0], *accepted_aux]

        request = TTSRequest(text=CALIBRATION_TEXT, voice_id=voice_id)
        base_params = _tts_params(request, seed_override=base_seed)
        base_params["aux_ref_audio_paths"] = accepted_aux
        candidate_dir.mkdir(parents=True, exist_ok=True)
        candidate_args = []
        with httpx.Client(timeout=180, trust_env=False) as client:
            for index, spec in enumerate(candidate_specs):
                _wait_for_calibration_idle()
                payload = dict(base_params)
                payload.update({
                    key: value for key, value in spec.items()
                    if key in {"seed", "top_k", "top_p", "temperature", "repetition_penalty"}
                })
                # Calibration is lower priority than a real preview/live
                # request. Never leave an operator waiting behind a background
                # candidate; the queued worker can retry it after the next
                # startup or a later clone operation.
                try:
                    with _tts_lock(priority="background", timeout=0.25):
                        _ensure_model_profile_loaded(_voice_model_profile(voice_id), client)
                        response = client.post(_tts_endpoint(), json=payload)
                        response.raise_for_status()
                except TimeoutError:
                    # A live request arrived while this candidate was waiting.
                    # Defer the optional refinement instead of adding latency to
                    # the foreground request or abandoning calibration forever.
                    _calibration_status(voice_id, 'queued', '播报优先，校准等待空闲后继续')
                    return CALIBRATION_DEFERRED
                if len(response.content) <= 128:
                    raise RuntimeError('校准候选未返回有效音频')
                label = str(spec.get("label") or spec.get("seed") or index)
                safe_label = re.sub(r"[^A-Za-z0-9_.-]+", "-", label).strip("-.") or str(index)
                candidate_path = candidate_dir / f"candidate-{safe_label}.wav"
                candidate_path.write_bytes(response.content)
                candidate_path.with_suffix('.request.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
                candidate_args.extend(("--candidate", f"{label}={candidate_path}"))
                _mark_tts_success()

        output = run_similarity(
            candidate_args
            + ["--candidate-specs-json", json.dumps(candidate_specs, ensure_ascii=False)]
        )
        if not output:
            raise RuntimeError('候选评分没有返回结果')
        best_candidate = output.get("best_candidate") or {}
        # ``--candidate`` is intentionally a local scoring mode. The mapping
        # above lets the helper include the original sampling profile in its
        # stability tie-break; restore it here as a defensive check before
        # persisting calibration.
        specs_by_label = {str(spec.get("label")): spec for spec in candidate_specs}
        for score in output.get("scores", []):
            label = str(score.get("label", ""))
            if label in specs_by_label:
                score["sampling"] = dict(specs_by_label[label])
        best_sampling = specs_by_label.get(str(best_candidate.get("label"))) or best_candidate.get("sampling") or {}
        best_seed = int(best_sampling.get("seed", base_seed))
        best_top_k = int(best_sampling.get("top_k", 15))
        best_top_p = float(best_sampling.get("top_p", DEFAULT_SAMPLING_PROFILE["top_p"]))
        best_temperature = float(best_sampling.get("temperature", DEFAULT_SAMPLING_PROFILE["temperature"]))
        accepted_aux = output.get("accepted_aux_paths", [])
        with conn() as c:
            current = c.execute(
                "SELECT reference_path,aux_reference_paths,prompt_text FROM voices WHERE id=?",
                (voice_id,),
            ).fetchone()
        if not current or (current[0], current[1] or "[]", current[2] or "") != calibration_signature:
            # The host edited the reference while this background run was in
            # flight. Never apply a result computed for the old speaker data.
            _calibration_status(voice_id, 'queued', '参考已修改，旧校准结果未应用')
            return CALIBRATION_DEFERRED
        calibration_details = {
            "calibration_version": CALIBRATION_VERSION,
            "scores": output.get("scores", []),
            "reference_scores": screened.get("reference_scores", []),
            "withheld_aux": withheld_aux,
            "reference_prosody": output.get("reference_prosody", {}),
            "reference_prosody_risk": output.get("reference_prosody_risk", {}),
            "evidence_dir": str(candidate_dir),
        }
        _save_voice_calibration(
            voice_id,
            best_seed,
            accepted_aux,
            top_k=best_top_k,
            top_p=best_top_p,
            temperature=best_temperature,
            details=calibration_details,
        )
        _mark_tts_success()
        _calibration_status(voice_id, 'complete', '参数与辅助参考筛选结果已保存；听感仍需试听确认')
        return {
            "seed": best_seed,
            "top_k": best_top_k,
            "top_p": best_top_p,
            "temperature": best_temperature,
            "scores": output.get("scores", []),
            "reference_scores": output.get("reference_scores", []),
            "accepted_aux_paths": accepted_aux,
        }
    except (OSError, ValueError, KeyError, IndexError, json.JSONDecodeError, httpx.HTTPError, subprocess.TimeoutExpired, sqlite3.Error, RuntimeError) as exc:
        _calibration_status(voice_id, 'failed', f'{type(exc).__name__}: {exc}')
        logging.getLogger('uvicorn.error').exception('voice_calibration failed voice=%s evidence=%s', voice_id, candidate_dir)
        return None
    finally:
        VOICE_CALIBRATING.discard(voice_id)
        # Keep the four candidates and scorer output so failures and selections
        # can be audited. Repeated page polling never starts another run.


def _checked_voice_probe(client, params, identity):
    """Use the production checked guard for creation and rollback acceptance."""
    async def generate(attempt, max_duration):
        data, limit, too_long = bytearray(), None, False
        with client.stream('POST', _tts_endpoint(), json=attempt, timeout=30) as response:
            response.raise_for_status()
            for chunk in response.iter_bytes():
                if too_long:
                    continue
                data.extend(chunk)
                offset = _wav_data_offset(data)
                if offset is not None and limit is None:
                    with wave.open(io.BytesIO(data), 'rb') as wav:
                        limit = offset + int(max_duration * wav.getframerate() * wav.getnchannels() * wav.getsampwidth())
                if limit is not None and len(data) > limit:
                    too_long = True
                    data.clear()
        if too_long:
            raise prompt_guard.RunawayAudio('验收音频超过文本时长上限')
        return bytes(data)
    return asyncio.run(prompt_guard.checked_sentence(params, generate, _probe_audio_payload, identity))


def _warm_single_voice(voice_id: str):
    """Warm one clone in the selected engine."""
    if _tts_provider() == 'gpt-sovits' and voice_adaptation.run_pending(voice_id):
        return
    if voice_id in VOICE_WARMED or voice_id in VOICE_WARMING:
        if voice_id in VOICE_WARMED:
            _schedule_voice_calibration(voice_id)
        return
    try:
        with conn() as c:
            status_row = c.execute(
                "SELECT synthesis_status,synthesis_message FROM voices WHERE id=?",
                (voice_id,),
            ).fetchone()
    except sqlite3.Error:
        status_row = None
    synthesis_status = (status_row[0] if status_row else "ready") or "ready"
    if synthesis_status == "failed":
        return
    needs_probe = synthesis_status == "pending"
    if not settings.gpt_sovits_url:
        return
    VOICE_WARMING.add(voice_id)
    try:
        with httpx.Client(timeout=180, trust_env=False) as client:
            # New clones use one complete validation request. It both loads the
            # profile and proves that a real phrase is usable, avoiding a
            # redundant short warm-up synthesis before the acceptance probe.
            if needs_probe:
                with _tts_lock(priority="background", timeout=30):
                    _ensure_model_profile_loaded(_voice_model_profile(voice_id), client)
                    probe_params = _live_unit_params(TTSRequest(text=CLONE_VALIDATION_TEXT, voice_id=voice_id), 0)
                    if _voice_prompt_policy(voice_id) == 'checked':
                        probe_audio, check = _checked_voice_probe(client, probe_params, ['clone-probe', _voice_model_profile(voice_id)])
                        if probe_audio is None:
                            raise RuntimeError('逐句检查和回退后仍未通过实际合成验收')
                    else:
                        probe_params.update(streaming_mode=False, parallel_infer=True)
                        probe_timeout = httpx.Timeout(30, connect=10, read=20, write=10, pool=10)
                        probe_response = client.post(_tts_endpoint(), json=probe_params, timeout=probe_timeout)
                        probe_response.raise_for_status()
                        probe_audio = probe_response.content
                    if len(probe_audio) > 1_000_000:
                        raise RuntimeError("实际合成验收输出异常过长")
                    probe = _probe_audio_payload(probe_audio)
                if probe["status"] != "ready":
                    _set_voice_synthesis_status(voice_id, "failed", probe["message"], probe.get("duration"))
                    VOICE_WARMED.discard(voice_id)
                    return
                _set_voice_synthesis_status(voice_id, "ready", probe["message"], probe.get("duration"))
            else:
                request = TTSRequest(text=WARMUP_TEXT, voice_id=voice_id)
                with _tts_lock(priority="background", timeout=30):
                    _ensure_model_profile_loaded(_voice_model_profile(voice_id), client)
                    with client.stream("POST", _tts_endpoint(), json=_tts_params(request, streaming=True)) as response:
                        response.raise_for_status()
                        received = sum(len(chunk) for chunk in response.iter_raw(8192))
                if received <= 128:
                    raise RuntimeError("预热未返回有效音频")
            _mark_tts_success()
            VOICE_WARMED.add(voice_id)
            # Reference screening and seed selection continue in the
            # background after the real synthesis probe succeeds.
            VOICE_WARMING.discard(voice_id)
            _schedule_voice_calibration(voice_id)
    except Exception as exc:
        if needs_probe:
            _set_voice_synthesis_status(voice_id, "failed", f"实际合成验收失败：{str(exc)[:300]}")
        VOICE_WARMED.discard(voice_id)
    finally:
        VOICE_WARMING.discard(voice_id)


def _prime_voice_profile(voice_id: str):
    """Pay the global GPT-SoVITS weight-switch cost before live playback.

    GPT-SoVITS keeps one model pair in process-global state. A clone bound to
    a custom installed profile otherwise makes the first live sentence wait
    for a weight swap. Priming runs the same locked, tiny request used by
    warmup so the subsequent live request starts with the selected profile
    already resident.
    """
    if not settings.gpt_sovits_url or voice_id in VOICE_PROFILE_WARMING:
        return
    profile = _voice_model_profile(voice_id)
    VOICE_PROFILE_WARMING.add(voice_id)
    try:
        _voice_synthesis_gate(voice_id)
        request = TTSRequest(text=CLONE_VALIDATION_TEXT, voice_id=voice_id)
        with _tts_lock(priority="foreground", timeout=20):
            with httpx.Client(timeout=180, trust_env=False) as client:
                # Re-check after acquiring the scheduler. Calibration or a
                # different foreground request may have changed the global
                # profile since the caller inspected it.
                _ensure_model_profile_loaded(profile, client)
                params = _live_unit_params(request, 0)
                response = client.post(_tts_endpoint().removesuffix('/tts') + '/runtime/prime', json=params)
                if response.is_success:
                    _mark_tts_success()
                    VOICE_WARMED.add(voice_id)
                    return
                if response.status_code != 404:
                    response.raise_for_status()
                if voice_id in VOICE_WARMED:
                    return
                with client.stream("POST", _tts_endpoint(), json=params) as response:
                    response.raise_for_status()
                    for _ in response.iter_raw(8192):
                        pass
                VOICE_WARMED.add(voice_id)
    except Exception:
        pass
    finally:
        VOICE_PROFILE_WARMING.discard(voice_id)
        if voice_id in VOICE_WARMED:
            _schedule_voice_calibration(voice_id)


def _voice_calibration_worker():
    """Run background calibration one voice at a time.

    GPT-SoVITS owns one global model and one inference lock. Starting one
    calibration thread per voice made several threads queue behind that lock,
    which could delay a real live request after a page refresh. A small queue
    keeps the work background-only while preserving service responsiveness.
    """
    global VOICE_CALIBRATION_WORKER_RUNNING
    # Warmup owns the same inference lock and schedules jobs while it is still
    # loading voices. Wait until all model/profile priming is complete so a
    # calibration candidate cannot time out simply because startup is active.
    while not TTS_WARMUP.wait(timeout=30):
        logging.getLogger('uvicorn.error').info('voice_calibration worker waiting for live_path_warmed')
    while True:
        with VOICE_CALIBRATION_QUEUE_LOCK:
            if not VOICE_CALIBRATION_QUEUE:
                VOICE_CALIBRATION_WORKER_RUNNING = False
                return
            voice_id = VOICE_CALIBRATION_QUEUE.pop()
            force = voice_id in VOICE_CALIBRATION_FORCED
            VOICE_CALIBRATION_FORCED.discard(voice_id)
        try:
            # Do not even start the local screening subprocess while a live
            # request owns a foreground ticket. This avoids CPU contention and
            # lets the next candidate begin cleanly after playback is idle.
            _wait_for_calibration_idle(grace=True)
            result = _calibrate_voice(voice_id, force=force)
            if result is CALIBRATION_DEFERRED:
                with VOICE_CALIBRATION_QUEUE_LOCK:
                    VOICE_CALIBRATION_QUEUE.add(voice_id)
                    if force:
                        VOICE_CALIBRATION_FORCED.add(voice_id)
                time.sleep(0.5)
        except Exception as exc:
            # Optional calibration must never terminate the worker or affect an
            # already usable clone.
            _calibration_status(voice_id, 'failed', str(exc))
            logging.getLogger('uvicorn.error').exception('voice_calibration worker failed voice=%s', voice_id)


def _wait_for_calibration_idle(*, grace=False):
    deadline = time.monotonic() + settings.gpt_sovits_calibration_idle_seconds if grace else 0
    with TTS_SCHEDULER:
        while (TTS_FOREGROUND_ACTIVE or TTS_FOREGROUND_WAITERS
               or time.monotonic() < max(deadline, _TTS_LAST_FOREGROUND_AT + settings.gpt_sovits_calibration_idle_seconds)):
            TTS_SCHEDULER.wait(timeout=1.0)


def _schedule_voice_calibration(voice_id: str, *, force=False):
    """Queue calibration without making the first usable clone wait."""
    global VOICE_CALIBRATION_WORKER_RUNNING
    if voice_id in PRESET_VOICES or not settings.gpt_sovits_calibration_enabled or voice_id in VOICE_CALIBRATING:
        return
    with conn() as c:
        row = c.execute('SELECT * FROM voices WHERE id=?', (voice_id,)).fetchone()
    if not row or not row['reference_path'] or row['model_profile'] != 'base' or row['synthesis_status'] != 'ready':
        return
    if not force and _calibration_complete(row):
        return
    if voice_adaptation.get_job(voice_id):
        # Per-voice adaptation is a separate acoustic experiment; do not
        # silently follow it with a sampling-parameter search.
        return
    with VOICE_CALIBRATION_QUEUE_LOCK:
        if voice_id in VOICE_CALIBRATION_QUEUE:
            return
        VOICE_CALIBRATION_QUEUE.add(voice_id)
        if force:
            VOICE_CALIBRATION_FORCED.add(voice_id)
        _calibration_status(voice_id, 'queued', '已排队，等待直播预热完成和播报空闲')
        if VOICE_CALIBRATION_WORKER_RUNNING:
            return
        VOICE_CALIBRATION_WORKER_RUNNING = True
    threading.Thread(
        target=_voice_calibration_worker,
        name="calibrate-voice-worker",
        daemon=True,
    ).start()


def _voice_config(voice_id: str):
    if voice_id in PRESET_VOICES:
        voice = PRESET_VOICES[voice_id]
        ref = preset_reference_paths(voice_id)[0]
        if _audio_quality(ref).get('status') != 'ready':
            raise HTTPException(503, f"内置音色“{voice['name']}”的资源缺失或损坏，请恢复 data/preset_voices")
        return ref, voice['prompt_text'], voice['prompt_lang']
    with conn() as c:
        row = c.execute("SELECT * FROM voices WHERE id=?", (voice_id,)).fetchone()
    global_reference = _tts_global_reference_audio()
    ref = row["reference_path"] if row and row["reference_path"] else global_reference
    default_prompt = settings.gpt_sovits_prompt_text
    prompt_text = row["prompt_text"] if row and row["prompt_text"] else default_prompt
    prompt_text = _normalize_prompt_text(prompt_text)
    prompt_lang = row["prompt_lang"] if row and row["prompt_lang"] else (
        settings.gpt_sovits_prompt_language
    )
    # A deployment may keep the global reference setting empty and store the
    # usable reference only in the voice library. In that case direct API calls
    # using a preset/default id should still resolve to a valid cloned speaker.
    preset_reference_unusable = (
        (not row or not row["reference_path"])
        and (
            not prompt_text.strip()
            or _audio_quality(ref).get("status") in {"missing", "invalid", "needs-review"}
        )
    )
    if not ref or preset_reference_unusable:
        with conn() as c:
            fallback = c.execute(
                "SELECT reference_path,prompt_text,prompt_lang FROM voices "
                "WHERE reference_path<>'' AND TRIM(prompt_text)<>'' ORDER BY created_at DESC"
            ).fetchall()
        for candidate in fallback:
            candidate_quality = _audio_quality(candidate[0])
            if candidate_quality.get("status") == "ready":
                ref = candidate[0]
                prompt_text = _normalize_prompt_text(candidate[1])
                prompt_lang = candidate[2] or prompt_lang
                break
    if not ref:
        raise HTTPException(503, "尚未配置参考音频")
    # GPT-SoVITS can infer without prompt text, but supplying the real transcript is
    # the single biggest factor in retaining the uploaded speaker's pronunciation
    # and prosody. Do not silently fall back for an uploaded clone.
    if row and row["reference_path"] and not prompt_text.strip():
        raise HTTPException(422, "该克隆音色尚未填写参考音频原文，请补充与录音完全一致的文本后再播报")
    quality = _audio_quality(ref)
    if quality["status"] in {"invalid", "missing", "needs-review"}:
        raise HTTPException(422, f"参考音频格式无效：{quality['message']}")
    if row and row["reference_path"] and quality["status"] == "needs-review":
        raise HTTPException(422, f"参考音频需要优化后再播报：{quality['message']}")
    return ref, prompt_text, prompt_lang


def _voice_aux_reference_paths(voice_id: str):
    """Return validated auxiliary references for the selected speaker.

    GPT-SoVITS uses the first reference for semantic/prompt conditioning and
    these additional files for multi-reference timbre fusion.  Preset ids are
    resolved to the same fallback clone as ``_voice_config`` so the reference
    and auxiliary set can never come from different speakers.
    """
    if voice_id in PRESET_VOICES:
        return [path for path in preset_reference_paths(voice_id)[1:] if Path(path).is_file()]
    row = None
    screened_column = True
    with conn() as c:
        try:
            row = c.execute(
                "SELECT reference_path,aux_reference_paths,validated_aux_reference_paths FROM voices WHERE id=?",
                (voice_id,),
            ).fetchone()
        except sqlite3.Error:
            screened_column = False
            row = c.execute(
                "SELECT reference_path,aux_reference_paths FROM voices WHERE id=?",
                (voice_id,),
            ).fetchone()
        global_reference = _tts_global_reference_audio()
        if (not row or not row[0]) and global_reference:
            columns = "reference_path,aux_reference_paths,validated_aux_reference_paths" if screened_column else "reference_path,aux_reference_paths"
            row = c.execute(
                f"SELECT {columns} FROM voices WHERE reference_path=? LIMIT 1",
                (global_reference,),
            ).fetchone()
        if not row or not row[0]:
            columns = "reference_path,aux_reference_paths,validated_aux_reference_paths" if screened_column else "reference_path,aux_reference_paths"
            candidates = c.execute(
                f"SELECT {columns} FROM voices "
                "WHERE reference_path<>'' AND TRIM(prompt_text)<>'' ORDER BY created_at DESC"
            ).fetchall()
            for candidate in candidates:
                if _audio_quality(candidate[0]).get("status") == "ready":
                    row = candidate
                    break
    if not row:
        return []
    try:
        # NULL means a legacy/new voice that has not been screened yet. Use
        # only the transcript-bearing primary clip until screening completes;
        # averaging an unverified auxiliary can make the first usable preview
        # less like the speaker. An explicit [] means calibration rejected all
        # auxiliaries and must not silently fall back to averaging them again.
        encoded_paths = row[2] if screened_column and row[2] is not None else "[]" if screened_column else row[1]
        paths = json.loads(encoded_paths or "[]")
    except (TypeError, ValueError):
        paths = []
    if not isinstance(paths, list):
        return []
    return [
        str(path)
        for path in paths[:MAX_AUX_REFERENCE_AUDIO]
        if isinstance(path, str)
        and Path(path).is_file()
        and _audio_quality(path).get("status") == "ready"
    ]


def _voice_sampling_seed(voice_id: str):
    """Return one deterministic seed per speaker, independent of script text."""
    if voice_id in PRESET_VOICES:
        return PRESET_VOICES[voice_id]['sampling']['seed']
    row = None
    try:
        with conn() as c:
            row = c.execute(
                "SELECT sampling_seed,reference_path FROM voices WHERE id=?",
                (voice_id,),
            ).fetchone()
            if not row or not row[1]:
                global_reference = _tts_global_reference_audio()
                if global_reference:
                    row = c.execute(
                        "SELECT sampling_seed,reference_path FROM voices WHERE reference_path=? LIMIT 1",
                        (global_reference,),
                    ).fetchone()
                if not row or not row[1]:
                    candidates = c.execute(
                        "SELECT sampling_seed,reference_path FROM voices "
                        "WHERE reference_path<>'' AND TRIM(prompt_text)<>'' ORDER BY created_at DESC"
                    ).fetchall()
                    for candidate in candidates:
                        if _audio_quality(candidate[1]).get("status") == "ready":
                            row = candidate
                            break
    except sqlite3.Error:
        # Unit callers and older installations may reach inference before the
        # additive migration has run. The normal lifespan migrates before HTTP.
        row = None
    if row and row[0] is not None:
        return int(row[0]) & 0x7FFFFFFF
    identity = row[1] if row and row[1] else voice_id
    return zlib.crc32(f"voice-seed\0{identity}".encode("utf-8")) & 0x7FFFFFFF


def _voice_sampling_profile(voice_id: str):
    """Load fixed preset parameters or the uploaded speaker's calibration."""
    if voice_id in PRESET_VOICES:
        return {**PRESET_VOICES[voice_id]['sampling'], 'calibrated': True}
    try:
        with conn() as c:
            row = c.execute(
                "SELECT sampling_top_k,sampling_top_p,sampling_temperature,sampling_model_version,calibration_details FROM voices WHERE id=?",
                (voice_id,),
            ).fetchone()
    except sqlite3.Error:
        row = None
    try:
        details = json.loads(row[4] or "{}") if row else {}
    except (TypeError, ValueError):
        details = {}
    if row and row[1] is not None and row[2] is not None and row[3] == settings.gpt_sovits_model_version and details.get("calibration_version") == CALIBRATION_VERSION:
        return {
            # Rows created before top-k calibration used the official default
            # 15. Keep those records valid while new clones persist the
            # selected value explicitly.
            "top_k": max(1, min(30, int(row[0] or 15))),
            "top_p": max(0.1, min(1.0, float(row[1]))),
            "temperature": max(0.1, min(2.0, float(row[2]))),
            "repetition_penalty": DEFAULT_SAMPLING_PROFILE["repetition_penalty"],
            "calibrated": True,
        }
    return {**DEFAULT_SAMPLING_PROFILE, "calibrated": False}


def _save_voice_calibration(
    voice_id: str,
    seed: int,
    accepted_aux_paths: list[str],
    *,
    top_k: int = 15,
    top_p: float = DEFAULT_SAMPLING_PROFILE["top_p"],
    temperature: float = DEFAULT_SAMPLING_PROFILE["temperature"],
    details: dict | None = None,
):
    with conn() as c:
        c.execute(
            "UPDATE voices SET sampling_seed=?,validated_aux_reference_paths=?,sampling_top_k=?,sampling_top_p=?,"
            "sampling_temperature=?,sampling_model_version=?,calibration_details=? WHERE id=?",
            (
                int(seed) & 0x7FFFFFFF,
                json.dumps(accepted_aux_paths, ensure_ascii=False),
                max(1, min(30, int(top_k))),
                max(0.1, min(1.0, float(top_p))),
                max(0.1, min(2.0, float(temperature))),
                settings.gpt_sovits_model_version,
                json.dumps({"calibration_version": CALIBRATION_VERSION, **(details or {})}, ensure_ascii=False),
                voice_id,
            ),
        )


def _voice_model_profile(voice_id: str):
    """Resolve the model family for a voice without changing inference state."""
    if voice_id in PRESET_VOICES:
        return PRESET_VOICES[voice_id]['model_profile']
    profiles = model_profiles(settings)
    with conn() as c:
        row = c.execute("SELECT model_profile,reference_path FROM voices WHERE id=?", (voice_id,)).fetchone()
    # Preset voices and newly uploaded samples use the v2ProPlus base model.  A
    # fine-tuned model is opt-in per database row so it cannot leak into other
    # speakers through GPT-SoVITS' process-global weight state.
    if not row or not row[1]:
        global_reference = _tts_global_reference_audio()
        if global_reference:
            with conn() as c:
                global_row = c.execute(
                    "SELECT model_profile FROM voices WHERE reference_path=? LIMIT 1",
                    (global_reference,),
                ).fetchone()
            if global_row and global_row[0] in profiles:
                return global_row[0]
        # Preset ids intentionally reuse the newest validated clone when no
        # global reference is configured. Resolve that same fallback here so
        # its reference and model family cannot get out of sync.
        with conn() as c:
            candidates = c.execute(
                "SELECT model_profile,reference_path FROM voices "
                "WHERE reference_path<>'' AND TRIM(prompt_text)<>'' ORDER BY created_at DESC"
            ).fetchall()
        for candidate in candidates:
            if _audio_quality(candidate[1]).get("status") == "ready":
                return candidate[0] if candidate[0] in profiles else "base"
        return "base"
    if row[0] and row[0] not in profiles:
        raise HTTPException(503, '该音色的模型配置缺失，请恢复已安装的音色模型')
    return row[0] or 'base'


def _model_profile_weights(profile: str):
    return profile_weights(settings, profile)


def _ensure_model_profile_loaded(profile: str, client: httpx.Client):
    """Load and verify a model pair while TTS_LOCK is held.

    The GPT-SoVITS process owns the actual global weights.  The backend's
    cached profile is only a hint: the upstream process can survive a backend
    restart, an external model switch, or a failed request.  Always let the
    transport compare ``/runtime/status`` before inference so a stale cached
    profile can never authorize synthesis with another speaker's weights.
    """
    global _ACTIVE_MODEL_PROFILE
    if not settings.gpt_sovits_url:
        return
    try:
        gpt_sovits_engine.load_weights(settings, profile, client)
    except (httpx.HTTPError, RuntimeError) as exc:
        raise HTTPException(503, f"切换 GPT-SoVITS {profile} 权重失败：{exc}") from exc
    _ACTIVE_MODEL_PROFILE = profile
    _mark_tts_success()


async def _ensure_model_profile_loaded_async(profile: str, client: httpx.AsyncClient):
    """Async counterpart used by the live streaming generator.

    Do not short-circuit from ``_ACTIVE_MODEL_PROFILE``; the service status is
    the authority for the process-global model state.
    """
    global _ACTIVE_MODEL_PROFILE
    if not settings.gpt_sovits_url:
        return
    try:
        await gpt_sovits_engine.load_weights_async(settings, profile, client)
    except httpx.HTTPError as exc:
        raise RuntimeError(f"切换 GPT-SoVITS {profile} 权重失败：{exc}") from exc
    except RuntimeError:
        raise
    _ACTIVE_MODEL_PROFILE = profile
    _mark_tts_success()


def _preferred_clone_voice_id():
    """Return the newest validated clone for warmup and acceptance checks."""
    with conn() as c:
        rows = c.execute(
            "SELECT id,reference_path,prompt_text FROM voices "
            "WHERE reference_path<>'' AND TRIM(prompt_text)<>'' ORDER BY created_at DESC"
        ).fetchall()
    for row in rows:
        if _audio_quality(row[1]).get("status") == "ready":
            return row[0]
    return "steady"


CN_DIGITS = "零一二三四五六七八九"


def _cn_section(number: int, *, omit_one_ten: bool = True) -> str:
    """Convert one 0..9999 section without duplicate or trailing 零."""
    if number == 0:
        return ""
    units = ("", "十", "百", "千")
    digits = list(reversed(f"{number:04d}"))
    out, pending_zero = [], False
    for index in range(3, -1, -1):
        digit = int(digits[index])
        if digit:
            if pending_zero and out:
                out.append("零")
            out.append(CN_DIGITS[digit] + units[index])
            pending_zero = False
        elif out:
            pending_zero = True
    result = "".join(out)
    # Only standalone 10–19 omit the leading 一; 110 must remain 一百一十.
    return result.replace("一十", "十", 1) if omit_one_ten and number < 20 else result


def _number_to_cn(value: str) -> str:
    """Use natural Chinese cardinal numbers; decimals read digit-by-digit."""
    if "." in value:
        left, right = value.split(".", 1)
        return _number_to_cn(left) + "点" + "".join(CN_DIGITS[int(x)] for x in right)
    number = int(value)
    if number == 0:
        return "零"
    sections = []
    while number:
        sections.append(number % 10000)
        number //= 10000
    result, need_zero = [], False
    big_units = ("", "万", "亿", "兆")
    for index in range(len(sections) - 1, -1, -1):
        section = sections[index]
        if not section:
            if result:
                need_zero = True
            continue
        if result and (need_zero or section < 1000):
            if result[-1] != "零":
                result.append("零")
        result.append(
            _cn_section(section, omit_one_ten=not result)
            + (big_units[index] if index < len(big_units) else "")
        )
        need_zero = False
    return "".join(result).replace("零零", "零").rstrip("零")


def normalize_tts_text(text: str) -> str:
    """Prepare speakable text; also accept already-normalized browser phrases."""
    # Citation numbers belong to the evidence display, never the spoken copy.
    # Strip before number conversion, including grouped citations such as [1,2].
    citation = r"\s*\d+(?:\s*[,，、\-–—]\s*\d+)*\s*"
    text = re.sub(rf"\[{citation}\]|【{citation}】|［{citation}］", "", text)
    text = re.sub(r"^根据当前车型资料[:：]\s*", "", text.strip())
    text = re.sub(r"^根据资料[:：]\s*", "", text)

    # Chinese percentages put the unit BEFORE the value. Handle these before
    # the generic number+unit conversion so 30% never becomes 三十百分之.
    text = re.sub(r"(\d+(?:\.\d+)?)[ \t]*[%％]",
                  lambda m: "百分之" + _number_to_cn(m.group(1)), text)

    # Dates first, otherwise the year/month/day would be converted independently.
    def date_replace(match):
        year, month, day = match.groups()
        return "".join(CN_DIGITS[int(x)] for x in year) + f"年{_number_to_cn(month)}月{_number_to_cn(day)}日"
    text = re.sub(r"(?<!\d)(19\d{2}|20\d{2})[-/]([01]?\d)[-/]([0-3]?\d)(?!\d)", date_replace, text)

    # Model years are conventionally read digit-by-digit: 二零二六款.
    text = re.sub(
        r"(?<!\d)(19\d{2}|20\d{2})(?=\s*(?:年|款|版|型号))",
        lambda m: "".join(CN_DIGITS[int(x)] for x in m.group(1)),
        text,
    )

    units = {
        "km/h": "公里每小时", "km": "公里", "kWh": "千瓦时", "kW": "千瓦", "N·m": "牛米", "Nm": "牛米",
        "L": "升", "万元": "万元", "万": "万元", "元": "元",
        "公里": "公里", "毫米": "毫米", "小时": "小时", "秒": "秒",
    }
    # Longest units first prevents 13.38万元 becoming “万元元”.
    unit_pattern = r"万元|km/h|kWh|N·m|公里|毫米|小时|秒|km|kW|Nm|L|万|元"
    pattern = re.compile(rf"(?<![\d点])(\d+(?:\.\d+)?)[ \t]*({unit_pattern})?", re.I)

    def replace(match):
        number, unit = match.group(1), match.group(2) or ""
        return _number_to_cn(number) + units.get(unit, unit)

    text = pattern.sub(replace, text)
    # Spaces in the live script are formatting, not speech pauses. Keeping
    # them between Chinese, digits and model abbreviations makes the upstream
    # tokenizer insert audible breaks such as “五 EV 二零二六”.
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[:：]+", "，", text)
    text = re.sub(r"，{2,}", "，", text)
    return text.strip()


# Keep quantities and their units together when a long live sentence is split.
# Splitting `580km` into `580` + `km` makes the upstream model read the value
# as disconnected digits instead of the natural Chinese quantity.
_TTS_ATOMIC_TOKEN = re.compile(
    r"(?:[A-Za-z]+(?:/[A-Za-z]+)?|百分之[零一二三四五六七八九十百千万亿兆点\d.]+"
    r"|\d+(?:\.\d+)?\s*(?:km/h|kWh|N·m|km|kW|Nm|公里|毫米|小时|秒|%|L|万元|万|元|年|款|版|型号|EV)?"
    r"|[零一二三四五六七八九十百千万亿兆点]+(?:公里每小时|千瓦时|牛米|公里|毫米|小时|秒|百分之|升|万元|元|年|款|版|型号|EV)?)",
    re.IGNORECASE,
)


def _safe_tts_cut(text: str, limit: int) -> int:
    """Return a split point that does not cut a number/unit token in half."""
    if len(text) <= limit:
        return len(text)
    end = max(1, limit)
    for match in _TTS_ATOMIC_TOKEN.finditer(text):
        if match.start() < end < match.end():
            # If the token starts the chunk, keep the whole token even when it
            # is slightly longer than the nominal chunk size.
            end = match.end() if match.start() == 0 else match.start()
            break
    return max(1, end)


def _voice_prompt_policy(voice_id: str):
    if voice_id in PRESET_VOICES:
        return 'off'
    try:
        with conn() as c:
            row=c.execute('SELECT prompt_policy FROM voices WHERE id=?',(voice_id,)).fetchone()
    except sqlite3.Error:
        return 'off'
    return row[0] if row else 'off'


def _tts_params(req: TTSRequest, *, streaming: bool = False, seed_override: int | None = None):
    ref, prompt_text, prompt_lang = _voice_config(req.voice_id)
    aux_refs = _voice_aux_reference_paths(req.voice_id)
    model_profile = _voice_model_profile(req.voice_id)
    model_config = model_profiles(settings)[model_profile]
    speaker_reference = model_config.get('speaker_ref_audio_path')
    live_model_mode = settings.gpt_sovits_streaming_mode if streaming else False
    builtin = req.voice_id in PRESET_VOICES
    profile = {}
    delivery = 'natural' if builtin else req.delivery
    if streaming and not builtin and req.delivery == 'natural':
        delivery = settings.gpt_sovits_live_delivery
    delivery_profiles = {'natural': (1.0,0.0), 'warm': (0.98,0.025), 'lively': (1.025,0.05), 'calm': (0.94,-0.025)}
    delivery_speed,delivery_temperature=delivery_profiles[delivery]
    profile['speed']=profile.get('speed',1.0)*delivery_speed
    profile['temperature']=profile.get('temperature',0.0)+delivery_temperature
    speed_factor = max(0.6, min(1.6, req.speed_factor * profile.get("speed", 1.0)))
    sampling = _voice_sampling_profile(req.voice_id)
    requested_top_k = sampling.get("top_k", 15) if builtin or req.top_k is None else req.top_k
    requested_top_p = sampling["top_p"] if builtin or req.top_p is None else req.top_p
    requested_temperature = sampling["temperature"] if builtin or req.temperature is None else req.temperature
    # Validated sampling must survive live playback unchanged. Floors only
    # supply an uncalibrated fallback; explicit overrides still win.
    if streaming and not builtin and not sampling.get('calibrated'):
        if req.top_k is None:
            requested_top_k = max(int(requested_top_k), settings.gpt_sovits_clone_live_top_k)
        if req.top_p is None:
            requested_top_p = max(float(requested_top_p), settings.gpt_sovits_clone_live_top_p)
        if req.temperature is None:
            requested_temperature = max(float(requested_temperature), settings.gpt_sovits_clone_live_temperature)
    requested_repetition = sampling["repetition_penalty"] if builtin or req.repetition_penalty is None else req.repetition_penalty
    top_p = max(0.1, min(1.0, float(requested_top_p)))
    temperature = max(0.1, min(2.0, float(requested_temperature) + profile.get("temperature", 0.0)))
    repetition_penalty = max(0.5, min(3.0, float(requested_repetition) + profile.get("repetition", 0.0)))
    normalized_text = normalize_tts_text(req.text)
    if (streaming and not builtin and not settings.gpt_sovits_live_use_prompt_text and _voice_prompt_policy(req.voice_id) != 'checked'
            and settings.gpt_sovits_model_version in {'v2', 'v2Pro', 'v2ProPlus'}
            and _voice_model_profile(req.voice_id) == 'base'):
        # Reference-free semantics can help diagnose a mismatched transcript,
        # at the cost of losing the reference's speaking style.
        prompt_text = ""
    if seed_override is not None:
        seed = int(seed_override) & 0x7FFFFFFF
    else:
        seed = _voice_sampling_seed(req.voice_id)
        if streaming and not builtin and settings.gpt_sovits_live_seed_mode == 'text-low8':
            seed ^= zlib.crc32(normalized_text.encode('utf-8')) & 255
    return {
        **({'speaker_ref_audio_path': speaker_reference} if speaker_reference else {}),
        **({'speaker_aux_ref_audio_paths': model_config['speaker_aux_ref_audio_paths'],
            'speaker_similarity_threshold': settings.gpt_sovits_aux_similarity_threshold}
           if 'speaker_aux_ref_audio_paths' in model_config else {}),
        "ref_audio_path": ref,
        "aux_ref_audio_paths": aux_refs,
        "prompt_text": prompt_text,
        "prompt_lang": prompt_lang,
        "text": normalized_text,
        "text_lang": (
            settings.gpt_sovits_live_text_language
            if streaming else settings.gpt_sovits_text_language
        ),
        # The browser already feeds short, punctuation-aligned live units. Keeping
        # each unit intact avoids a second server-side split and makes revisions
        # take effect at the next natural pause.
        "text_split_method": (
            settings.gpt_sovits_live_text_split_method if streaming
            else "cut2"
        ),
        "media_type": "wav",
        # Mode 1 returns clean phrase-sized PCM fragments using the quality-
        # first decoder, which keeps Chinese articulation stable.
        "streaming_mode": live_model_mode,
        "top_k": max(1, min(30, int(requested_top_k))),
        "top_p": top_p,
        "temperature": temperature,
        "seed": seed,
        "speed_factor": speed_factor,
        # Mode 1 uses the model's quality-first fragment path.
        # Mode 1 is fragment-return mode, not the naive streaming decoder, so
        # it can safely use the faster parallel path. Modes 2/3 require the
        # incremental decoder and must remain serial.
        "parallel_infer": False if live_model_mode in (2, 3) else True,
        "repetition_penalty": repetition_penalty,
        # Used by the upstream v3/v4 vocoder branch; v2ProPlus ignores this.
        "sample_steps": 32,
        # Both live fragments and quality preview audio already contain natural
        # punctuation pauses. Appending GPT-SoVITS' synthetic zero tail creates
        # an avoidable gap at every unit boundary.
        "fragment_interval": 0.0,
        "overlap_length": 2,
        # Keep the model's default semantic block size for stable prosody.
        "min_chunk_length": (
            max(8, min(32, int(settings.gpt_sovits_live_min_chunk_length)))
            if streaming else 16
        ),
    }


def _live_unit_params(req: TTSRequest, unit_index: int):
    """Use the same progressive decoder for playback, priming and benchmarks."""
    delivery = 'natural' if req.voice_id in PRESET_VOICES else req.delivery
    if req.voice_id not in PRESET_VOICES and req.delivery == 'natural':
        delivery = settings.gpt_sovits_live_delivery
    text,speed=delivery_text(req.text,delivery,req.speed_factor)
    params = _tts_params(req.model_copy(update={'text':text,'speed_factor':speed}), streaming=True)
    params['streaming_mode'] = settings.gpt_sovits_live_streaming_mode
    params['parallel_infer'] = params['streaming_mode'] not in (2, 3)
    return params


def _effective_live_settings(voice_id: str, params: dict | None = None):
    params = params or _live_unit_params(TTSRequest(voice_id=voice_id, text=CLONE_VALIDATION_TEXT), 0)
    return {**{key: params[key] for key in ('top_k','top_p','temperature','seed','speed_factor','streaming_mode')},
            'use_prompt_text': bool(params['prompt_text']),
            'model_profile': _voice_model_profile(voice_id),
            'prompt_policy': _voice_prompt_policy(voice_id),
            'speaker_anchor': ('scored-fusion' if params.get('speaker_aux_ref_audio_paths') is not None
                               else 'original-reference') if params.get('speaker_ref_audio_path') else 'reference-list',
            'seed_mode': 'fixed' if voice_id in PRESET_VOICES else settings.gpt_sovits_live_seed_mode}


def _tts_text_units(text: str, max_chars: int = 40):
    """Split live copy at natural pauses while keeping tiny lead-ins attached.

    A standalone two- or three-character opener is an especially weak cloning
    prompt. It tends to produce a clipped initial consonant and adds an extra
    request boundary, so short clauses are merged with the following clause
    within a bounded natural phrase.
    """
    short_clause_chars = 10
    # Keep a short salutation attached to context without creating a long
    # first request that delays the first PCM fragment beyond three seconds.
    clause_merge_max = max_chars + 8
    min_tail_chars = 8
    normalized = normalize_tts_text(text)
    raw = re.findall(r"[^。！？；.!?]+[。！？；.!?]?", normalized)
    units = []
    for sentence in raw or [normalized]:
        sentence = sentence.strip()
        if not sentence:
            continue
        clauses = re.findall(r"[^，,、：:]+[，,、：:]?", sentence)
        merged_clauses = []
        index = 0
        while index < len(clauses):
            clause = clauses[index].strip()
            following = clauses[index + 1].strip() if index + 1 < len(clauses) else ""
            # Preserve punctuation-aligned phrases whenever they fit.  The
            # former two-sided >=10 check could never fit its 16-character
            # cap, leaving every comma clause as a separate request.
            can_merge = following and len(clause) + len(following) <= clause_merge_max
            if can_merge:
                combined = clause + following
                index += 2
                while index < len(clauses) and len(combined) + len(clauses[index].strip()) <= clause_merge_max:
                    combined += clauses[index].strip()
                    index += 1
                merged_clauses.append((combined, True))
                continue
            # If a salutation is followed by a long clause, still attach a
            # bounded phonetic context to it. A standalone "老板" is prone to
            # clipped initials, while sending the whole long clause defeats
            # the first-audio latency target.
            if following and len(clause) < short_clause_chars:
                room = clause_merge_max - len(clause) - 1  # reserve the pause comma
                prefix_end = _safe_tts_cut(following, max(1, room))
                prefix = following[:prefix_end].rstrip("，,、：:")
                if prefix:
                    merged_clauses.append((clause + prefix + "，", True))
                    remainder = following[prefix_end:].lstrip()
                    clauses[index + 1] = remainder
                    index += 1
                    continue
            merged_clauses.append((clause, False))
            index += 1
        if len(sentence) <= max_chars or len(clauses) <= 1:
            while len(sentence) > max_chars:
                cut = _safe_tts_cut(sentence, max_chars)
                # Never leave a tiny tail such as the final character of a
                # model name. Keep the complete phrase together instead.
                if len(sentence) - cut < min_tail_chars:
                    units.append(sentence)
                    sentence = ""
                    break
                units.append(sentence[:cut].rstrip("，,、：:") + "，")
                sentence = sentence[cut:].lstrip()
            if sentence:
                units.append(sentence)
            continue
        current = ""
        for clause, relaxed in merged_clauses:
            clause_limit = clause_merge_max if relaxed else max_chars
            while len(clause) > clause_limit:
                if current:
                    units.append(current)
                    current = ""
                cut = _safe_tts_cut(clause, max_chars)
                if len(clause) - cut < min_tail_chars:
                    current = clause
                    clause = ""
                    break
                units.append(clause[:cut].rstrip("，,、：:") + "，")
                clause = clause[cut:].lstrip()
            if current and len(current) + len(clause) > max_chars:
                units.append(current)
                current = ""
            current += clause
        if current:
            units.append(current)
    return units or ([normalized] if normalized else [])


def _tts_stream_units(req: TTSRequest):
    """Resolve stream units once, without splitting an already queued phrase.

    The browser owns live queue boundaries because it also uses them for
    progress, playback scheduling, and safe dynamic revisions. Direct API
    callers can omit ``unitized`` and retain the server-side fallback splitter.
    """
    if req.unitized and req.stream_batch:
        normalized = normalize_tts_text(req.text)
        # The browser sends a bounded batch (at most a few natural phrases).
        # Do not split it again here: a second split would recreate the exact
        # per-phrase HTTP lifecycle that causes prosody discontinuities.
        return [normalized] if normalized else []
    if req.unitized:
        normalized = normalize_tts_text(req.text)
        return _tts_text_units(
            normalized,
            max_chars=max(24, min(96, int(settings.gpt_sovits_live_max_chars))),
        )
    return _tts_text_units(req.text)


def _wav_data_offset(buffer: bytes):
    """Return the PCM data offset once a complete WAV header is buffered."""
    if len(buffer) < 4:
        return None
    if buffer[:4] != b"RIFF":
        raise ValueError("上游返回的不是 WAV 音频")
    if len(buffer) < 12:
        return None
    if buffer[8:12] != b"WAVE":
        raise ValueError("上游返回的不是 WAV 音频")
    offset = 12
    while offset + 8 <= len(buffer):
        chunk_size = struct.unpack_from("<I", buffer, offset + 4)[0]
        body = offset + 8
        if body + chunk_size > len(buffer):
            return None
        if buffer[offset:offset + 4] == b"data":
            return body
        offset = body + chunk_size + (chunk_size % 2)
    return None


def _iter_pcm_wav_payload(chunks, *, include_header: bool):
    """Yield PCM bytes from one complete/streaming WAV response.

    GPT-SoVITS returns a WAV container for every text unit. The live endpoint
    concatenates those units into one browser stream, so only the first unit
    may contribute its RIFF header. This helper also tolerates headers split
    across arbitrary HTTP chunks.
    """
    header_buffer = b""
    data_started = False
    saw_audio = False
    for chunk in chunks:
        if not chunk:
            continue
        if data_started:
            saw_audio = True
            yield chunk
            continue
        header_buffer += chunk
        data_offset = _wav_data_offset(header_buffer)
        if data_offset is None:
            continue
        if include_header:
            yield header_buffer[:data_offset]
        payload = header_buffer[data_offset:]
        if payload:
            saw_audio = True
            yield payload
        data_started = True
    if not data_started:
        raise RuntimeError("GPT-SoVITS 未返回完整 WAV 音频")
    if not saw_audio:
        raise RuntimeError("GPT-SoVITS 未返回 PCM 音频数据")


def _tts_endpoint():
    try:
        return gpt_sovits_engine.endpoint(settings)
    except RuntimeError as exc:
        raise HTTPException(503, str(exc)) from exc


def _ensure_tts_ready():
    if not _tts_has_reference():
        raise HTTPException(503, "GPT-SoVITS 尚未配置可用参考音频")
    if not _gpt_sovits_reachable():
        raise HTTPException(503, "GPT-SoVITS 服务未启动或不可访问")


INSUFFICIENT_ANSWER = "知识库中暂时没有足够依据回答这个问题。"

# Keep question aliases separate from source labels. Broad aliases such as
# “功率” previously matched the wrong line (150kW instead of 204Ps), while
# many configuration fields were not covered at all.
FIELD_SPECS = [
    {"name": "对外放电功率", "aliases": ("对外放电功率", "外放电功率"), "source": ("对外交流放电功率",), "unit": "kW"},
    {"name": "快充电量范围", "aliases": ("快充电量范围",), "source": ("电池快充电量范围",), "unit": "%"},
    {"name": "电能当量燃料消耗", "aliases": ("电能当量燃料消耗", "电能当量油耗"), "source": ("电能当量燃料消耗量",), "unit": "L/100km"},
    {"name": "前电机型号", "aliases": ("前电机型号", "前电动机型号"), "source": ("前电动机型号",), "unit": ""},
    {"name": "前电机品牌", "aliases": ("前电机品牌", "前电动机品牌"), "source": ("前电动机品牌",), "unit": ""},
    {"name": "慢充接口位置", "aliases": ("慢充接口位置",), "source": ("慢充接口位置",), "unit": ""},
    {"name": "单电机布局", "aliases": ("单电机布局", "电机布局"), "source": ("电机布局",), "unit": ""},
    {"name": "电池冷却方式", "aliases": ("电池冷却方式",), "source": ("电池冷却方式",), "unit": ""},
    {"name": "充电功能", "aliases": ("充电功能", "快充功能"), "source": ("快充功能",), "unit": ""},
    {"name": "车门开启方式", "aliases": ("车门开启方式",), "source": ("车门开启方式",), "unit": ""},
    {"name": "后排侧气囊", "aliases": ("后排侧气囊", "侧气囊"), "source": ("前/后排侧气囊",), "unit": ""},
    {"name": "车道偏离预警", "aliases": ("车道偏离预警",), "source": ("车道偏离预警系统",), "unit": ""},
    {"name": "导航系统", "aliases": ("导航系统", "卫星导航"), "source": ("卫星导航系统",), "unit": ""},
    {"name": "蓝牙车载电话", "aliases": ("蓝牙车载电话", "蓝牙电话", "车载电话"), "source": ("蓝牙/车载电话",), "unit": ""},
    {"name": "自动泊车", "aliases": ("自动泊车", "辅助泊车"), "source": ("辅助泊车入位", "遥控泊车"), "unit": ""},
    {"name": "语音连续识别", "aliases": ("语音连续识别",), "source": ("语音连续识别",), "unit": ""},
    {"name": "辅助驾驶灯", "aliases": ("辅助驾驶灯",), "source": ("辅助驾驶灯",), "unit": ""},
    {"name": "能量回收系统", "aliases": ("能量回收系统",), "source": ("能量回收系统",), "unit": ""},
    {"name": "全液晶仪表盘", "aliases": ("全液晶仪表盘",), "source": ("全液晶仪表盘",), "unit": ""},
    {"name": "多功能方向盘", "aliases": ("多功能方向盘",), "source": ("多功能方向盘",), "unit": ""},
    {"name": "OTA升级", "aliases": ("OTA升级", "OTA"), "source": ("OTA升级",), "unit": ""},
    {"name": "Wi-Fi热点", "aliases": ("Wi-Fi热点", "wifi热点", "热点"), "source": ("Wi-Fi热点",), "unit": ""},
    {"name": "驾驶模式", "aliases": ("驾驶模式",), "source": ("驾驶模式",), "unit": "", "pattern": r"驾驶模式[^。；;]*?(经济、标准、运动)"},
    {"name": "快充时间", "aliases": ("快充时间",), "source": ("电池快充时间",), "unit": "小时"},
    {"name": "慢充时间", "aliases": ("慢充时间",), "source": ("电池慢充时间",), "unit": "小时"},
    {"name": "电动机功率", "aliases": ("电动机功率", "电动机马力"), "source": ("电动机(Ps)", "电动机总马力"), "unit": "Ps"},
    {"name": "最大功率", "aliases": ("最大功率", "总功率"), "source": ("最大功率", "电动机总功率", "前电动机最大功率"), "unit": "kW"},
    {"name": "最大扭矩", "aliases": ("最大扭矩", "总扭矩"), "source": ("最大扭矩", "电动机总扭矩"), "unit": "N·m"},
    {"name": "厂商指导价", "aliases": ("厂商指导价", "指导价", "价格", "售价", "多少钱"), "source": ("厂商指导价",), "unit": "万元"},
    {"name": "电池容量", "aliases": ("电池容量", "电池能量"), "source": ("电池能量",), "unit": "kWh"},
    {"name": "后备厢", "aliases": ("后备厢容积", "后备箱容积", "后备厢", "后备箱"), "source": ("后备厢容积",), "unit": "L"},
    {"name": "续航", "aliases": ("CLTC纯电续航", "纯电续航", "续航"), "source": ("CLTC纯电续航里程",), "unit": "公里"},
    {"name": "轴距", "aliases": ("轴距",), "source": ("轴距",), "unit": "毫米"},
    {"name": "最高车速", "aliases": ("最高车速", "最高速度"), "source": ("最高车速",), "unit": "km/h"},
    {"name": "整备质量", "aliases": ("整备质量",), "source": ("整备质量",), "unit": "kg"},
    {"name": "最大满载质量", "aliases": ("最大满载质量", "满载质量"), "source": ("最大满载质量",), "unit": "kg"},
    {"name": "车身长度", "aliases": ("车身长度", "车长", "长度"), "source": ("长度",), "unit": "mm"},
    {"name": "车身宽度", "aliases": ("车身宽度", "车宽", "宽度"), "source": ("宽度",), "unit": "mm"},
    {"name": "车身高度", "aliases": ("车身高度", "车高", "高度"), "source": ("高度",), "unit": "mm"},
    {"name": "前轮距", "aliases": ("前轮距",), "source": ("前轮距",), "unit": "mm"},
    {"name": "后轮距", "aliases": ("后轮距",), "source": ("后轮距",), "unit": "mm"},
]


def _pick_source(sources, keywords):
    for source in sources:
        content = source.get("content", "")
        if any(keyword in content for keyword in keywords):
            return source
    return sources[0] if sources else None


def _extract_line(text, keywords):
    matches = []
    for raw in re.split(r"[\n\r]+", text):
        line = raw.strip()
        if not line:
            continue
        if any(keyword in line for keyword in keywords):
            value_match = re.search(r"[:：]\s*([^\n\r。；;]+)", line)
            value = value_match.group(1).strip() if value_match else ""
            # Prefer a complete numeric value over a chunk-truncated prefix.
            numeric_score = len(re.findall(r"\d", value))
            exact_score = max((len(keyword) for keyword in keywords if keyword in line), default=0)
            matches.append((numeric_score, exact_score, len(value), line))
    return max(matches, default=(0, 0, 0, ""))[-1]


def _field_spec(question):
    normalized = re.sub(r"[\s_?？：:，,。]+", "", normalize_question(question)).lower()
    return next(
        (spec for spec in sorted(FIELD_SPECS, key=lambda item: max(map(len, item["aliases"])), reverse=True)
         if any(alias.lower() in normalized for alias in spec["aliases"])),
        None,
    )


def _field_value(text, spec):
    candidates = []
    for raw in re.split(r"[\n\r]+", text):
        line = raw.strip()
        if not line:
            continue
        pattern = spec.get("pattern")
        if pattern:
            match = re.search(pattern, line)
            if match:
                candidates.append((len(match.group(0)), 0, len(match.group(1)), match.group(1)))
        for keyword in sorted(spec["source"], key=len, reverse=True):
            if keyword not in line:
                continue
            value_match = re.search(r"[:：]\s*([^\n\r。；;]+)", line[line.find(keyword) + len(keyword):])
            if not value_match:
                continue
            value = value_match.group(1).strip().rstrip("，,、")
            if value:
                candidates.append((len(keyword), len(re.findall(r"\d", value)), len(value), value))
            break
    return max(candidates, default=(0, 0, 0, ""))[-1]


def _short_answer(question, sources):
    if re.search(r'优惠|补贴|免息|贷款|分期|对比|区别',question) and sources:
        excerpt=excerpt_answer(question,sources)
        if excerpt:return excerpt
    spec = _field_spec(question)
    if not sources:
        return INSUFFICIENT_ANSWER
    if not spec:
        return excerpt_answer(question,sources) or INSUFFICIENT_ANSWER
    combined = "\n".join(s.get("content", "") for s in sources)
    value = _field_value(combined, spec)
    if not value:
        return excerpt_answer(question,sources) or INSUFFICIENT_ANSWER
    unit = spec["unit"]
    suffix = "" if not unit or unit in value or (unit == "万元" and "万" in value) else unit
    label = 'CLTC纯电续航' if spec['name'] == '续航' else spec['name']
    return f"这款车的{label}是{value}{suffix}。"


def _llm_endpoint():
    return llm_gateway.endpoint()


def _llm_answer(question, sources):
    return llm_gateway.generate(question,sources)


def _spoken_answer_text(text):
    """Keep fallback and model answers natural when shown or sent to TTS."""
    text = re.sub(r'\[\d+\]', '', text or '')
    text = re.sub(r'^\s*根据当前车型资料[：:]\s*', '', text)
    text = re.sub(r'^\s*资料摘录（[^）]*）[：:]\s*', '目前能确认的是：', text)
    text = re.sub(r'^\s*相关资料对比依据[：:]\s*', '对比来看：', text)
    text = text.replace('销售资料记载的该动力版本指导价为', '这款车该动力版本的指导价是')
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


@app.post("/api/tts/synthesize")
def synthesize_tts(req: TTSRequest):
    _ensure_tts_ready()
    _voice_synthesis_gate(req.voice_id)
    try:
        with _foreground_ticket():
            with _tts_lock(priority="foreground"):
                with httpx.Client(timeout=120, trust_env=False) as client:
                    _ensure_model_profile_loaded(_voice_model_profile(req.voice_id), client)
                    response = gpt_sovits_engine.synthesize(settings, client, _tts_params(req))
                    _mark_tts_success()
    except httpx.HTTPError as exc:
        raise HTTPException(502, f"GPT-SoVITS 服务调用失败：{exc}") from exc
    return Response(content=response.content, media_type=response.headers.get("content-type", "audio/wav"))


@asynccontextmanager
async def _live_tts_client():
    client = getattr(app.state, 'tts_client', None)
    if client is not None:
        yield client
    else:
        # Direct unit tests do not run the application's lifespan.
        async with httpx.AsyncClient(timeout=120, trust_env=False) as fallback:
            yield fallback


async def _guarded_prompt_stream(req: TTSRequest, request: Request, units: list[str]):
    # Check full sentences, retaining commas and the browser's first phrase.
    # Every sentence is fully buffered before playback; future PCM is never
    # described as validated before it has actually been synthesized.
    sentences = [part for unit in units for part in re.findall(r'[^。！？!?；;]+[。！？!?；;]?', unit) if part.strip()]
    if not sentences:
        raise HTTPException(422, '播报文本为空')
    request_id=uuid4().hex
    checks=[]
    profile=_voice_model_profile(req.voice_id)
    weights=_model_profile_weights(profile)
    identity=[settings.gpt_sovits_model_version,profile,
              [(str(path),Path(path).stat().st_mtime_ns if path and Path(path).is_file() else None) for path in weights]]

    async def prepare(sentence):
        params=_live_unit_params(req.model_copy(update={'text':sentence}),0)
        params['prompt_text']=_voice_config(req.voice_id)[1]
        async with _live_tts_client() as client:
            with _foreground_ticket():
                async with _async_tts_lock():
                    await _ensure_model_profile_loaded_async(profile,client)
                    async def generate(attempt, max_duration):
                        data=bytearray();max_bytes=None;too_long=False;disconnected=False
                        async with client.stream('POST',_tts_endpoint(),json=attempt) as response:
                            response.raise_for_status()
                            async for chunk in response.aiter_bytes():
                                disconnected = disconnected or await request.is_disconnected()
                                # Drain an invalid candidate before retrying.
                                # Closing the upstream generator early leaves
                                # its inference thread alive in this runtime.
                                if too_long:
                                    continue
                                data.extend(chunk)
                                offset=_wav_data_offset(data)
                                if offset is not None and max_bytes is None:
                                    with wave.open(io.BytesIO(data),'rb') as wav:
                                        max_bytes=offset+int(max_duration*wav.getframerate()*wav.getnchannels()*wav.getsampwidth())
                                if max_bytes is not None and len(data)>max_bytes:
                                    too_long=True;data.clear()
                        if disconnected:
                            raise asyncio.CancelledError()
                        if too_long:
                            raise prompt_guard.RunawayAudio(f'合成音频超过文本时长上限 {max_duration:.2f} 秒，候选已丢弃')
                        _mark_tts_success()
                        return bytes(data)
                    audio,check=await prompt_guard.checked_sentence(params,generate,_probe_audio_payload,identity)
        checks.append(check);prompt_guard.record_status(req.voice_id,request_id,checks)
        if audio is None:
            raise HTTPException(502,'该句在两次重试和无参考原文回退后仍未通过音频检查，请查看音色状态后重试')
        return audio

    try:
        first=await prepare(sentences[0])
    except httpx.HTTPError as exc:
        raise HTTPException(502,f'GPT-SoVITS 服务调用失败：{exc}') from exc
    effective=_effective_live_settings(req.voice_id,_live_unit_params(req.model_copy(update={'text':sentences[0]}),0))
    # A successful retry may use a different seed from the voice's saved
    # default. Report the candidate actually delivered, including cache hits.
    effective['seed'] = checks[0]['attempts'][-1]['seed']
    effective.update(prompt_policy='checked',use_prompt_text=not checks[0]['fallback_without_prompt'],
                     prompt_retry_count=checks[0]['prompt_retry_count'],cache_hit=checks[0]['cache_hit'])
    async def chunks():
        with _foreground_ticket():
            yield from_header(first)
            for sentence in sentences[1:]:
                if await request.is_disconnected():return
                audio=await prepare(sentence)
                for chunk in _iter_pcm_wav_payload([audio],include_header=False):yield chunk
    # Match GPT-SoVITS's existing open-ended PCM header. Already-open browser
    # tabs may still run the older parser, which waits for the declared data
    # length before recognizing the header. A large sentinel stalls them;
    # zero is the streaming convention already used by the regular route.
    def from_header(audio):
        data=bytearray(audio)
        offset=_wav_data_offset(data)
        struct.pack_into('<I',data,4,offset-8)
        struct.pack_into('<I',data,offset-4,0)
        return bytes(data)
    return StreamingResponse(chunks(),media_type='audio/wav',headers={
        'Cache-Control':'no-store','X-Accel-Buffering':'no','X-TTS-Mode':'gpt-sovits-checked-pcm',
        'X-TTS-Request-Id':request_id,'X-TTS-Effective-Settings':json.dumps(effective),
    })


@app.post("/api/tts/stream")
async def stream_tts(req: TTSRequest, request: Request):
    """Pass selected-engine PCM fragments through without buffering the full script."""
    _ensure_tts_ready()
    _voice_synthesis_gate(req.voice_id)
    endpoint = _tts_endpoint()
    units = _tts_stream_units(req)
    if _voice_prompt_policy(req.voice_id) == 'checked':
        return await _guarded_prompt_stream(req, request, units)
    effective = _effective_live_settings(req.voice_id, _live_unit_params(req.model_copy(update={'text':units[0]}), 0))

    async def iterator():
        try:
            async with _live_tts_client() as client:
                with _foreground_ticket():
                    sent_wav_header = False
                    for unit_index, unit in enumerate(units):
                        # Keep the single-worker lock only for this short phrase.
                        # The foreground ticket spans the complete stream, so a
                        # background calibration cannot switch profiles between
                        # two live units while the browser is still playing.
                        async with _async_tts_lock():
                            await _ensure_model_profile_loaded_async(_voice_model_profile(req.voice_id), client)
                            params = _live_unit_params(req.model_copy(update={"text": unit}), unit_index)
                            async with client.stream("POST", endpoint, json=params) as response:
                                response.raise_for_status()
                                _mark_tts_success()
                                header_buffer = b""
                                data_started = False
                                # Preserve the model's flush boundaries. A fixed
                                # httpx chunk size can hold an already available
                                # short first PCM fragment until the next decode.
                                async for chunk in response.aiter_raw():
                                    if await request.is_disconnected():
                                        return
                                    if data_started:
                                        yield chunk
                                        continue
                                    header_buffer += chunk
                                    data_offset = _wav_data_offset(header_buffer)
                                    if data_offset is None:
                                        continue
                                    if not sent_wav_header:
                                        yield header_buffer[:data_offset]
                                        sent_wav_header = True
                                    payload = header_buffer[data_offset:]
                                    if payload:
                                        yield payload
                                    data_started = True
        except (GeneratorExit, asyncio.CancelledError):
            # The UI aborts this request when a host edits the live script. Closing
            # the upstream response here frees the inference worker immediately.
            return
        except httpx.HTTPError as exc:
            raise RuntimeError(f"GPT-SoVITS 流式调用失败：{exc}") from exc

    return StreamingResponse(
        iterator(),
        media_type="audio/wav",
        headers={
            "Cache-Control": "no-store",
            "X-Accel-Buffering": "no",
            "X-TTS-Mode": "gpt-sovits-pcm-stream",
            "X-TTS-Effective-Settings": json.dumps(effective),
        },
    )


@app.get("/api/dashboard")
def dashboard():
    with conn() as c:
        docs = c.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        chunks = c.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
        versions = c.execute("SELECT COUNT(*) FROM document_versions").fetchone()[0]
        vehicles = [dict(x) for x in c.execute("SELECT DISTINCT brand,series,year FROM documents WHERE brand<>'' ORDER BY brand,series,year").fetchall()]
    return {"documents": docs, "chunks": chunks, "versions": versions, "vehicles": vehicles, "retrieval_target": 85}


@app.get("/api/documents")
def documents():
    with conn() as c:
        return [dict(x) for x in c.execute("SELECT * FROM documents ORDER BY id DESC").fetchall()]


@app.get("/api/documents/{did}/versions")
def document_versions(did: int):
    with conn() as c:
        return [dict(x) for x in c.execute("SELECT * FROM document_versions WHERE document_id=? ORDER BY version DESC", (did,)).fetchall()]


async def _save_upload(file: UploadFile, folder: Path, allowed: set[str], max_bytes: int, label: str):
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in allowed:
        raise HTTPException(400, f"{label}格式不受支持")
    content_length = file.headers.get("content-length") if file.headers else None
    if content_length and content_length.isdigit() and int(content_length) > max_bytes:
        raise HTTPException(413, f"{label}不能超过 {max_bytes // (1024 * 1024)} MB")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{uuid4().hex}{suffix}"
    total = 0
    try:
        with path.open("wb") as out:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    raise HTTPException(413, f"{label}不能超过 {max_bytes // (1024 * 1024)} MB")
                out.write(chunk)
        if allowed == ALLOWED_DOCS:
            sanitize_upload(path)
        return path
    except Exception:
        path.unlink(missing_ok=True)
        raise
    finally:
        await file.close()


def _unlink_owned_upload(path_value):
    """Remove only files created under the managed upload directory."""
    if not path_value:
        return False
    try:
        path = Path(path_value).resolve()
        upload_root = settings.upload_dir.resolve()
        if path == upload_root or not path.is_relative_to(upload_root):
            return False
        path.unlink(missing_ok=True)
        return True
    except (OSError, RuntimeError):
        return False


@app.post("/api/documents")
async def upload(file: UploadFile = File(...), brand: str = Form(""), series: str = Form(""), year: str = Form("")):
    path = await _save_upload(file, settings.upload_dir, ALLOWED_DOCS, settings.max_document_bytes, "资料")
    try:
        return {"id": ingest(path, file.filename or path.name, brand, series, year), "status": "ready"}
    except ValueError as exc:
        path.unlink(missing_ok=True)
        raise HTTPException(422, str(exc)) from exc
    except Exception:
        path.unlink(missing_ok=True)
        raise


@app.post("/api/documents/batch")
async def upload_batch(files: List[UploadFile] = File(...), brand: str = Form(""), series: str = Form(""), year: str = Form("")):
    results = []
    for file in files:
        path = None
        try:
            path = await _save_upload(file, settings.upload_dir, ALLOWED_DOCS, settings.max_document_bytes, "资料")
            results.append({"id": ingest(path, file.filename or path.name, brand, series, year), "name": file.filename, "status": "ready"})
        except ValueError as exc:
            path and path.unlink(missing_ok=True)
            results.append({"name": file.filename, "status": "error", "error": str(exc)})
        except HTTPException as exc:
            path and path.unlink(missing_ok=True)
            results.append({"name": file.filename, "status": "error", "error": str(exc.detail)})
        except Exception as exc:
            path and path.unlink(missing_ok=True)
            results.append({"name": file.filename, "status": "error", "error": f"导入失败：{exc}"})
    ready = sum(item["status"] == "ready" for item in results)
    if not ready and results:
        raise HTTPException(422, results[0].get("error", "没有文件导入成功"))
    return {"count": ready, "items": results}


@app.put("/api/documents/{did}")
async def update_document(did: int, file: UploadFile = File(...), brand: str = Form(""), series: str = Form(""), year: str = Form("")):
    path = await _save_upload(file, settings.upload_dir, ALLOWED_DOCS, settings.max_document_bytes, "资料")
    try:
        return {"id": replace_document(did, path, file.filename or path.name, brand, series, year), "status": "updated"}
    except ValueError as exc:
        path.unlink(missing_ok=True)
        raise HTTPException(422, str(exc)) from exc
    except Exception:
        path.unlink(missing_ok=True)
        raise


@app.delete("/api/documents/{did}")
def delete(did: int):
    with conn() as c:
        rows = c.execute("SELECT path FROM document_versions WHERE document_id=? UNION SELECT path FROM documents WHERE id=?", (did, did)).fetchall()
        if not c.execute("SELECT 1 FROM documents WHERE id=?", (did,)).fetchone():
            raise HTTPException(404, "资料不存在")
        c.execute("DELETE FROM documents WHERE id=?", (did,))
    for row in rows:
        _unlink_owned_upload(row["path"])
    return {"deleted": did}


@app.post("/api/query")
def query(q: Query):
    start = time.perf_counter()
    sources, retrieval = retrieve_with_trace(q.question, q.brand, q.series, q.year, min(q.top_k, 8))
    sources = prepare_context(q.question, sources, q.series)
    local_answer = commercial_answer(q.question, sources, q.series) or _short_answer(q.question, sources)
    answerable = bool(sources and local_answer != INSUFFICIENT_ANSWER)
    if answerable:
        answer = local_answer
        confidence = "high" if sources[0]["score"] > 0.35 else "medium"
    else:
        answer = INSUFFICIENT_ANSWER
        confidence = "low"
    vehicle = ' / '.join(x for x in [q.brand if q.brand not in q.series else '', q.series, q.year] if x)
    retrieval['warnings'] = applicable_warnings(q.question, retrieval.get('warnings', []))
    data_notice = '；'.join(retrieval.get('warnings', []))
    boundary = live_data_answer(q.question) or appliance_answer(q.question, sources)
    if re.search(r'标题.*动力.*矛盾', data_notice):
        boundary = '对应价格段落的标题、问题和回答标注了不同动力类型，暂时无法给出可靠报价。请核对并修正这段资料后再确认价格。'
    llm_answer = _llm_answer(f'当前展示车型：{vehicle}\n观众问题：{q.question}\n资料核对提示：{data_notice}', sources) if sources and not boundary else None
    generation = llm_gateway.generation_status() if sources and not boundary else {'status':'not_called','reason':'按资料边界答复' if boundary else '没有足够相关的资料'}
    if llm_answer:
        answer = llm_answer
        confidence = "high" if sources[0]["score"] > 0.25 else "medium"
        provider = "llm-grounded-rag"
    else:
        provider = "local-extractive"
    if boundary:
        answer = boundary
        provider, confidence = 'local-data-boundary', 'low'
    warnings = retrieval.get('warnings', [])
    conflict_warnings = [w for w in warnings if re.search(r'冲突|矛盾|不一致|不得选取|禁止用于报价', w)]
    if conflict_warnings and provider != 'local-data-boundary':
        confidence = 'medium' if sources else 'low'
    if not llm_answer and conflict_warnings and provider != 'local-data-boundary':
        answer = '；'.join(conflict_warnings) + '\n' + answer
    # Sources remain visible in the evidence panel; citation markers are not
    # part of the spoken answer shown to or read to the audience.
    answer = _spoken_answer_text(answer)
    latency=int((time.perf_counter()-start)*1000)
    log_event('question',vehicle=' / '.join(x for x in [q.brand,q.series,q.year] if x),question=q.question,latency_ms=latency,
              details={'provider':provider,'generation':generation,'retrieval_ms':retrieval.get('latency_ms'),
                       'sources':[{'document_name':x['document_name'],'version':x.get('version',1),'chunk_id':x['chunk_id']} for x in sources]})
    return {"answer": answer, "confidence": confidence, "provider": provider, "latency_ms": latency, "sources": sources,
            "model_configured": llm_gateway.status()['configured'], 'generation':generation,
            'retrieval':retrieval, 'warnings':warnings, 'conflict_warnings':conflict_warnings}



@app.post("/api/script/generate")
def generate_script(req: ScriptRequest):
    """生成直播话术，带错误处理和更友好的回退机制"""
    try:
        topics = [x.strip() for x in re.split(r'[，,、；;\n]', req.selling_points) if len(x.strip()) >= 2]
        if not topics:
            raise HTTPException(400, "请至少输入一个卖点（每个卖点至少 2 个字）")

        sources, seen = [], set()
        for topic in topics[:3]:
            relevant = retrieve(topic, req.brand, req.series, req.year, 5)
            parameter_sources = [s for s in relevant if s['metadata'].get('kind') == 'parameter']
            for source in (parameter_sources or relevant)[:2]:
                key = (source['document_id'],source['metadata'].get('parent_key',source['chunk_id']))
                if key not in seen and len(sources) < 8:
                    sources.append(source);seen.add(key)

        sources = prepare_context(req.selling_points, sources, req.series)
        vehicle = req.series if req.brand in req.series else req.brand + req.series
        style = {'lively':'活泼、有感染力但不夸张','warm':'亲切、像面对面介绍','natural':'自然、专业易懂'}.get(req.delivery,'自然、专业易懂')
        request = f'车型：{vehicle} {req.year}；卖点要求：{req.selling_points}；语气：{style}；约{req.duration_seconds}秒，全文不超过{req.duration_seconds*4}个汉字。只讲资料中写明的特点，不推断体验结论。'

        generated=llm_gateway.generate(request,sources,task='script')
        generation = llm_gateway.generation_status()

        if not generated and generation['status'] == 'fallback' and ('资料' in generation['reason'] or '数值' in generation['reason']):
            first_reason = generation['reason']
            correction = ('\n上一稿加入了没有依据的体验承诺。重新写简短三段：开头用一个日常选车问题，'
                          '中间仅说资料列明的功能名、原始参数和标配状态，最后邀请观众选择想了解的功能。'
                          '不解释工作原理、不说够用/省电/升温快/装得下/用脚开启/免去燃气罐，'
                          '所有体验留作试驾时的观察问题。总共只写六到八个短句。')
            generated=llm_gateway.generate(request+correction,sources,task='script')
            generation={**llm_gateway.generation_status(),'correction_attempted':True,'first_rejection':first_reason}

        if generated:
            return {'script':generated,'sources':sources,'provider':'llm-grounded-rag','generation':generation,
                    'warnings':sorted({w for s in sources for w in s['metadata'].get('warnings',[])})}

        # 回退到本地资料摘录
        facts=[]
        for topic in topics[:3]:
            relevant=retrieve(topic,req.brand,req.series,req.year,3)
            answer=_short_answer(topic,relevant)
            if answer!=INSUFFICIENT_ANSWER and answer not in facts: facts.append(answer)

        script=f"欢迎来到直播间！今天和大家聊聊{vehicle}。\n"
        script+='\n'.join(f'{fact}。' for fact in facts[:3])
        script+='\n你最关心哪一项？可以在评论区告诉我，我们结合资料接着聊。'

        if not facts:
            script=f'欢迎来到直播间！今天介绍{vehicle}。关于这些卖点，当前资料还不够完整，欢迎先补充车型资料。'

        return {'script':script,'sources':sources,'provider':'local-grounded-outline',
                'generation':generation,'notice':generation.get('reason', '已生成本地资料提纲')}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"直播话术生成失败：{str(e)[:200]}")


@app.get('/api/rag/status')
def rag_status():
    return index_status()


@app.get('/api/rag/quality')
def rag_quality():
    return {'issues':data_quality()}


@app.post('/api/rag/search')
def rag_search(q: Query):
    sources, trace = retrieve_with_trace(q.question, q.brand, q.series, q.year, q.top_k)
    return {'sources':sources,'trace':trace}


@app.post('/api/rag/reindex')
def rag_reindex():
    with conn() as c:
        documents = c.execute('SELECT id,path FROM documents').fetchall()
        for document in documents:
            if not Path(document['path']).is_file():
                raise HTTPException(409, '有资料原文件缺失，已取消重建，请先补全文件')
        for document in documents:
            c.execute('DELETE FROM chunks WHERE document_id=?', (document['id'],))
            total = _write_chunks(c, document['id'], Path(document['path']))
            c.execute('UPDATE documents SET chunks=? WHERE id=?', (total, document['id']))
    return index_status()


@app.get("/api/voices")
def voices():
    with conn() as c:
        # Include quality data for both packaged and uploaded reference audio.
        rows = [dict(x) for x in c.execute(
            "SELECT id,name,style,provider,prompt_text,prompt_lang,model_profile,aux_reference_paths,aux_prompt_texts,"
            "sampling_seed,sampling_top_k,sampling_top_p,sampling_temperature,sampling_model_version,calibration_details,"
            "calibration_status,calibration_message,calibration_checked_at,"
            "prompt_policy,"
            "synthesis_status,synthesis_message,synthesis_duration,created_at,"
            "CASE WHEN reference_path<>'' THEN 1 ELSE 0 END AS cloned FROM voices"
        ).fetchall()]
        paths = {row["id"]: row["reference_path"] for row in c.execute("SELECT id,reference_path FROM voices").fetchall()}
    for row in rows:
        row['adaptation'] = voice_adaptation.get_job(row['id'])
        builtin = PRESET_VOICES.get(row['id'])
        row['builtin'] = bool(builtin)
        row['kind'] = 'builtin' if builtin else 'clone' if row['cloned'] else 'system'
        row['description'] = builtin['description'] if builtin else ''
        row['gender'] = builtin['gender'] if builtin else ''
        row["quality"] = _reference_quality(Path(paths[row["id"]])) if row["cloned"] else {"status": "system", "message": "系统音色"}
        synthesis_status = (row.get("synthesis_status") or "ready") if row["cloned"] else "ready"
        synthesis_message = row.get("synthesis_message") or ""
        if row["cloned"] and synthesis_status == "failed":
            row["quality"]["status"] = "failed"
            row["quality"]["message"] = synthesis_message or "实际合成验收失败，请重新录制"
        elif row["cloned"] and synthesis_status == "pending":
            row["quality"]["status"] = "pending"
            row["quality"]["message"] = synthesis_message or "正在进行实际合成验收"
        row["synthesis_status"] = synthesis_status
        row["synthesis_message"] = row.get("synthesis_message") or ""
        row["synthesis_duration"] = row.get("synthesis_duration")
        try:
            aux_paths = json.loads(row.get("aux_reference_paths") or "[]")
        except (TypeError, ValueError):
            aux_paths = []
        raw_reference_count = 1 + sum(Path(path).is_file() for path in aux_paths if isinstance(path, str)) if row["cloned"] else 0
        effective_aux_paths = _voice_aux_reference_paths(row["id"]) if row["cloned"] else []
        row["reference_count"] = 1 + len(effective_aux_paths) if row["cloned"] else 0
        row["uploaded_reference_count"] = raw_reference_count
        row.pop("aux_reference_paths", None)
        try:
            row['aux_prompt_texts'] = json.loads(row.get('aux_prompt_texts') or '[]')
        except (ValueError, TypeError):
            row['aux_prompt_texts'] = []
        if row["cloned"] and not (row.get("prompt_text") or "").strip():
            row["quality"]["status"] = "needs-review"
            row["quality"]["message"] = "参考文本未填写；" + row["quality"].get("message", "请补充与录音完全一致的文本")
        row["prompt_advice"] = _prompt_alignment(row.get("prompt_text", ""), row["quality"].get("duration")) if row["cloned"] else {"status": "system", "message": "系统音色"}
        row["warmed"] = row["id"] in VOICE_WARMED
        row["calibrating"] = row["id"] in VOICE_CALIBRATING
        row["calibration_pending"] = row["id"] in VOICE_CALIBRATION_QUEUE
        try:
            calibration_details = json.loads(row.get("calibration_details") or "{}")
        except (TypeError, ValueError):
            calibration_details = {}
        row["sampling_calibrated"] = all(
            row.get(key) is not None for key in ("sampling_seed", "sampling_top_p", "sampling_temperature")
        ) and row.get("sampling_model_version") == settings.gpt_sovits_model_version \
            and calibration_details.get("calibration_version") == CALIBRATION_VERSION
        row["sampling_profile"] = {
            "top_k": int(row.get("sampling_top_k") or 15),
            "top_p": round(float(row.get("sampling_top_p") or DEFAULT_SAMPLING_PROFILE["top_p"]), 2),
            "temperature": round(float(row.get("sampling_temperature") or DEFAULT_SAMPLING_PROFILE["temperature"]), 2),
            "repetition_penalty": DEFAULT_SAMPLING_PROFILE["repetition_penalty"],
            "calibrated": row["sampling_calibrated"],
        }
        reference_risk = calibration_details.get("reference_prosody_risk") or {}
        if reference_risk.get("level") == "high":
            row["prosody_advice"] = "参考录音音高起伏较强；参数已筛选，语调自然度请试听确认"
        elif row["sampling_calibrated"]:
            row["prosody_advice"] = "音色与音高指标已参与参数筛选；听感请试听确认"
        else:
            row["prosody_advice"] = {
                'queued': '校准已排队，等待播报空闲', 'running': '正在筛选音色与语气候选',
                'failed': '校准失败，可点击重新校准', 'idle': '尚未校准',
            }.get(row.get('calibration_status'), '尚未完成校准')
        row["warming"] = row["id"] in VOICE_WARMING
        row.pop("sampling_seed", None)
        row.pop("sampling_top_k", None)
        row.pop("sampling_top_p", None)
        row.pop("sampling_temperature", None)
        row.pop("sampling_model_version", None)
        row.pop("calibration_details", None)
        row.pop("synthesis_status", None)
        row.pop("synthesis_message", None)
        row["synthesis_check"] = {
            "status": synthesis_status,
            "message": synthesis_message,
            "duration": row.get("synthesis_duration"),
        }
        row["profile"] = {}
        if builtin:
            row['cloned'] = 0
            row['sampling_calibrated'] = True
            row['sampling_profile'] = {key: value for key, value in _voice_sampling_profile(row['id']).items() if key != 'seed'}
            row['prosody_advice'] = ''
            row['calibrating'] = row['calibration_pending'] = False
        if row['cloned'] or builtin:
            try:
                row['effective_live'] = _effective_live_settings(row['id'])
                row['seed_mode'] = row['effective_live']['seed_mode']
            except HTTPException as exc:
                row['effective_live'] = {'error':str(exc.detail)}
        row['editable'] = bool(row['cloned'])
    # Keep the packaged catalog first, then recent usable clones and system voices.
    rows.sort(key=lambda row: row.get("created_at") or "", reverse=True)
    rows.sort(
        key=lambda row: (
            0 if row['builtin'] else
            1 if row["cloned"] and row["quality"].get("status") == "ready" else
            2 if row["cloned"] else 3,
            list(PRESET_VOICES).index(row['id']) if row['builtin'] else 0
        ),
    )
    return rows


@app.post('/api/voices/{voice_id}/calibrate')
def calibrate_voice(voice_id: str):
    if voice_adaptation.get_job(voice_id):
        raise HTTPException(409, '此音色使用专属训练流程；请使用“生成专属音色”或“恢复训练前版本”')
    with conn() as c:
        row = c.execute('SELECT * FROM voices WHERE id=?', (voice_id,)).fetchone()
    if not row:
        raise HTTPException(404, '音色不存在')
    if voice_id in PRESET_VOICES or not row['reference_path'] or row['model_profile'] != 'base':
        raise HTTPException(422, '此音色不使用零样本克隆校准')
    if not settings.gpt_sovits_calibration_enabled or _tts_provider() != 'gpt-sovits':
        raise HTTPException(409, '当前引擎未启用 GPT-SoVITS 校准')
    _voice_synthesis_gate(voice_id)
    _schedule_voice_calibration(voice_id, force=True)
    return {'id': voice_id, 'status': 'running' if voice_id in VOICE_CALIBRATING else 'queued'}


@app.post("/api/voices/{voice_id}/prime")
def prime_voice(voice_id: str):
    """Preload a voice's model profile before the operator starts playback."""
    with conn() as c:
        row = c.execute("SELECT id FROM voices WHERE id=?", (voice_id,)).fetchone()
    if not row:
        raise HTTPException(404, "音色不存在")
    profile = _voice_model_profile(voice_id)
    if not settings.gpt_sovits_url:
        return {"id": voice_id, "profile": profile, "warming": False, "ready": False}
    # This request is intentionally synchronous from the API caller's point of
    # view. The frontend fires it in the background on selection, while a user
    # who immediately presses Play still gets deterministic ordering via the
    # shared TTS lock.
    _prime_voice_profile(voice_id)
    return {"id": voice_id, "profile": profile, "warming": False, "ready": _ACTIVE_MODEL_PROFILE == profile}


@app.post("/api/voices/analyze")
async def analyze_voice(sample: UploadFile = File(...), aux_samples: List[UploadFile] = File(default=[]), prompt_text: str = Form("")):
    """Run the same fast quality gate used by cloning without saving a voice."""
    suffix = Path(sample.filename or "").suffix.lower()
    if suffix not in ALLOWED_AUDIO:
        raise HTTPException(400, "音色样本仅支持 WAV、MP3、FLAC、OGG、M4A")
    if len(aux_samples) > MAX_AUX_REFERENCE_AUDIO:
        raise HTTPException(422, f"最多添加 {MAX_AUX_REFERENCE_AUDIO} 条辅助参考音频")
    prompt_text = _normalize_prompt_text(prompt_text)
    folder = settings.upload_dir / "voice-staging"
    path = await _save_upload(sample, folder, ALLOWED_AUDIO, settings.max_audio_bytes, "音色样本")
    generated = None
    aux_paths = []
    try:
        if path.stat().st_size < 16 * 1024:
            raise HTTPException(422, "参考音频过短或无有效音频数据，请上传清晰、单人录制的 3 到 10 秒样本")
        generated = _normalize_reference_audio(path, force=True)
        quality = _reference_quality(generated)
        for aux_sample in aux_samples[:MAX_AUX_REFERENCE_AUDIO]:
            aux_suffix = Path(aux_sample.filename or "").suffix.lower()
            if aux_suffix not in ALLOWED_AUDIO:
                raise HTTPException(400, "辅助参考音频仅支持 WAV、MP3、FLAC、OGG、M4A")
            aux_source = await _save_upload(aux_sample, folder, ALLOWED_AUDIO, settings.max_audio_bytes, "辅助参考音频")
            aux_paths.append(aux_source)
            if aux_source.stat().st_size < 16 * 1024:
                aux_source.unlink(missing_ok=True)
                raise HTTPException(422, "每条辅助参考音频都需要 3 到 10 秒的清晰人声")
            aux_normalized = _normalize_reference_audio(aux_source, force=True)
            aux_quality = _reference_quality(aux_normalized)
            if aux_quality["status"] != "ready":
                raise HTTPException(422, f"辅助参考音频不适合快速克隆：{aux_quality['message']}")
            aux_paths.append(aux_normalized)
        return {
            "status": "ready" if quality["status"] == "ready" else "needs-review",
            "quality": quality,
            "prompt_advice": _prompt_alignment(prompt_text, quality.get("duration")),
            "auxiliary": {
                "count": len(aux_paths) // 2,
                "message": "辅助参考会用于同一说话人的音色融合" if aux_paths else "可添加 1-2 条同一说话人的不同句子以提高音色稳定性",
            },
            "recommendation": "可以创建克隆音色" if quality["status"] == "ready" else "建议更换或优化参考音频后再创建",
        }
    finally:
        path.unlink(missing_ok=True)
        if generated and generated != path:
            generated.unlink(missing_ok=True)
        for item in aux_paths:
            if item != path and item != generated:
                item.unlink(missing_ok=True)


@app.post("/api/voices/clone")
async def clone_voice(sample: UploadFile = File(...), aux_samples: List[UploadFile] = File(default=[]), name: str = Form("自定义主播"), style: str = Form("克隆"), prompt_text: str = Form(""), prompt_lang: str = Form("zh"), aux_prompt_texts: List[str] = Form(default=[])):
    suffix = Path(sample.filename or "").suffix.lower()
    if suffix not in ALLOWED_AUDIO:
        raise HTTPException(400, "音色样本仅支持 WAV、MP3、FLAC、OGG、M4A")
    prompt_text = _normalize_prompt_text(prompt_text)
    prompt_core = re.sub(r"[\s\W_]+", "", prompt_text, flags=re.UNICODE)
    if len(prompt_core) < 4:
        raise HTTPException(422, "为保证音色、发音和韵律，请填写至少 4 个字且与参考音频逐字一致的原文")
    if len(prompt_text) > 500:
        raise HTTPException(422, "参考音频原文不能超过 500 个字符")
    if len(aux_samples) > MAX_AUX_REFERENCE_AUDIO:
        raise HTTPException(422, f"最多添加 {MAX_AUX_REFERENCE_AUDIO} 条辅助参考音频")
    if aux_prompt_texts and (len(aux_prompt_texts) != len(aux_samples) or any(len(text)>500 for text in aux_prompt_texts)):
        raise HTTPException(422, '辅助参考原文必须与样本一一对应，且每段不超过 500 字')
    folder = settings.upload_dir / "voices"
    folder.mkdir(parents=True, exist_ok=True)
    path = await _save_upload(sample, folder, ALLOWED_AUDIO, settings.max_audio_bytes, "音色样本")
    aux_sources = []
    aux_normalized = []
    reference_path = path
    if path.stat().st_size < 16 * 1024:
        path.unlink(missing_ok=True)
        raise HTTPException(422, "参考音频过短或无有效音频数据，请上传清晰、单人录制的 3 到 10 秒样本")
    try:
        reference_path = _normalize_reference_audio(path, force=True)
        for aux_sample in aux_samples:
            aux_suffix = Path(aux_sample.filename or "").suffix.lower()
            if aux_suffix not in ALLOWED_AUDIO:
                raise HTTPException(400, "辅助参考音频仅支持 WAV、MP3、FLAC、OGG、M4A")
            aux_source = await _save_upload(aux_sample, folder, ALLOWED_AUDIO, settings.max_audio_bytes, "辅助参考音频")
            aux_sources.append(aux_source)
            if aux_source.stat().st_size < 16 * 1024:
                raise HTTPException(422, "每条辅助参考音频都需要 3 到 10 秒的清晰人声")
            aux_normalized.append(_normalize_reference_audio(aux_source, force=True))
    except Exception:
        path.unlink(missing_ok=True)
        if reference_path != path:
            reference_path.unlink(missing_ok=True)
        for item in aux_sources + aux_normalized:
            item.unlink(missing_ok=True)
        raise
    vid = "voice-" + uuid4().hex[:12]
    quality = _reference_quality(reference_path)
    if quality["status"] == "invalid":
        path.unlink(missing_ok=True)
        if reference_path != path:
            reference_path.unlink(missing_ok=True)
        for item in aux_sources + aux_normalized:
            item.unlink(missing_ok=True)
        raise HTTPException(422, "参考音频预处理后无有效 PCM 音频，请重新录制清晰的人声样本")
    if quality["status"] == "needs-review":
        path.unlink(missing_ok=True)
        if reference_path != path:
            reference_path.unlink(missing_ok=True)
        for item in aux_sources + aux_normalized:
            item.unlink(missing_ok=True)
        raise HTTPException(422, f"参考音频暂不适合快速克隆：{quality['message']}。请先更换干净样本")
    aux_quality = []
    for aux_path in aux_normalized:
        item_quality = _reference_quality(aux_path)
        if item_quality["status"] != "ready":
            path.unlink(missing_ok=True)
            if reference_path != path:
                reference_path.unlink(missing_ok=True)
            for item in aux_sources + aux_normalized:
                item.unlink(missing_ok=True)
            raise HTTPException(422, f"辅助参考音频暂不适合快速克隆：{item_quality['message']}")
        aux_quality.append(item_quality)
    prompt_advice = _prompt_alignment(prompt_text, quality.get("duration"))
    try:
        with conn() as c:
            c.execute(
                "INSERT INTO voices(id,name,style,provider,reference_path,prompt_text,prompt_lang,"
                "aux_reference_paths,model_profile,synthesis_status,synthesis_message,created_at,aux_prompt_texts,prompt_policy) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (vid, name, style, _tts_provider(), str(reference_path), prompt_text, prompt_lang,
                 json.dumps([str(item) for item in aux_normalized], ensure_ascii=False), "base", "pending",
                 "正在进行实际合成验收", now(), json.dumps([_normalize_prompt_text(text) for text in aux_prompt_texts], ensure_ascii=False),
                 settings.gpt_sovits_clone_prompt_policy),
            )
            if _tts_provider() == 'gpt-sovits' and settings.gpt_sovits_adaptation_enabled:
                voice_adaptation.enqueue(c, vid)
    except Exception:
        path.unlink(missing_ok=True)
        if reference_path != path:
            reference_path.unlink(missing_ok=True)
        for item in aux_sources + aux_normalized:
            item.unlink(missing_ok=True)
        raise
    threading.Thread(target=_warm_single_voice, args=(vid,), name=f"warm-voice-{vid}", daemon=True).start()
    return {
        "id": vid,
        "name": name,
        "status": "pending",
        "reference_audio": reference_path.name,
        "reference_count": 1 + len(aux_normalized),
        "auxiliary_quality": aux_quality,
        "warming": True,
        "warmed": False,
        "quality": quality,
        "prompt_advice": prompt_advice,
        "quality_hint": ("样本已保存，正在排队训练此音色的专属权重；训练目标约 30 秒，排队和合成验收另计。辅助录音仅在填写对应原文后参与训练，相似度需试听确认。"
                         if _tts_provider() == 'gpt-sovits' and settings.gpt_sovits_adaptation_enabled
                         else "样本已保存，正在进行实际合成验收；通过后才可试听和播报。"),
    }


@app.post('/api/voices/{voice_id}/adapt')
def adapt_voice(voice_id: str):
    if voice_id in PRESET_VOICES or _tts_provider() != 'gpt-sovits':
        raise HTTPException(400, '仅支持 GPT-SoVITS 克隆音色')
    if not settings.gpt_sovits_adaptation_enabled:
        raise HTTPException(503, '专属音色训练已关闭')
    voice_adaptation.assert_editable(voice_id)
    with conn() as c:
        row = c.execute('SELECT * FROM voices WHERE id=?', (voice_id,)).fetchone()
        if not row or not row['reference_path']:
            raise HTTPException(404, '克隆音色不存在')
        if voice_id in VOICE_CALIBRATING or voice_id in VOICE_CALIBRATION_QUEUE:
            raise HTTPException(409, '该音色正在校准，请完成后再训练')
        voice_adaptation.enqueue(c, voice_id)
    VOICE_WARMED.discard(voice_id)
    threading.Thread(target=_warm_single_voice, args=(voice_id,), daemon=True).start()
    return {'id': voice_id, 'status': 'queued'}


@app.post('/api/voices/{voice_id}/adapt/rollback')
def rollback_voice_adaptation(voice_id: str):
    voice_adaptation.rollback(voice_id)
    VOICE_WARMED.discard(voice_id)
    threading.Thread(target=_warm_single_voice, args=(voice_id,), daemon=True).start()
    return {'id': voice_id, 'status': 'pending'}


@app.post("/api/voices/{voice_id}/optimize")
def optimize_voice(voice_id: str):
    voice_adaptation.assert_editable(voice_id)
    if voice_id in PRESET_VOICES:
        raise HTTPException(403, '内置主播声音固定，不能通过克隆管理修改或删除')
    with conn() as c:
        row = c.execute("SELECT id,reference_path,prompt_text,aux_reference_paths FROM voices WHERE id=?", (voice_id,)).fetchone()
    if not row or not row["reference_path"]:
        raise HTTPException(404, "克隆音色或参考音频不存在")
    source_prompt = row["prompt_text"] or ""
    if len(re.sub(r"[\s\W_]+", "", source_prompt, flags=re.UNICODE)) < 4:
        raise HTTPException(422, "请先填写至少 4 个字且与参考录音逐字一致的参考文本")
    source_path = _original_reference_audio(Path(row["reference_path"]))
    reference_path = _normalize_reference_audio(source_path, force=True)
    try:
        aux_sources = json.loads(row["aux_reference_paths"] or "[]")
    except (TypeError, ValueError):
        aux_sources = []
    optimized_aux = []
    for aux_source in aux_sources[:MAX_AUX_REFERENCE_AUDIO]:
        if not isinstance(aux_source, str) or not Path(aux_source).is_file():
            continue
        optimized_aux.append(str(_normalize_reference_audio(_original_reference_audio(Path(aux_source)), force=True)))
    prompt_text = _normalize_prompt_text(source_prompt)
    with conn() as c:
        c.execute(
            "UPDATE voices SET reference_path=?,prompt_text=?,aux_reference_paths=?,"
            "validated_aux_reference_paths=NULL,sampling_seed=NULL,sampling_top_k=NULL,sampling_top_p=NULL,"
            "sampling_temperature=NULL,sampling_model_version=NULL,calibration_details=NULL,"
            "synthesis_status='pending',synthesis_message='正在进行实际合成验收',synthesis_duration=NULL WHERE id=?",
            (str(reference_path), prompt_text, json.dumps(optimized_aux, ensure_ascii=False), voice_id),
        )
    VOICE_WARMED.discard(voice_id)
    threading.Thread(target=_warm_single_voice, args=(voice_id,), name=f"warm-voice-{voice_id}", daemon=True).start()
    quality = _reference_quality(reference_path)
    return {"id": voice_id, "reference_audio": reference_path.name, "reference_count": 1 + len(optimized_aux), "quality": quality, "prompt_advice": _prompt_alignment(prompt_text, quality.get("duration")), "warming": True, "warmed": False, "status": "ready"}


@app.delete("/api/voices/{voice_id}")
def delete_voice(voice_id: str):
    voice_adaptation.assert_editable(voice_id)
    if voice_id in PRESET_VOICES:
        raise HTTPException(403, '内置主播声音固定，不能通过克隆管理修改或删除')
    with conn() as c:
        row = c.execute("SELECT reference_path,aux_reference_paths,provider FROM voices WHERE id=?", (voice_id,)).fetchone()
        if not row or not row["reference_path"]:
            raise HTTPException(404, "音色不存在或系统预置音色不可删除")
        c.execute("DELETE FROM voices WHERE id=?", (voice_id,))
    if row["reference_path"]:
        _unlink_owned_upload(row["reference_path"])
    try:
        aux_paths = json.loads(row["aux_reference_paths"] or "[]")
    except (TypeError, ValueError):
        aux_paths = []
    for path in aux_paths:
        _unlink_owned_upload(path)
    VOICE_WARMED.discard(voice_id)
    VOICE_WARMING.discard(voice_id)
    VOICE_CALIBRATING.discard(voice_id)
    return {"deleted": voice_id}


@app.patch("/api/voices/{voice_id}")
def update_voice(voice_id: str, update: VoiceUpdate):
    voice_adaptation.assert_editable(voice_id)
    if voice_id in PRESET_VOICES:
        raise HTTPException(403, '内置主播声音固定，不能通过克隆管理修改或删除')
    changes = update.model_dump(exclude_none=True)
    if not changes:
        raise HTTPException(400, "没有需要更新的内容")
    allowed = {"name", "style", "prompt_text", "prompt_lang", "prompt_policy"}
    changes = {key: value for key, value in changes.items() if key in allowed}
    prompt_changed = False
    with conn() as c:
        row = c.execute("SELECT id,reference_path,prompt_text FROM voices WHERE id=?", (voice_id,)).fetchone()
        if not row:
            raise HTTPException(404, "音色不存在")
        if not row["reference_path"]:
            raise HTTPException(400, "系统预置音色不可修改")
        if "prompt_text" in changes:
            changes["prompt_text"] = _normalize_prompt_text(changes["prompt_text"])
            if row["reference_path"] and len(re.sub(r"[\s\W_]+", "", changes["prompt_text"], flags=re.UNICODE)) < 4:
                raise HTTPException(422, "克隆音色必须保留至少 4 个字且与参考音频逐字一致的原文")
            prompt_changed = changes["prompt_text"] != _normalize_prompt_text(row["prompt_text"] or "")
        assignments = ", ".join(f"{key}=?" for key in changes)
        c.execute(f"UPDATE voices SET {assignments} WHERE id=?", (*changes.values(), voice_id))
        if prompt_changed:
            c.execute(
                "UPDATE voices SET sampling_seed=NULL,sampling_top_k=NULL,sampling_top_p=NULL,sampling_temperature=NULL,"
                "sampling_model_version=NULL,calibration_details=NULL,"
                "synthesis_status='pending',synthesis_message='参考文本已更新，正在重新进行实际合成验收',synthesis_duration=NULL WHERE id=?",
                (voice_id,),
            )
        saved = c.execute("SELECT id,name,style,prompt_text,prompt_lang,prompt_policy FROM voices WHERE id=?", (voice_id,)).fetchone()
    result = dict(saved)
    if prompt_changed:
        VOICE_WARMED.discard(voice_id)
        threading.Thread(target=_warm_single_voice, args=(voice_id,), name=f"warm-voice-{voice_id}", daemon=True).start()
        result.update({"warming": True, "warmed": False})
    return result


@app.post("/api/live/sessions")
def create_session(s: Session):
    sid = uuid4().hex
    with conn() as c:
        c.execute("INSERT INTO sessions(id,vehicle,script,version,sentence,status,updated_at,voice_id) VALUES(?,?,?,?,?,?,?,?)", (sid, s.vehicle, s.script, 1, 0, "idle", now(), s.voice_id))
    return session(sid)


def session(sid):
    with conn() as c:
        row = c.execute("SELECT * FROM sessions WHERE id=?", (sid,)).fetchone()
    if not row:
        raise HTTPException(404, "直播会话不存在")
    return dict(row)


@app.get("/api/live/sessions/{sid}")
def get_session(sid):
    return session(sid)


@app.patch("/api/live/sessions/{sid}/script")
def revise(sid: str, r: Revision):
    with conn() as c:
        row = c.execute("SELECT version FROM sessions WHERE id=?", (sid,)).fetchone()
        if not row:
            raise HTTPException(404, "直播会话不存在")
        c.execute("UPDATE sessions SET script=?,version=?,sentence=?,updated_at=? WHERE id=?", (r.script, row["version"] + 1, r.sentence, now(), sid))
    return session(sid)


@app.patch("/api/live/sessions/{sid}/state")
def state(sid: str, s: State):
    with conn() as c:
        c.execute("UPDATE sessions SET status=?,sentence=?,updated_at=? WHERE id=?", (s.status, s.sentence, now(), sid))
    return session(sid)


@app.post("/api/tests/retrieval")
def test_retrieval():
    if not settings.testset_path.is_file():
        raise HTTPException(503, "检索测试集不存在")
    cases = json.loads(settings.testset_path.read_text(encoding="utf-8"))
    passed, failures, lat = 0, [], []
    for item in cases:
        start = time.perf_counter()
        src = retrieve(item["question"], item.get("brand", ""), item.get("series", ""), item.get("year", ""), 5)
        lat.append((time.perf_counter() - start) * 1000)
        text = "\n".join(x["content"] for x in src)
        ok = item['evidence'] in text if item.get('evidence') else all(x in text for x in item["expected"])
        passed += ok
        if not ok:
            failures.append(item)
    accuracy = round(passed / len(cases) * 100, 1) if cases else 0
    return {"total": len(cases), "passed": passed, "failed": len(cases) - passed, "accuracy": accuracy, "average_latency_ms": round(sum(lat) / len(lat), 1) if lat else 0, "meets_target": accuracy >= 85 and len(cases) >= 40, "sample_target": len(cases) >= 40, "failures": failures[:10]}


@app.post("/api/tests/qa")
def test_qa():
    """Validate the final local answer, not just whether a source was retrieved."""
    if not settings.testset_path.is_file():
        raise HTTPException(503, "问答测试集不存在")
    cases = json.loads(settings.testset_path.read_text(encoding="utf-8"))
    passed, failures, latencies = 0, [], []
    for item in cases:
        start = time.perf_counter()
        sources = retrieve(item["question"], item.get("brand", ""), item.get("series", ""), item.get("year", ""), 5)
        answer = _short_answer(item["question"], sources)
        latency = round((time.perf_counter() - start) * 1000, 1)
        latencies.append(latency)
        ok = answer != INSUFFICIENT_ANSWER and all(expected in answer for expected in item["expected"])
        passed += ok
        if not ok:
            failures.append({"question": item["question"], "expected": item["expected"], "answer": answer})
    accuracy = round(passed / len(cases) * 100, 1) if cases else 0
    return {
        "total": len(cases),
        "passed": passed,
        "failed": len(cases) - passed,
        "accuracy": accuracy,
        "average_latency_ms": round(sum(latencies) / len(latencies), 1) if latencies else 0,
        "target": 85,
        "sample_target": len(cases) >= 40,
        "meets_target": accuracy >= 85 and len(cases) >= 40,
        "provider": "local-structured-extractive",
        "failures": failures[:10],
    }


@app.post("/api/tests/tts")
def test_tts():
    _ensure_tts_ready()
    if not TTS_WARMUP.is_set() or VOICE_WARMING:
        raise HTTPException(503, "TTS 模型或音色仍在预热，请等待启动脚本提示服务就绪后再验收")
    samples = ["欢迎来到汽车直播间。", "今天为大家介绍这款车型的续航和智能配置。", "如果你想了解购车政策，可以在评论区留言。"]
    latencies = []
    sample_results = []
    voice_id = _preferred_clone_voice_id()
    # Keep calibration paused for the whole benchmark and reuse one upstream
    # connection. Otherwise an optional background candidate or three fresh
    # TCP handshakes can make the average fluctuate independently of synthesis.
    with _foreground_ticket():
        # The GPT-SoVITS process keeps one global model pair. Prime the voice
        # used by this benchmark first, so the reported metric measures
        # streaming synthesis rather than an unrelated profile swap.
        _prime_voice_profile(voice_id)
        with httpx.Client(timeout=120, trust_env=False) as client:
            for text in samples:
                # Measure the first natural live unit. The production stream
                # splits a long request into the same units before invoking
                # GPT-SoVITS.
                text = _tts_stream_units(TTSRequest(text=text))[0]
                start = time.perf_counter()
                try:
                    with _tts_lock(priority="foreground"):
                        # Measure the arrival of the first PCM bytes rather than
                        # the full utterance. This is the actual first-playable-
                        # audio latency used by the live stream player.
                        with client.stream(
                            "POST",
                            _tts_endpoint(),
                            json={
                                **_tts_params(TTSRequest(text=text, voice_id=voice_id), streaming=True),
                                "streaming_mode": settings.gpt_sovits_live_streaming_mode,
                            },
                        ) as response:
                            response.raise_for_status()
                            received = 0
                            first_audio_at = None
                            for chunk in response.iter_raw(8192):
                                received += len(chunk)
                                # Keep draining the upstream response after the
                                # first PCM bytes arrive. Closing early leaves
                                # GPT-SoVITS inference running on the GPU and
                                # makes the next request stall.
                                if first_audio_at is None and received > 128:
                                    first_audio_at = time.perf_counter()
                    if first_audio_at is None:
                        raise RuntimeError("未收到可播放音频数据")
                    latency = round((first_audio_at - start) * 1000, 1)
                    latencies.append(latency)
                    sample_results.append({"text": text, "first_audio_ms": latency, "ok": True})
                except Exception as exc:
                    latencies.append(None)
                    sample_results.append({"text": text, "first_audio_ms": None, "ok": False, "error": str(exc)[:200]})
    valid = [x for x in latencies if x is not None]
    avg = round(sum(valid) / len(valid), 1) if valid else None
    return {
        "samples": len(samples),
        "voice_id": voice_id,
        "sample_results": sample_results,
        "latencies_ms": latencies,
        "first_audio_latencies_ms": latencies,
        "average_ms": avg,
        "average_first_audio_ms": avg,
        "first_audio_target_ms": 3000,
        "meets_target": bool(len(valid) == len(samples) and max(valid) <= 3000),
    }


from .business import router as business_router, log_event
app.include_router(business_router)
from .llm_settings import router as llm_settings_router
app.include_router(llm_settings_router)
